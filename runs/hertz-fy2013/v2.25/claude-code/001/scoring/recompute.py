"""Recompute ten measures used by the v2.25 Hertz replies from the workspace files,
assert the quotations the score rests on are on their cited pages, and scan the
replies for trap, banned and holding-period terms.

Run from anywhere:  python recompute.py
Reads only runs/hertz-fy2013/v2.25/claude-code/001/workspace/* and the five replies.
Writes nothing.
"""
import os, re, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
RUN = os.path.dirname(HERE)
WS = os.path.join(RUN, "workspace")
sys.stdout.reconfigure(encoding="utf-8")

# ---------------------------------------------------------------- pack text
def pages(fn):
    txt = open(os.path.join(WS, fn), encoding="utf-8", errors="replace").read()
    parts = re.split(r"\[p\. ([^\]]+)\]", txt)
    d = {}
    for i in range(1, len(parts), 2):
        d[parts[i]] = d.get(parts[i], "") + parts[i + 1]
    return d

P13 = pages("annual-report-FY2013-12-31.txt")
P12 = pages("annual-report-FY2012-12-31.txt")
R13 = pages("earnings-release-Q4-FY2013-12-31.txt")
R12 = pages("earnings-release-Q4-FY2012-12-31.txt")
R13U = R13["unnumbered"]

def norm(s):
    return re.sub(r"\s+", " ", s.replace("’", "'").replace("‘", "'")
                  .replace("“", '"').replace("”", '"').replace("‑", "-")
                  .replace(chr(0x2013), "-").replace(chr(0x2014), "-"))

def on_page(doc, page, text):
    return norm(text) in norm(doc.get(page, ""))

QUOTES = [
    # K5 and K6 locators (scorer's own)
    (P13, "11", "our approximate average holding period for a rental car was eighteen months in the United States and thirteen months in our international operations"),
    (P12, "11", "our approximate average holding period for a rental car was eighteen months in the United States and fourteen months in our international operations"),
    (P12, "81", "Useful lives are as follows:"), (P12, "81", "Cars | 4 to 28 months"),
    (P13, "91", "Holding periods are as follows:"), (P13, "91", "Cars | 4 to 36 months"),
    (P13, "53", "vehicle maintenance costs of $9.6 million due to the expansion of our off-airport and leisure businesses, longer holding periods and the impact of recalls"),
    (P13, "126", "optimization of fleet holding periods related to the integration of Dollar Thrifty"),
    (P13, "42", "optimization of fleet holding periods related to the integration of Dollar Thrifty"),
    (P13, "126", "was also indicative of the residual values experienced in the U.S."),
    (P12, "113", "indicative of the strong residual values experienced in the U.S."),
    (P13, "126", "net losses of $48.2 million, and net gains of $100.6 million"),
    (P13, "126", "declining residual values from falling demand for used vehicles"),
    (P13, "126", "net decreases of $44.2 million, $139.4 million and $26.7 million"),
    (P13, "126", "$2,407.8 | $2,145.9 | $1,912.3"),
    (P12, "113", "$2,165.2 | $1,921.8"),
    (P12, "113", "net decreases of $130.6 million"),
    (P13, "54", "a deterioration in the used vehicle residual values"),
    (P12, "42", "due to improved residual values in the U.S."),
    (P13, "42", "due to mix optimization, improved procurement and remarketing efforts"),
    (P13, "43", "slight strengthening of used vehicle residual values"),
    (P13, "45", "industry residual value guidebooks"), (P13, "45", "Depreciation rates are adjusted prospectively"),
    (P13, "45", "relatively minor adjustments"),
    # other quotations the replies rest on
    (P13, "93", "range from two to fifteen years"), (P13, "91", "Other intangible assets | 3 to 10 years"),
    (P12, "81", "Other intangible assets | 3 to 20 years"), (P13, "99", "Customer relationships | 16 | $65.0"),
    (P13, "99", "Trademark | 20 | 7.0"), (P12, "89", "Customer relationships | 16 | $65.0"),
    (P13, "50", "certain one-time charges and nonoperational items"), (P13, "140", "$645.4 million"),
    (P13, "140", "Additional efficiency and cost saving initiatives are being developed"),
    (P13, "50", "Other(11) | (50.7) | (44.0)"), (P13, "50", "Debt-related charges"), (P12, "50", "non-cash debt charges"),
    (R12, "14", "Other (g) | 5.0 | 15.8 | 23.2 | 44.0"), (R13U, None, "Other unusual/non-recurring (c)"),
    (P13, "83", "out of period errors totaling $46.3 million"), (P13, "83", "management has revised in this filing"),
    (P13, "84", "Income before income taxes | 450.6 | (9.2) | 441.4"),
    (P13, "158", "Year ended December 31, 2013 | $29.3 | $41.9 | $(0.1) | $(40.3) | (a) | $30.8"),
    (P13, "158", "Charged to Expense"),
    (P12, "140", "Year ended December 31, 2012 | $20,282 | $34,144 | $28 | $(29,341) | (a) | $25,113"),
    (P13, "80", "Provision for losses on doubtful accounts | 45.9 | 38.3 | 28.2"),
    (P13, "137", "could potentially exceed $40.0 million"), (P13, "137", "We continue to believe the outcome of this case will not be material"),
    (P13, "137", "in March 2013, the court granted, in part, the plaintiffs' motion for partial summary judgment with respect to restitution"),
    (P12, "32", "The court has since entered a stay order and the parties will again be engaging in mediation"),
    (P13, "138", "material misrepresentations and/or omissions"), (P13, "138", "none of those reserves are material"),
    (P13, "159", "the Company implemented Oracle general ledger"), (P13, "94", "no impairment existed"),
    (P13, "24", "As of February 1, 2014, Simply Wheelz was in possession of approximately 9,000 of our vehicles"),
    (P13, "98", "Accounts payable | (43)"), (P13, "147", "$3.1 million for the executive officers"),
    (P13, "70", "2,727.7"), (P13, "70", "4,757.6"), (P13, "64", "due to the timing of our payments"), (P12, "60", "due to the timing of our payments"),
    (P13, "133", "decrease in fleet receivables"), (R13, "2", "lower than expected pricing and higher expenses related to carrying extra fleet"),
    (P13, "141", "Cash payments | (22.6) | " + chr(0x2014) + " | (1.6) | (12.6) | (36.8)"), (P12, "126", "Cash payments | (22.6) | " + chr(0x2014) + " | (0.9) | (3.3) | (26.8)"),
    (P12, "126", "$10.3 million for facility closures"), (P13, "141", "$0.5 million related to a goodwill write-off"),
    (P12, "22", "We expect this percentage to continue to decrease"), (P13, "23", "Increased fleet costs, either generally or due to declines in the value of the non-program cars"),
    (P12, "22", "Declines in the value of the non-program cars"), (P12, "141", 'which was acquired on November 19, 2012'),
    (P13, "42", "Percentage of non-program cars in our U.S. car rental operations | 91% | 95% | 83%"),
    (P13, "42", "197,700, 136,400 and 121,800 non-program cars"), (P13, "42", "the average age of our fleet will increase"),
    (P13, "49", "Average number of cars (Company-operated) | 468,500 | 358,000"), (P13, "49", "Average number of cars (Leased) | 21,500 | 1,100"),
    (P13, "131", "U.S. car rental | $1,269.3 | $940.6"), (P13, "131", "International car rental | 532.0 | 528.2"),
    (P13, "15", "43 months"), (P13, "98", "nine years"),
    (R12, "4", "$1.82 - $1.92"), (R13, "4", "$1.63"), (R12, "4", "Adjusted pre-tax income | 901.5 | 901.5 | 680.5 | 680.5"),
    (R12, "10", "Income before income taxes | 450.6 | 450.9 | 901.5 | 324.3 | 356.2 | 680.5"),
    (R13, "5", "Adjusted pre-tax income | 1,153.2 | 1,153.2 | 892.3 | 892.3"), (R13, "5", "Restructuring and related charges | 98.8 | 49.1"),
    (R13U, None, "Depreciation of revenue earning equipment | 333.4 | 102.0 | 75.7 | 109.1 | 620.2 | 245.9 | 111.9 | 73.9 | 102.1 | 533.8"),
    (R13U, None, "Depreciation of revenue earning equipment and lease charges | 643.0 | 550.0"),
    (R12, "17", "Other (c) | 5.0 | 15.8 | 5.7 | 26.5 | " + chr(0x2014) + " | " + chr(0x2014) + " | 62.4 | 62.4"),
]
bad = []
for doc, page, q in QUOTES:
    ok = (norm(q) in norm(doc)) if page is None else on_page(doc, page, q)
    if not ok:
        bad.append((page, q))
print(f"Quotations asserted on their pages: {len(QUOTES) - len(bad)} of {len(QUOTES)}")
for page, q in bad:
    print("  NOT FOUND:", page, q)

# ------------------------------------------------------------- workbook
wb = openpyxl.load_workbook(os.path.join(WS, "statements.xlsx"), data_only=True)
def sheet(name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    out = {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        out[r[0]] = dict(zip(hdr[1:], r[1:]))
    return out, hdr
IS, ish = sheet("IS annual"); BS, bsh = sheet("BS annual"); CF, cfh = sheet("CF annual")
BSQ, bsqh = sheet("BS quarterly"); ISQ, isqh = sheet("IS quarterly")
FY = {y: f"FY (ended {y}-12-31)" for y in (2009, 2010, 2011, 2012, 2013)}
def isv(line, y): return IS[line][FY[y]]
def bsv(line, y): return BS[line][FY[y]]
def cfv(line, y): return CF[line][FY[y]]
PTI = isv("Income before income taxes", 2013)   # 663.1
assert abs(PTI - 663.1) < 0.05
results = []
def rec(measure, label, reported, computed, tol):
    match = abs(computed - reported) <= tol
    results.append((measure, label, reported, computed, match))
    print(f"  [{'ok ' if match else 'NO '}] {label}: reported {reported}, computed {computed:.4g}")

# ---- M1 rental depreciation rate (Note 8 charge before disposal adjustment over average gross fleet)
print("\nM1 depreciation rate on rental equipment")
gross = {y: bsv("Cars", y) + bsv("Other equipment", y) for y in (2010, 2011, 2012, 2013)}
gross12_rev = 12548.8 + bsv("Other equipment", 2012)  # Revisions sheet: cars 12,548.8 as shown later
dep = {2011: 1921.8, 2012: 2165.2, 2013: 2407.8}       # FY2012 p.113; FY2013 p.126
dep12_rev = 2145.9                                     # FY2013 p.126
rate11 = dep[2011] / ((gross[2010] + gross[2011]) / 2)
rate12 = dep[2012] / ((gross[2011] + gross[2012]) / 2)
rate13 = dep[2013] / ((gross[2012] + gross[2013]) / 2)
rate12r = dep12_rev / ((gross[2011] + gross12_rev) / 2)
rate13r = dep[2013] / ((gross12_rev + gross[2013]) / 2)
rec("M1", "FY2011 rate % (numbers-1 16.2)", 16.2, rate11 * 100, 0.05)
rec("M1", "FY2012 rate % first reported (numbers-1 15.3)", 15.3, rate12 * 100, 0.05)
rec("M1", "FY2013 rate % (numbers-1 14.3)", 14.3, rate13r * 100, 0.06)
rec("M1", "FY2012 rate % revised (numbers-2 15.2)", 15.2, rate12r * 100, 0.05)
short_fr = rate12 * (gross[2012] + gross[2013]) / 2 - dep[2013]
short_rev = rate12r * (gross12_rev + gross[2013]) / 2 - dep[2013]
rec("M1", "charge not taken at FY2012 rate, first reported $m (174.5)", 174.5, short_fr, 0.3)
rec("M1", "same, revised $m (152.1)", 152.1, short_rev, 0.3)
rec("M1", "share of pre-tax % (26.3)", 26.3, short_fr / PTI * 100, 0.1)
rec("M1", "share of pre-tax % revised, report ranks on (22.9)", 22.9, short_rev / PTI * 100, 0.1)
rec("M1", "$39.6m net rate changes as % of revised shortfall (26)", 26, 39.6 / short_rev * 100, 0.6)
flat13 = (2407.8 + 37.2) / ((gross12_rev + gross[2013]) / 2)
flat12 = (2145.9 - 96.8) / ((gross[2011] + gross12_rev) / 2)
rec("M1", "rate with disposal adjustment added back FY2013 % (14.5)", 14.5, flat13 * 100, 0.06)
rec("M1", "same FY2012 % (14.5)", 14.5, flat12 * 100, 0.06)
rec("M1", "implied life to zero FY2012 months (text-4 79)", 79, 12 / rate12r, 0.6)
rec("M1", "implied life to zero FY2013 months (text-4 84)", 84, 12 / rate13r, 0.6)
# Q4 same-quarter rate, which the run did not compute (AK12): release line over quarter-end fleet
def q(col):
    return BSQ["Cars"][col] + BSQ["Other equipment"][col]
q4_13 = 620.2 * 365 / 92 / ((q("Q3 (ended 2013-09-30)") + q("Q4 (ended 2013-12-31)")) / 2)
q4_12 = 533.8 * 366 / 92 / ((q("Q3 (ended 2012-09-30)") + q("Q4 (ended 2012-12-31)")) / 2)
print(f"  scorer's Q4 same-quarter rate, annualised: Q4 2012 {q4_12*100:.2f}%  Q4 2013 {q4_13*100:.2f}%  (not computed by the run)")

# ---- M2 stated rate changes and disposal results as shares
print("\nM2 stated rate changes, disposal results, range")
rec("M2", "$44.2m as % of pre-tax (6.7)", 6.7, 44.2 / PTI * 100, 0.05)
rec("M2", "$48.2m as % (7.3)", 7.3, 48.2 / PTI * 100, 0.05)
rec("M2", "net cuts 44.2 - 5.0 + 0.4 (39.6)", 39.6, 44.2 - 5.0 + 0.4, 0.01)
rec("M2", "$39.6m as % (6.0)", 6.0, 39.6 / PTI * 100, 0.05)
rec("M2", "2012 net cuts 139.4 - 8.8 - 0.5 (130.1)", 130.1, 139.4 - 8.8 - 0.5, 0.01)
rec("M2", "U.S. disposal swing 100.6 + 48.2 (148.8)", 148.8, 100.6 + 48.2, 0.01)
rec("M2", "swing as % (22.4)", 22.4, 148.8 / PTI * 100, 0.05)
rec("M2", "range top 28 to 36 months, rise % (29)", 29, (36 / 28 - 1) * 100, 0.6)
rec("M2", "range midpoint FY2012 (16)", 16, (4 + 28) / 2, 0.01)
rec("M2", "range midpoint FY2013 (20)", 20, (4 + 36) / 2, 0.01)

# ---- M3 per-vehicle and fleet-age tests in text-4 and numbers-2
print("\nM3 per-vehicle charges and fleet age (text-4 table, numbers-2)")
us13, us12 = 1269.3, 940.6; intl13, intl12 = 532.0, 528.2       # p.131 segment depreciation and lease charges
co13, co12, le13, le12 = 468500, 358000, 21500, 1100           # p.49
in_co13, in_co12, in_le13, in_le12 = 159700, 153700, 1600, 1400  # p.49 international
rec("M3", "U.S. charge per car-month incl. leased FY2012 $ (218)", 218, us12 * 1e6 / ((co12 + le12) * 12), 0.6)
rec("M3", "same FY2013 $ (216)", 216, us13 * 1e6 / ((co13 + le13) * 12), 0.6)
pc13 = us13 * 1e6 / (co13 * 12); pc12 = us12 * 1e6 / (co12 * 12)
rec("M3", "company-operated only, change % (numbers-2 +2.9)", 2.9, (pc13 / pc12 - 1) * 100, 0.1)
rec("M3", "international per car-month FY2012 $ (284)", 284, intl12 * 1e6 / ((in_co12 + in_le12) * 12), 0.6)
rec("M3", "same FY2013 $ (275)", 275, intl13 * 1e6 / ((in_co13 + in_le13) * 12), 0.6)
rec("M3", "car accumulated depreciation / gross FY2012 revised % (14.7)", 14.7, 1850.4 / 12548.8 * 100, 0.06)
rows = list(wb["BS annual"].iter_rows(values_only=True))
acc_rows = [r for r in rows if r[0] == "Less accumulated depreciation"]   # cars, other equipment, property
rec("M3", "same FY2013 % (18.5)", 18.5, -acc_rows[0][4] / bsv("Cars", 2013) * 100, 0.06)
oe_acc = {y: BS["Other equipment"][FY[y]] for y in (2012, 2013)}
oe_accdep = {2012: -acc_rows[1][3], 2013: -acc_rows[1][4]}
rec("M3", "equipment accumulated depreciation / gross FY2012 % (32.2)", 32.2, oe_accdep[2012] / oe_acc[2012] * 100, 0.06)
rec("M3", "same FY2013 % (31.2)", 31.2, oe_accdep[2013] / oe_acc[2013] * 100, 0.06)
proxy12 = co12 * (0.83 + 0.95) / 2 / 136400 * 12; proxy13 = co13 * (0.95 + 0.91) / 2 / 197700 * 12
rec("M3", "holding proxy FY2012 months (text-4 about 28)", 28, proxy12, 0.6)
rec("M3", "holding proxy FY2013 months (text-4 about 26)", 26, proxy13, 0.6)

# ---- M4 allowance
print("\nM4 allowance for doubtful accounts")
net = {y: bsv("Receivables, less allowance for doubtful accounts of $30.8 and $29.3", y) for y in (2010, 2011, 2012, 2013)}
allow = {2010: 19.7, 2011: 20.3, 2012: 25.1, 2013: 30.8}   # Schedule II first reported (FY2012 p.140 in thousands; FY2013 p.158)
for y, rep in ((2010, 1.43), (2011, 1.24), (2012, 1.31), (2013, 2.00)):
    rec("M4", f"allowance % of gross receivables {y} ({rep})", rep, allow[y] / (net[y] + allow[y]) * 100, 0.006)
rec("M4", "allowance / write-offs (0.76)", 0.76, 30.8 / 40.3, 0.006)
rec("M4", "charge / write-offs (1.04)", 1.04, 41.9 / 40.3, 0.006)
rec("M4", "cash-flow provision less Schedule II charge $m (4.0)", 4.0, cfv("Provision for losses on doubtful accounts", 2013) - 41.9, 0.01)
rec("M4", "that gap as % of pre-tax (0.6)", 0.6, 4.0 / PTI * 100, 0.05)
rec("M4", "2012 allowance revised 25.1 to 29.3: rise $m (4.2)", 4.2, 29.3 - 25.113, 0.05)

# ---- M5 adjusted measures
print("\nM5 adjusted pre-tax income and add-backs")
adj = {2011: 680.5, 2012: 892.3, 2013: 1153.2}; pti = {2011: 324.3, 2011.5: 305.6, 2012: 441.4, 2013: PTI}
rec("M5", "gap 2013 $m (490.1)", 490.1, adj[2013] - PTI, 0.01)
rec("M5", "gap 2013 % of pre-tax (73.9)", 73.9, (adj[2013] - PTI) / PTI * 100, 0.05)
rec("M5", "gap 2012 revised $m (450.9)", 450.9, adj[2012] - 441.4, 0.01)
rec("M5", "gap 2012 % (102)", 102, (adj[2012] - 441.4) / 441.4 * 100, 0.6)
rec("M5", "gap 2011 $m (356.2)", 356.2, adj[2011] - 324.3, 0.01)
rec("M5", "gap 2011 % of reported pre-tax (numbers-2 105)", 105, (adj[2011] - 324.3) / 324.3 * 100, 0.6)
print(f"       (on revised 2011 pre-tax 305.6, with the same adjusted figure: {(adj[2011]-324.3)/305.6*100:.1f}%; neither gives 105)")
rec("M5", "gap 2011 % of adjusted (52)", 52, (adj[2011] - 324.3) / adj[2011] * 100, 0.6)
rec("M5", "gap 2012 % of adjusted (51)", 51, (adj[2012] - 441.4) / adj[2012] * 100, 0.6)
rec("M5", "gap 2013 % of adjusted (42)", 42, (adj[2013] - PTI) / adj[2013] * 100, 0.6)
rec("M5", "net-income gap 2013 749.6 - 346.2 (403.4)", 403.4, 749.6 - 346.2, 0.01)
rec("M5", "net-income gap 2012 588.9 - 238.6 (350.3)", 350.3, 588.9 - 238.6, 0.01)
rec("M5", "restructuring plus related $98.8m as % (14.9)", 14.9, 98.8 / PTI * 100, 0.05)
rec("M5", "restructuring since 2007, average per year (92.2)", 92.2, 645.4 / 7, 0.05)
new = 40.0 + 7.8 + 44.0 + 28.7 + 50.7
rec("M5", "exclusions new in 2013 incl. whole Other line $m (171.2)", 171.2, new, 0.01)
rec("M5", "that as % (25.8)", 25.8, new / PTI * 100, 0.05)
rec("M5", "new categories ex premiums and Other $m (91.8)", 91.8, 40.0 + 7.8 + 44.0, 0.01)
rec("M5", "that as % (13.8)", 13.8, 91.8 / PTI * 100, 0.05)
print(f"       four line items absent in 2012 (ex the Other line present both years): {40.0+7.8+44.0+28.7:.1f} = {(40.0+7.8+44.0+28.7)/PTI*100:.1f}%")
rec("M5", "2012 adjusted pre-tax restated 901.5 to 892.3, change (9.2)", 9.2, 901.5 - 892.3, 0.01)
rec("M5", "Corporate EBITDA 1,635.6 to 1,626.4, change (9.2)", 9.2, 1635.6 - 1626.4, 0.01)

# ---- M6 error corrections
print("\nM6 out-of-period errors and revisions")
rec("M6", "46.3 as % of pre-tax (7.0)", 7.0, 46.3 / PTI * 100, 0.05)
rec("M6", "2013 part 46.3 - 34.7 (11.6)", 11.6, 46.3 - 34.7, 0.01)
rec("M6", "11.6 as % (1.7)", 1.7, 11.6 / PTI * 100, 0.05)
rec("M6", "2012 pre-tax 450.6 to 441.4, change (9.2)", 9.2, 450.6 - 441.4, 0.01)
rec("M6", "workbook 2012 pre-tax (450.5; PK10)", 450.5, isv("Income before income taxes", 2012), 0.01)
rec("M6", "2012 Schedule II charge 34.1 to 38.3, change (4.2)", 4.2, 38.3 - 34.144, 0.05)

# ---- M7 payables
print("\nM7 payables")
cost = {y: isv("Direct operating", y) + isv("Depreciation of revenue earning equipment and lease charges", y) for y in (2012, 2013)}
cost12_rev = 4806.0 + 2128.9
ap = {2012: bsv("Accounts payable", 2012), 2013: bsv("Accounts payable", 2013)}; ap12_rev = 1003.2
dpo12 = ap[2012] / cost[2012] * 366; dpo12r = ap12_rev / cost12_rev * 366; dpo13 = ap[2013] / cost[2013] * 365
rec("M7", "days payable FY2012 first reported (52.7)", 52.7, dpo12, 0.06)
rec("M7", "days payable FY2012 revised (52.9)", 52.9, dpo12r, 0.06)
rec("M7", "days payable FY2013 (42.7)", 42.7, dpo13, 0.06)
q4cost13 = 1439.3 + 643.0; q4cost12 = 1257.7 + 550.0        # release Table 1
rec("M7", "Q4 2012 days payable (50.8)", 50.8, ap[2012] / q4cost12 * 92, 0.06)
rec("M7", "Q4 2013 days payable (42.8)", 42.8, ap[2013] / q4cost13 * 92, 0.06)
rec("M7", "payables growth less cost growth, pts (-22.3)", -22.3, ((ap[2013] / ap[2012]) - (cost[2013] / cost[2012])) * 100, 0.06)
size_fr = cost[2013] * dpo12 / 365 - ap[2013]; size_rev = cost[2013] * dpo12r / 365 - ap[2013]
rec("M7", "size at FY2012 days, first reported $m (226.3)", 226.3, size_fr, 0.3)
rec("M7", "size as % (34.1)", 34.1, size_fr / PTI * 100, 0.06)
rec("M7", "size revised $m (232.8)", 232.8, size_rev, 0.3)
rec("M7", "size revised as % (35.1)", 35.1, size_rev / PTI * 100, 0.06)
rec("M7", "Dollar Thrifty payables $43m as % of year-end 2012 (4.3)", 4.3, 43 / ap[2012] * 100, 0.06)
dpo12_exdt = (ap12_rev - 43) / cost12_rev * 366
short_exdt = cost[2013] * dpo12_exdt / 365 - ap[2013]
rec("M7", "share of shortfall the acquisition covers % (22)", 22, (1 - short_exdt / size_rev) * 100, 0.6)
rec("M7", "payables ex fleet and capex purchases, % of sales 2013 (5.8)", 5.8, (ap[2013] - 289.1 - 55.6) / isv("Total revenues", 2013) * 100, 0.06)
rec("M7", "same 2012 revised (8.0)", 8.0, (ap12_rev - 247.0 - 35.0) / 9024.9 * 100, 0.06)
rec("M7", "balance fall 1003.2 - 967.9 (35.3)", 35.3, ap12_rev - ap[2013], 0.01)
rec("M7", "cash-flow line (23.2) plus balance fall = unreconciled (58.5)", 58.5, cfv("Accounts payable", 2013) + (ap12_rev - ap[2013]), 0.01)
rec("M7", "payables % of sales 2013 (9.0)", 9.0, ap[2013] / isv("Total revenues", 2013) * 100, 0.06)
ind_gap = ap[2013] - 0.045 * isv("Total revenues", 2013)
rec("M7", "excess over industry 4.5% of sales $m (483.2)", 483.2, ind_gap, 0.3)
rec("M7", "that as % (72.9)", 72.9, ind_gap / PTI * 100, 0.06)

# ---- M8 receivables
print("\nM8 receivables")
rev = {y: isv("Total revenues", y) for y in (2012, 2013)}
rec("M8", "DSO FY2012 (76.5)", 76.5, net[2012] / rev[2012] * 366, 0.06)
rec("M8", "DSO FY2013 (51.3)", 51.3, net[2013] / rev[2013] * 365, 0.06)
rec("M8", "Q4 2012 DSO (74.8)", 74.8, net[2012] / 2319.7 * 92, 0.06)
rec("M8", "Q4 2013 DSO (54.4)", 54.4, net[2013] / 2556.3 * 92, 0.06)
rec("M8", "receivables % of sales 2013 (14.0)", 14.0, net[2013] / rev[2013] * 100, 0.06)
rgap = net[2013] - 0.115 * rev[2013]
rec("M8", "excess over industry 11.5% $m (273.8)", 273.8, rgap, 0.3)
rec("M8", "that as % (41.3)", 41.3, rgap / PTI * 100, 0.06)
rec("M8", "ex fleet-sale receivables % of sales (10.7)", 10.7, (net[2013] - 357.3) / rev[2013] * 100, 0.06)
rec("M8", "fall in receivables, revised basis 1879.7 - 1512.6 (367.1)", 367.1, 1879.7 - net[2013], 0.01)
rec("M8", "fall in fleet-sale receivables 618.6 - 357.3 (261.3)", 261.3, 618.6 - 357.3, 0.01)
rec("M8", "fleet share of the fall % (71)", 71, 261.3 / 367.1 * 100, 0.6)
rec("M8", "fleet-sale receivables % of 2013 receivables (24)", 24, 357.3 / net[2013] * 100, 0.6)
rec("M8", "receivables growth less revenue growth, pts (-39.2)", -39.2, ((net[2013] / net[2012]) - (rev[2013] / rev[2012])) * 100, 0.06)

# ---- M9 margins
print("\nM9 margins")
gm13q = (2556.3 - 1439.3 - 643.0) / 2556.3; gm12q = (2319.7 - 1257.7 - 550.0) / 2319.7
rec("M9", "Q4 2012 gross margin % (22.1)", 22.1, gm12q * 100, 0.06)
rec("M9", "Q4 2013 gross margin % (18.5)", 18.5, gm13q * 100, 0.06)
rec("M9", "fall times Q4 2013 revenue $m (90.2)", 90.2, (gm12q - gm13q) * 2556.3, 0.3)
rec("M9", "that as % (13.6)", 13.6, (gm12q - gm13q) * 2556.3 / PTI * 100, 0.06)
gm13 = (rev[2013] - cost[2013]) / rev[2013]
rec("M9", "FY2013 gross margin % (23.2)", 23.2, gm13 * 100, 0.06)
om13 = (rev[2013] - cost[2013] - isv("Selling, general and administrative", 2013)) / rev[2013]
rec("M9", "FY2013 operating margin % (13.7)", 13.7, om13 * 100, 0.06)
rec("M9", "utilization change from +22.4% fleet and +16.1% days % (-5.1)", -5.1, (1.161 / 1.224 - 1) * 100, 0.1)
rec("M9", "explained share of the Q4 fall: 20.7 + 19.4 over 90.2 % (about 45)", 45, (20.7 + 19.4) / 90.2 * 100, 1.0)

# ---- M10 balances and exposures
print("\nM10 balances, exposures, small items")
rec("M10", "Sobel $40.0m as % (6.0)", 6.0, 40.0 / PTI * 100, 0.06)
extra = 65.0 / 15 - 65.0 / 16 + 7.0 / 15 - 7.0 / 20
rec("M10", "Donlen intangibles over 15 years instead: extra $m a year (0.4)", 0.4, extra, 0.06)
rec("M10", "that as % (0.06)", 0.06, extra / PTI * 100, 0.006)
rec("M10", "public liability reserve 347.7 as % of pre-tax (52)", 52, 347.7 / PTI * 100, 0.6)
rec("M10", "goodwill 1,347.5 (workbook)", 1347.5, bsv("Goodwill", 2013), 0.01)
rec("M10", "lease and concession commitments 2,727.7 tie (p.70)", 2727.7, 601.9 + 789.1 + 455.4 + 881.3, 0.1)

n_ok = sum(1 for r in results if r[4]); print(f"\nSub-figures: {n_ok} of {len(results)} match at the reply's rounding")
for r in results:
    if not r[4]:
        print("  mismatch:", r[0], r[1], "reported", r[2], "computed", round(r[3], 3))

# ------------------------------------------------------------- reply scans
print("\nReply scans")
replies = {n: open(os.path.join(RUN, n), encoding="utf-8").read() for n in
           ("numbers-1.md", "numbers-2.md", "text-3.md", "text-4.md", "report-5.md")}
banned = r"\bfraud|manipulat|red flag|misstat|suspicious|material weakness"
traps = (r"planned holding|21 to|25 months|24 or 30|disposed cars|affirmative|attorney|rolling|spreadsheet error|"
         r"9\.3 million|tone at the top|\b235\b|June 2014|May 2014|November 2014|2015|1\.78|1\.68|1\.72|1\.66|1\.65|budget|"
         r"\b1\.82\b|extended holding|p\. ?53\b")
for n, t in replies.items():
    b = re.findall(banned, t, re.I); tr = re.findall(traps, t, re.I)
    print(f"  {n}: banned {b}; trap or key terms {tr}; 'longer holding' {len(re.findall('longer holding', t))}; "
          f"'eighteen' {len(re.findall('eighteen', t))}; '18 months' {len(re.findall('18 months', t))}; "
          f"'optimization of fleet' {len(re.findall('optimization of fleet', t))}; "
          f"'longer|lengthen|extend' {re.findall(r'longer|lengthen|extend', t)}")
t4 = replies["text-4.md"]
print("  text-4 hit list present (counts of pages per search word)?",
      bool(re.search(r'"residual"[^\n]*\d+ hits|hits? (on|at) pp?\.|\| residual \|', t4)))
print("  FY2013 'longer holding' pages in the pack:", [p for p, t in P13.items() if "longer holding" in t],
      "FY2012:", [p for p, t in P12.items() if "longer holding" in t])
