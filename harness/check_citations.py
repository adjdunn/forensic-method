"""Check one prompt reply against the company's pack.

Usage:
    python harness/check_citations.py <reply.md|.txt|.pdf> --pack test-data/<company> [--out report.md] [--json report.json]

What it checks:
- Quotations: every "quoted passage" in the reply, whether it appears word for word in a pack document, and
  whether it is on the page the reply cites (page markers [p. N] in the pack files).
- Figures: every money amount and figure with a decimal or a thousands separator, whether it is printed in a
  pack document or in the statements workbook. Percentages are usually computed by the model and are listed
  separately. A figure not found is not an error on its own: sums and shares are derived. The list is for a reader.
- Arithmetic: where a sentence gives an amount and its share of pre-tax income, the share is recomputed.
- Banned words: fraud, manipulation, red flag; em-dashes.
- Post-cutoff facts: years later than the pack's cutoff date mentioned in the reply, with context.
- Length: words and an estimate of pages.

The pack's config.json names the documents, their year ends and the cutoff. Without it, the file names are used.
Exit code is 0; the report is the product. A runner reads the JSON.
"""
import argparse, glob, json, os, re, sys
from pathlib import Path

PAGE = re.compile(r"\[p\. (\d+)\]")
BANNED = ["fraud", "manipulat", "red flag", "red-flag"]


# ---------- loading ----------

def read_reply(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        text = "\n".join(pg.extract_text() or "" for pg in PdfReader(str(path)).pages)
        text = re.sub(r"(?m)^Page \d+\s*$", "", text)  # renderer headers
        return text
    return path.read_text(encoding="utf-8", errors="replace")


def norm(s: str) -> str:
    s = (s.replace("‘", "'").replace("’", "'").replace("“", '"').replace("”", '"')
         .replace("–", "-").replace("—", "-").replace("−", "-").replace(" ", " ")
         .replace("…", "...").replace("ﬁ", "fi").replace("ﬂ", "fl"))
    return re.sub(r"\s+", " ", s)


def load_pack(pack: Path) -> dict:
    cfg_path = pack / "config.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8")) if cfg_path.exists() else {}
    docs = []
    if cfg.get("documents"):
        for d in cfg["documents"]:
            docs.append(dict(d))
    else:  # infer from file names
        for f in sorted(glob.glob(str(pack / "annual-report-FY*.txt"))):
            name = os.path.basename(f)
            if "-key-sections" in name or "-chat" in name:
                continue
            ye = re.search(r"FY(\d{4}-\d{2}-\d{2})", name).group(1)
            docs.append({"file": name, "kind": "report", "year_end": ye})
        for f in sorted(glob.glob(str(pack / "earnings-release-*.txt"))):
            name = os.path.basename(f)
            ye = re.search(r"FY(\d{4}-\d{2}-\d{2})", name).group(1)
            docs.append({"file": name, "kind": "release", "year_end": ye})
    for d in docs:
        raw = (pack / d["file"]).read_text(encoding="utf-8", errors="replace")
        d["text"] = norm(raw)
        d["year"] = int(d["year_end"][:4])
        d["fiscal_year"] = int(d.get("fiscal_label", d["year_end"][:4])[-4:]) if str(d.get("fiscal_label", "")).strip() else d["year"]
    reports = [d for d in docs if d["kind"] == "report"]
    for d in reports:
        d["role"] = "latest" if d is max(reports, key=lambda x: x["year_end"]) else "prior"
    return {"cfg": cfg, "docs": docs, "cutoff": cfg.get("cutoff")}


def page_of(doc_text: str, pos: int) -> int | None:
    ms = PAGE.findall(doc_text[:pos])
    return int(ms[-1]) if ms else None


def _pages_of_needle(text: str, needle: str) -> list[int]:
    pages, start = [], 0
    while True:
        i = text.find(needle, start)
        if i < 0:
            break
        p = page_of(text, i)
        if p is not None and p not in pages:
            pages.append(p)
        start = i + 1
        if len(pages) > 40:
            break
    return pages


def find_pages(doc: dict, needle: str) -> list[int]:
    """Pages where the quotation appears. An elided quotation ("A... B") counts on a page that holds every part."""
    parts = [p.strip(" ,;") for p in needle.split("...") if p.strip(" ,;")]
    if not parts:
        return []
    sets = [set(_pages_of_needle(doc["text"], p)) for p in parts]
    common = set.intersection(*sets) if sets else set()
    return sorted(common)


def present(doc: dict, needle: str) -> bool:
    parts = [p.strip(" ,;") for p in needle.split("...") if p.strip(" ,;")]
    return bool(parts) and all(p in doc["text"] for p in parts)


# ---------- quotations ----------

def cited_pages(ctx: str) -> list[int]:
    out = []
    for m in re.finditer(r"\bpp?\.\s*(\d{1,3})(?:\s*(?:-|to|and)\s*(\d{1,3}))?((?:,\s*\d{1,3})*)", ctx):
        a = int(m.group(1)); b = int(m.group(2)) if m.group(2) else a
        if b < a or b - a > 6:
            b = a
        out.extend(range(a, b + 1))
        out.extend(int(x) for x in re.findall(r"\d{1,3}", m.group(3) or ""))
    return sorted(set(out))


def doc_hint(ctx: str, pack: dict) -> list[dict]:
    """Which pack documents a citation context points at. Returns [] when the context says nothing."""
    docs = pack["docs"]
    years = {int(y) for y in re.findall(r"\b(?:FY\s?|fiscal )?((?:19|20)\d{2})\b", ctx)}
    wants_release = bool(re.search(r"\breleases?\b", ctx, re.I))
    wants_report = bool(re.search(r"\b(?:report|10-K|40-F|annual|Note \d|Item \d|MD&A|management discussion)\b", ctx, re.I))
    if re.search(r"last year'?s? report|prior year'?s? report", ctx, re.I):
        return [d for d in docs if d["kind"] == "report" and d["role"] == "prior"]
    if re.search(r"this year'?s? report|latest report", ctx, re.I):
        return [d for d in docs if d["kind"] == "report" and d["role"] == "latest"]
    if not years and not wants_release and not wants_report:
        return []
    cands = []
    for d in docs:
        if wants_release and not wants_report and d["kind"] != "release":
            continue
        if wants_report and not wants_release and d["kind"] == "release":
            continue
        if years and not ({d["year"], d["fiscal_year"]} & years):
            continue
        cands.append(d)
    return cands


SENTENCE_END = re.compile(r"(?<!\bp)(?<!\bpp)(?<!\bNo)(?<!\bvs)[.;]\s")


def check_quotes(reply: str, pack: dict) -> list[dict]:
    text = norm(reply)
    out = []
    for m in re.finditer(r'"([^"]{2,400})"', text):
        q = m.group(1).strip()
        if q.startswith("[") or q.lower() in {"continue", "most likely", "possibly"} or re.fullmatch(r"(?:FY)?\d{4} (?:report|10-K|release)", q):
            continue
        after = text[m.end(): m.end() + 320]
        if re.match(r"\s*(?:is|means|refers to)\s", after):  # the reply defining its own label
            continue
        before = text[max(0, m.start() - 120): m.start()]
        # the citation sits in a parenthesis after the quote, usually right after it; otherwise read to the sentence end
        par = re.match(r"\s*\(([^)]{0,160})\)", after)
        ctx = par.group(1) if par else SENTENCE_END.split(after, maxsplit=1)[0]
        pages = cited_pages(ctx)
        if not pages:  # the citation may follow a later quote in the same sentence: '"A" and "B" (report, p. 34)'
            later = re.search(r"\(([^)]*\bpp?\.\s*\d+[^)]*)\)", after)
            if later and not SENTENCE_END.search(after[:later.start()]):
                ctx = later.group(1); pages = cited_pages(ctx)
        if not pages:  # a page given just before the quote: 'p. 42: "..."'
            pages = cited_pages(before[-40:])
        cands = doc_hint(ctx, pack) or doc_hint(before[-80:], pack) or pack["docs"]
        found = []
        for d in cands:
            ps = find_pages(d, q)
            if ps or present(d, q):
                found.append((d["file"], ps))
        elsewhere = []
        if not found:
            for d in pack["docs"]:
                if d in cands:
                    continue
                ps = find_pages(d, q)
                if ps or present(d, q):
                    elsewhere.append((d["file"], ps))
        if found:
            on_page = any(set(pages) & set(ps) for _, ps in found) if pages else None
            status = "exact, on cited page" if on_page else ("exact, page differs" if pages else "exact, no page cited")
        elif elsewhere:
            status = "found in another document"
        else:
            status = "not found"
        out.append({"quote": q, "cited_pages": pages, "status": status,
                    "found": found or elsewhere, "context": ctx[:120]})
    return out


# ---------- figures ----------

NUM = re.compile(r"(?<![\w.])(\$?)(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d{4,})\s*(%|bn|billion|m\b|M\b|million|k\b|K\b)?")


def load_workbook_values(pack_dir: Path) -> set[str]:
    vals = set()
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(pack_dir / "statements.xlsx"), data_only=True, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if isinstance(c, (int, float)):
                        for d in (0, 1, 2):
                            vals.add(f"{round(c, d):,.{d}f}")
                            vals.add(f"{round(c, d):.{d}f}")
    except Exception:
        pass
    return vals


def check_figures(reply: str, pack: dict, pack_dir: Path) -> dict:
    text = norm(reply)
    # strip page citations, fiscal labels and list numbering before scanning
    scan = re.sub(r"\bpp?\.\s*\d{1,3}(?:\s*(?:-|to|and)\s*\d{1,3})?", " ", text)
    scan = re.sub(r"\b(?:FY|Q[1-4]|fiscal|Note|Item|Schedule|Form)\s?\d+\S*", " ", scan)
    scan = re.sub(r"\b(?:19|20)\d{2}\b", " ", scan)
    wb = load_workbook_values(pack_dir)
    all_text = " ".join(d["text"] for d in pack["docs"])
    seen, results = set(), {"found": [], "computed_percent": [], "not_found": []}
    for m in NUM.finditer(scan):
        raw, unit = m.group(2), (m.group(3) or "")
        key = raw + unit
        if key in seen:
            continue
        seen.add(key)
        if unit == "%":
            if raw in all_text:
                results["found"].append(raw + "%")
            else:
                results["computed_percent"].append(raw + "%")
            continue
        plain = raw.replace(",", "")
        variants = {raw, plain}
        if "." not in raw:
            variants |= {raw + ".0", f"{int(plain):,}"}
        hit = any(v in all_text for v in variants) or any(v in wb for v in variants)
        (results["found"] if hit else results["not_found"]).append(m.group(1) + raw + (" " + unit if unit else ""))
    return results


# ---------- arithmetic ----------

AMOUNT = re.compile(r"\$\s?(\d{1,3}(?:,\d{3})*(?:\.\d+)?|\d+(?:\.\d+)?)\s*(bn|billion|m|M|million)?")


def to_millions(val: str, unit: str) -> float:
    x = float(val.replace(",", ""))
    return x * 1000 if unit and unit.lower().startswith("b") else x


def workbook_pre_tax_income(pack_dir: Path) -> float | None:
    """The latest fiscal year's income before taxes from the statements workbook, in millions."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(pack_dir / "statements.xlsx"), data_only=True, read_only=True)
        for name in wb.sheetnames:
            if "IS" not in name or "annual" not in name.lower():
                continue
            for row in wb[name].iter_rows(values_only=True):
                if row and isinstance(row[0], str) and re.search(r"(?i)(income|earnings).*before.*tax", row[0]) and not re.search(r"(?i)per share", row[0]):
                    nums = [c for c in row[1:] if isinstance(c, (int, float))]
                    if nums:
                        return float(nums[-1])
    except Exception:
        return None
    return None


PTI_STATED = re.compile(r"pre-tax income(?: of| is| was| =|:)?\s*(?:FY\s?\S+\s*(?:=|of)?\s*)?\$?\s?(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d{3,})\s*(bn|billion|m|M|million)?", re.I)
NUMBER = r"(?<![\d.,])\$\s?(\d{1,3}(?:,\d{3})*(?:\.\d+)?|\d+(?:\.\d+)?)\s*(bn|billion|m|M|million)?(?![\d.,%])"
# "$746m, 113% of pre-tax income" / "$607m, 3.3% of pre-tax" / "$1,334m (7.2% of pre-tax income)"
PAIR_AMT_PCT = re.compile(NUMBER + r"(?:[,:]|\s+or|\s+which is|\s+=)?\s*\(?\s*(?<![\d.])(\d{1,3}(?:\.\d+)?)% of (?:(?:FY\s?\S+|reported|latest year's|that) )?pre-tax", re.I)
# "113% of pre-tax income ($746m: ...)" / "8.1% of pre-tax income (net income less free cash flow, 1,510)"
PAIR_PCT_AMT = re.compile(r"(?<![\d.])(\d{1,3}(?:\.\d+)?)% of (?:(?:FY\s?\S+|reported|latest year's|that) )?pre-tax(?: income)?\s*\(([^)]{0,80})\)", re.I)


def check_arithmetic(reply: str, pack_dir: Path) -> dict:
    text = norm(reply)
    stated = PTI_STATED.search(text)
    pti_stated = to_millions(stated.group(1), stated.group(2) or "") if stated else None
    pti_wb = workbook_pre_tax_income(pack_dir)
    pti = pti_stated or pti_wb
    note = ""
    if pti_stated and pti_wb and abs(pti_stated - pti_wb) / pti_wb > 0.02:
        note = f"reply states pre-tax income {pti_stated}; workbook latest year {pti_wb} (a revised figure, or a different year)"
    if not pti:
        return {"pre_tax_income": None, "checks": [], "note": "no pre-tax income found"}
    checks, seen = [], set()

    def add(amt, pct, sent):
        key = (amt, pct)
        if key in seen:
            return
        seen.add(key)
        calc = amt / pti * 100
        ok = abs(calc - pct) <= max(0.25, 0.015 * pct)
        checks.append({"amount": amt, "stated": pct, "recomputed": round(calc, 1), "ok": ok, "sentence": sent[:160]})

    for sent in re.split(r"(?<=[.;])\s+(?=[A-Z•\-•])", text):
        for m in PAIR_AMT_PCT.finditer(sent):
            add(to_millions(m.group(1), m.group(2) or ""), float(m.group(3)), sent)
        for m in PAIR_PCT_AMT.finditer(sent):
            inner = AMOUNT.search(m.group(2)) or re.search(r"(?<![\d.,])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+)\s*(bn|m|M|million)?(?![\d.,%])", m.group(2))
            if inner:
                add(to_millions(inner.group(1), inner.group(2) or ""), float(m.group(1)), sent)
    return {"pre_tax_income": pti, "source": "reply" if pti_stated else "workbook", "note": note, "checks": checks}


# ---------- words, cutoff, length ----------

def check_banned(reply: str) -> dict:
    low = reply.lower()
    hits = {w: len(re.findall(re.escape(w), low)) for w in BANNED}
    return {"words": {k: v for k, v in hits.items() if v}, "em_dashes": reply.count("—")}


def check_cutoff(reply: str, cutoff: str | None) -> list[dict]:
    if not cutoff:
        return []
    cy = int(cutoff[:4])
    out = []
    for m in re.finditer(r"\b(20\d{2}|19\d{2})\b", reply):
        y = int(m.group(1))
        if y > cy:
            out.append({"year": y, "context": norm(reply[max(0, m.start() - 70): m.end() + 50])})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("reply"); ap.add_argument("--pack", required=True)
    ap.add_argument("--out"); ap.add_argument("--json")
    a = ap.parse_args()
    reply_path, pack_dir = Path(a.reply), Path(a.pack)
    reply = read_reply(reply_path)
    pack = load_pack(pack_dir)
    quotes = check_quotes(reply, pack)
    figures = check_figures(reply, pack, pack_dir)
    arith = check_arithmetic(reply, pack_dir)
    banned = check_banned(reply)
    post = check_cutoff(reply, pack["cutoff"])
    words = len(reply.split())
    summary = {
        "reply": str(reply_path), "pack": pack_dir.name, "words": words, "pages_est": round(words / 450, 1),
        "quotes": {"total": len(quotes), "exact_on_page": sum(q["status"] == "exact, on cited page" for q in quotes),
                   "exact_page_differs": sum(q["status"] == "exact, page differs" for q in quotes),
                   "exact_no_page": sum(q["status"] == "exact, no page cited" for q in quotes),
                   "other_document": sum(q["status"] == "found in another document" for q in quotes),
                   "not_found": sum(q["status"] == "not found" for q in quotes)},
        "figures": {k: len(v) for k, v in figures.items()},
        "arithmetic": {"pre_tax_income": arith["pre_tax_income"], "checked": len(arith["checks"]),
                       "mismatches": sum(not c["ok"] for c in arith["checks"])},
        "banned": banned, "post_cutoff": len(post),
    }
    lines = [f"# Citation check: {reply_path.name} against {pack_dir.name}", "",
             f"{words} words, about {summary['pages_est']} pages.", "",
             "| Check | Result |", "|---|---|",
             f"| Quotations | {summary['quotes']['exact_on_page']} of {len(quotes)} exact and on the cited page; "
             f"{summary['quotes']['exact_page_differs']} exact with a different page; {summary['quotes']['exact_no_page']} exact, no page cited; "
             f"{summary['quotes']['other_document']} found in another document; {summary['quotes']['not_found']} not found |",
             f"| Figures | {summary['figures']['found']} printed in the pack or workbook; {summary['figures']['not_found']} not printed (derived or to check); "
             f"{summary['figures']['computed_percent']} percentages not printed (computed) |",
             f"| Arithmetic | pre-tax income {arith['pre_tax_income']} ({arith.get('source', 'none')}); {summary['arithmetic']['checked']} shares recomputed, {summary['arithmetic']['mismatches']} mismatches{'; ' + arith['note'] if arith.get('note') else ''} |",
             f"| Banned words | {banned['words'] or 'none'}; em-dashes {banned['em_dashes']} |",
             f"| Post-cutoff years | {len(post)} mentions after {pack['cutoff'] or 'no cutoff in config'} |", ""]
    bad = [q for q in quotes if q["status"] != "exact, on cited page"]
    if bad:
        lines += ["## Quotations to look at", ""]
        for q in bad:
            where = "; ".join(f"{f} {ps}" for f, ps in q["found"]) if q["found"] else "nowhere"
            lines.append(f"- **{q['status']}**: \"{q['quote'][:90]}\" cited {q['cited_pages'] or 'no page'}, found {where}")
        lines.append("")
    if figures["not_found"]:
        lines += ["## Figures not printed in the pack (derived, rounded differently, or wrong)", "",
                  ", ".join(figures["not_found"][:60]), ""]
    mism = [c for c in arith["checks"] if not c["ok"]]
    if mism:
        lines += ["## Arithmetic mismatches", ""]
        for c in mism:
            lines.append(f"- stated {c['stated']}%, recomputed {c['recomputed']}% from {c['amount']}: {c['sentence']}")
        lines.append("")
    if post:
        lines += ["## Mentions after the cutoff (a recommendation to read a later filing is fine; a fact from one is not)", ""]
        for x in post[:20]:
            lines.append(f"- {x['year']}: ...{x['context']}...")
        lines.append("")
    report = "\n".join(lines)
    if a.out:
        Path(a.out).write_text(report, encoding="utf-8")
    if a.json:
        Path(a.json).write_text(json.dumps({"summary": summary, "quotes": quotes, "figures": figures, "arithmetic": arith,
                                            "post_cutoff": post}, indent=1), encoding="utf-8")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(report)


if __name__ == "__main__":
    main()
