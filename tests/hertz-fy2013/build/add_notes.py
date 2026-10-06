"""Add note data from the pack's text files to statements.xlsx NOTES (run after build_workbook.py).

- Allowance rollforward (Schedule II), each year as originally reported in its own 10-K;
  FY2009 and FY2010 from the FY2011 10-K (the earliest 10-K in the text pack).
- Employees at year end, from Item 1 of each 10-K.
Values are USD millions; the FY2011 and FY2012 schedules are printed in thousands.
"""
import re
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).parent
PACK = HERE.parent / "pack"
man = (PACK / "manifest.md").read_text(encoding="utf-8")


def pid(name):
    m = re.search(r"\|\s*(P\d+)\s*\|\s*`?" + re.escape(name), man)
    return m.group(1) if m else "?"


def num(s):
    s = s.replace("$", "").replace(",", "").strip()
    if s in ("—", "-", ""):
        return 0.0
    neg = s.startswith("(")
    v = float(s.strip("()"))
    return -v if neg else v


def schedule(year_file, fy):
    txt = (PACK / f"AR-FY{year_file}_schedule-ii.txt").read_text(encoding="utf-8")
    scale = 0.001 if "Thousands" in txt else 1.0
    page = re.search(r"\[p\. (\d+)\]", txt).group(1)
    sec = txt.split("Allowance for doubtful accounts:")[1].split("Tax valuation")[0]
    for line in sec.splitlines():
        if f"December 31, {fy}" in line:
            cells = [c.strip() for c in line.split("|")[1:] if c.strip() and c.strip() != "(a)"]
            opening, charged, fx, deductions, closing = (num(c) * scale for c in cells[:5])
            return {"opening": opening, "charged": charged, "fx": fx, "deductions": deductions, "closing": closing, "page": page}
    return None


def main():
    wb = load_workbook(PACK / "statements.xlsx")
    ns, lc = wb["NOTES"], wb["LOCATORS"]
    hdr = [c.value for c in ns[1]]
    rows = {ns.cell(r, 1).value: r for r in range(2, ns.max_row + 1)}
    sources = {2009: "2011", 2010: "2011", 2011: "2011", 2012: "2012", 2013: "2013"}
    for fy, yf in sources.items():
        s = schedule(yf, fy)
        col = hdr.index(f"FY{fy}") + 1
        f = f"AR-FY{yf}_schedule-ii.txt"
        loc = f"{pid(f)} ({f}), Schedule II, p. {s['page']}"
        for line, key in (("Allowance: opening", "opening"), ("Allowance: provision", "charged"), ("Allowance: write-offs", "deductions"), ("Allowance: other", "fx"), ("Allowance: closing", "closing")):
            ns.cell(rows[line], col).value = round(s[key], 1)
            lc.append(["NOTES", line, f"FY{fy}", loc, "Schedule II, allowance for doubtful accounts", "N",
                       "Write-offs are 'amounts written off, net of recoveries' as the schedule labels them; recoveries are not shown separately"
                       if key == "deductions" else ("Schedule II 'charged to expense'; the cash flow statement's provision line can differ" if key == "charged" else "")])
    for r in range(2, len(hdr) + 1):
        if hdr[r - 1].startswith("Q") and ns.cell(rows["Allowance: provision"], r).value is not None:
            lc.append(["NOTES", "Allowance: provision", hdr[r - 1], "P1 cash flow statement, provision for losses on doubtful accounts", "quarterly values from the cash flow statement", "Y", "Quarterly provision is the cash-flow line, discrete quarters derived as for CF_Q; annual columns are Schedule II"])
    emp = {}
    for yf in ("2011", "2012", "2013"):
        f = f"AR-FY{yf}_business.txt"
        t = (PACK / f).read_text(encoding="utf-8")
        m = re.search(r"As of December 31, " + yf + r", we employed approximately ([\d,]+) persons", t)
        pages = [(mm.start(), mm.group(1)) for mm in re.finditer(r"\[p\. (\d+)\]", t)]
        pg = [p for pos, p in pages if pos < m.start()][-1] if pages else "?"
        emp[yf] = (float(m.group(1).replace(",", "")), f"{pid(f)} ({f}), Item 1, p. {pg}")
    for yf, (v, loc) in emp.items():
        ns.cell(rows["Employees at year end"], hdr.index(f"FY{yf}") + 1).value = v
        lc.append(["NOTES", "Employees at year end", f"FY{yf}", loc, "Item 1 Employees", "N", "'approximately', as stated"])
    wb.save(PACK / "statements.xlsx")
    print("notes added; schedules:", {fy: schedule(sources[fy], fy) for fy in sources})
    print("employees:", emp)


if __name__ == "__main__":
    main()
