"""Working-capital measures from statements.xlsx (Under Armour, as originally reported)."""
import re, sys
from datetime import date
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
wb = openpyxl.load_workbook("statements.xlsx", data_only=True)

def load(sheet):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    ends = {}
    for j, h in enumerate(hdr):
        m = re.search(r"ended (\d{4})-(\d{2})-(\d{2})", str(h or ""))
        if m:
            ends[j] = date(*map(int, m.groups()))
    out = {}
    for r in rows[1:]:
        if r[0] and r[0] not in out:          # first occurrence of a label
            out[r[0]] = {ends[j]: r[j] for j in ends}
    return out

def days_in(end, months):
    y, m = end.year, end.month - months + 1
    while m <= 0:
        m += 12; y -= 1
    return (end - date(y, m, 1)).days + 1

def measures(IS, BS, CF, months):
    rev, cogs = IS["Net revenues"], IS["Cost of goods sold"]
    ar, inv, ap = BS["Accounts receivable, net"], BS["Inventories"], BS["Accounts payable"]
    res = {}
    for end in sorted(ar):
        d = days_in(end, months)
        prior = date(end.year - 1, end.month, end.day)
        g = lambda s: (s[end] / s[prior] - 1) * 100 if prior in s and s.get(prior) else None
        gap = lambda a, b: (g(a) - g(b)) if g(a) is not None and g(b) is not None else None
        res[end] = dict(
            days=d, rev=rev[end], cogs=cogs[end], ar=ar[end], inv=inv[end], ap=ap[end],
            dso=ar[end] / rev[end] * d, dio=inv[end] / cogs[end] * d, dpo=ap[end] / cogs[end] * d,
            g_rev=g(rev), g_cogs=g(cogs), g_ar=g(ar), g_inv=g(inv), g_ap=g(ap),
            gap_ar=gap(ar, rev), gap_inv=gap(inv, cogs), gap_ap=gap(ap, cogs),
            cfo=CF["Net cash provided by (used in) operating activities"].get(end),
            cf_ap=CF["Accounts payable"].get(end), cf_ar=CF["Accounts receivable"].get(end),
            cf_res=CF["Changes in reserves and allowances"].get(end),
            acq=CF["Purchase of businesses, net of cash acquired"].get(end),
        )
    return res

A = measures(load("IS annual"), load("BS annual"), load("CF annual"), 12)
Q = measures(load("IS quarterly"), load("BS quarterly"), load("CF quarterly"), 3)

f = lambda v, p=1: "n/a" if v is None else f"{v:,.{p}f}"
for name, R in (("ANNUAL", A), ("QUARTERLY", Q)):
    print(name)
    for end, m in R.items():
        print(end, " ".join(f"{k}={f(v)}" for k, v in m.items()))

# Extra checks used in the flags
q = lambda y, mth, dd: Q[date(y, mth, dd)]
print("Q4 AP: 2014", q(2014,12,31)["ap"], "2015", q(2015,12,31)["ap"], "2016", q(2016,12,31)["ap"])
print("AP change Q3->Q4:", {y: round(q(y,12,31)["ap"] - q(y,9,30)["ap"], 1) for y in (2014, 2015, 2016)})
print("Q4 2016 CFO ex AP inflow:", round(q(2016,12,31)["cfo"] - q(2016,12,31)["cf_ap"], 1))
print("FY2016 CFO ex AP inflow:", round(A[date(2016,12,31)]["cfo"] - A[date(2016,12,31)]["cf_ap"], 1))
print("FY2016 DPO if AP had grown with COGS:",
      round(A[date(2015,12,31)]["ap"] * (1 + A[date(2016,12,31)]["g_cogs"]/100) / A[date(2016,12,31)]["cogs"] * 366, 1))
gw = load("BS quarterly")["Goodwill"]
print("Goodwill:", {str(k): v for k, v in gw.items()})
gm = lambda IS, e: (1 - IS["Cost of goods sold"][e] / IS["Net revenues"][e]) * 100
ISq, ISa = load("IS quarterly"), load("IS annual")
print("Gross margin Q:", {str(e): round(gm(ISq, e), 1) for e in ISq["Net revenues"]})
print("Gross margin A:", {str(e): round(gm(ISa, e), 1) for e in ISa["Net revenues"]})

# Revisions touching the lines used
ws = wb["Revisions"]
keys = ("receivable", "Inventor", "Accounts payable", "Net revenues", "Cost of goods")
hits = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(k in str(r[1]) for k in keys)]
print("Revisions rows touching AR / inventory / AP / revenue / COGS:", len(hits))
