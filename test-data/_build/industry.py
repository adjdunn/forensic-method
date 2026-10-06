"""Add an "Industry" sheet to a company workbook: the matching industry row from Aswath Damodaran's public
industry tables (NYU Stern), from the January update nearest the company's last fiscal year end.

Source: https://pages.stern.nyu.edu/~adamodar/New_Home_Page/data.html (current year) and the archive page for earlier
years. Files are cached in industry/damodaran/. Archive suffix NN is the January (NN+1) update, which covers fiscal
year NN: wcdata13.xls is dated January 2014 and is the table for a December 2013 year end.

Peer sets built from FMP were tried and dropped (2 October 2026): they had to be hand-picked, which is a judgment the
tool should not make for the user.

Usage: imported by build.py; or  python industry.py <slug>  to refresh one workbook's sheet.
"""
import re, sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent
ROOT = HERE.parent
DAMO = HERE / "industry" / "damodaran"

INDUSTRY = {
    "hertz-fy2013": ("Transportation", "13", "Rental and leasing has no industry row of its own; Transportation is the nearest bucket."),
    "under-armour-fy2016": ("Apparel", "16", ""),
    "kraft-heinz-fy2017": ("Food Processing", "17", ""),
    "home-depot-fy2025": ("Retail (Building Supply)", "", ""),
    "cae-fy2026": ("Aerospace/Defense", "", "CAE reports under IFRS in Canadian dollars; the industry row is US companies in US dollars. Ratios are currency-free."),
}


def damodaran_row(industry, suffix):
    rows = {}
    for f, cols in (("wcdata", ["Acc Rec/ Sales", "Inventory/Sales", "Acc Pay/ Sales"]),
                    ("margin", ["Gross Margin", "Pre-tax Unadjusted Operating Margin"]),
                    ("capex", ["Cap Ex/Deprecn"])):
        p = DAMO / f"{f}{suffix}.xls"
        if not p.exists():
            continue
        x = pd.ExcelFile(p)
        sheet = next((n for n in x.sheet_names if "Industry" in n), x.sheet_names[-1])
        df = x.parse(sheet, header=None)
        hdr = next(i for i in range(25) if any(str(v).strip().startswith("Industry") for v in df.iloc[i].tolist()))
        head = [str(c).strip() for c in df.iloc[hdr].tolist()]
        data = df.iloc[hdr + 1:]
        data.columns = head
        updated = next((str(v) for r in range(hdr) for v in df.iloc[r].tolist()[1:2] if "20" in str(v)), "")
        hit = data[data[head[0]].astype(str).str.strip() == industry]
        if hit.empty:
            rows[f] = {"missing": f"no row named {industry!r} in {p.name}"}
            continue
        r = hit.iloc[0]
        rows[f] = {"file": p.name, "updated": updated[:10], "firms": int(r.get("Number of firms", r.get("Number of Firms", 0)) or 0)}
        for c in cols:
            key = next((h for h in head if h.replace(" ", "") == c.replace(" ", "")), None)
            if key:
                v = r[key]
                rows[f][c] = round(float(v) * (100 if f != "capex" else 1), 1) if pd.notna(v) else None
    return rows


def add_sheet(slug, workbook_path):
    cfg = INDUSTRY.get(slug)
    if not cfg:
        return None
    industry, suffix, note = cfg
    wb = load_workbook(workbook_path)
    if "Industry" in wb.sheetnames:
        del wb["Industry"]
    ws = wb.create_sheet("Industry")
    ws.column_dimensions["A"].width = 60
    ws.append([f"Industry averages: {industry}"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Aswath Damodaran, NYU Stern. US industry averages computed across all firms in the group, so they are aggregates, not medians, and carry no range. Use them to judge whether a level is ordinary for the industry; the statement sheets show the company's own movements."])
    if note:
        ws.append([note])
    ws.append([])
    rows = damodaran_row(industry, suffix)
    for f, r in rows.items():
        if "missing" in r:
            ws.append([r["missing"]]); continue
        ws.append([f"{r['file']}: data updated {r['updated']}, {r['firms']} firms in the group"])
        ws[f"A{ws.max_row}"].font = Font(bold=True)
        for k, v in r.items():
            if k not in ("file", "updated", "firms"):
                ws.append([k, v])
        ws.append([])
    ws.append(["Shares of sales are in percent. Cap Ex/Deprecn is a ratio. To turn receivables over sales into days: share divided by 100, times 365. The industry inventory and payables shares are over sales, while the company's days inventory and days payable use cost of sales, so compare those through the company's own inventory-to-sales and payables-to-sales shares."])
    ws.append(["Source: https://pages.stern.nyu.edu/~adamodar/New_Home_Page/data.html (current year); https://pages.stern.nyu.edu/~adamodar/New_Home_Page/dataarchived.html (earlier years)."])
    wb.save(workbook_path)
    return {f: r.get("updated", r.get("missing")) for f, r in rows.items()}


if __name__ == "__main__":
    for slug in sys.argv[1:] or INDUSTRY:
        print(slug, add_sheet(slug, ROOT / slug / "statements.xlsx"))
