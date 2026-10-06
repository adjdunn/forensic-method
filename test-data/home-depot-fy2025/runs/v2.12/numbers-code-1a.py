# Scan the numbers: Home Depot. All figures USD millions.
# Run from the folder holding statements.xlsx and the four text files.
import openpyxl, re, datetime as dt
wb = openpyxl.load_workbook('statements.xlsx')

def sheet(name):
    ws = wb[name]
    heads = [c.value for c in ws[1]]
    data, fills = {}, {}
    for row in ws.iter_rows(min_row=2):
        label = row[0].value
        if label is None: continue
        for h, c in zip(heads[1:], row[1:]):
            if h == 'XBRL element': continue
            end = re.search(r'ended (\d{4}-\d{2}-\d{2})', h).group(1)
            data.setdefault(label, {})[end] = c.value
            fills.setdefault(label, {})[end] = c.fill.fgColor.rgb if c.fill.fill_type else None
    return data, fills, heads

ISa, _, _ = sheet('IS annual'); BSa, _, _ = sheet('BS annual'); CFa, _, _ = sheet('CF annual')
ISq, ISqf, ISqh = sheet('IS quarterly'); BSq, _, _ = sheet('BS quarterly'); CFq, CFqf, CFqh = sheet('CF quarterly')
print('IS quarterly headings marked mixed:', [h for h in ISqh if 'mixed' in str(h)])
print('CF quarterly headings marked mixed:', [h for h in CFqh if 'mixed' in str(h)])
ORANGE = '00F8D7B0'
for lab in CFq:
    for e in ('2025-02-02', '2026-02-01'):
        if CFqf[lab].get(e) == ORANGE: print('orange (left out of every measure):', lab, e)

# The company names its fiscal years by the calendar year they mostly cover.
FY = {'FY2021': '2022-01-30', 'FY2022': '2023-01-29', 'FY2023': '2024-01-28', 'FY2024': '2025-02-02', 'FY2025': '2026-02-01'}
names = list(FY)
def D(s): return dt.date.fromisoformat(s)
days = {}
for i, n in enumerate(names):
    if i: days[n] = (D(FY[n]) - D(FY[names[i-1]])).days
print('days in year', days)
REV, COGS, GP, SGA, OI, PTI, NI = ('Net sales', 'Cost of sales', 'Gross profit', 'Selling, general and administrative',
    'Operating income', 'Earnings before provision for income taxes', 'Net earnings')
AR, INV, AP, TA, CA, PPE, CL, LTD = ('Receivables, net', 'Merchandise inventories', 'Accounts payable', 'Total assets',
    'Total current assets', 'Net property and equipment', 'Total current liabilities', 'Long-term debt, excluding current installments')
OCF, CAPEX, ACQ = 'Net cash provided by operating activities', 'Capital expenditures', 'Payments for businesses acquired, net'

# Property note (annual reports, p. 58): cost less land and construction in progress;
# depreciation and finance lease amortization expense (the one depreciation figure used everywhere).
def txt(f): return open(f, encoding='utf-8', errors='replace').read()
AR25, AR26 = txt('annual-report-FY2025-02-02-key-sections.txt'), txt('annual-report-FY2026-02-01-key-sections.txt')
def page(t, p):
    a = t.index('[p. %d]' % p); b = t.find('[p. %d]' % (p+1), a)
    return t[a:b if b > 0 else None]
def row(t, p, label):
    m = re.search(r'^' + re.escape(label) + r' \| (.+)$', page(t, p), re.M)
    return [float(x.replace('$', '').replace(',', '')) for x in m.group(1).split(' | ')]
gross = {}
for t, cols in ((AR25, ('FY2024', 'FY2023')), (AR26, ('FY2025', 'FY2024'))):
    cost, land, cip = row(t, 58, 'Property and equipment, at cost'), row(t, 58, 'Land'), row(t, 58, 'Construction in progress')
    for k, n in enumerate(cols):
        g = cost[k] - land[k] - cip[k]
        assert gross.get(n, g) == g; gross[n] = g
d25 = row(AR25, 58, 'Depreciation and finance lease amortization expense')
d26 = row(AR26, 58, 'Depreciation and finance lease amortization expense')
DEP = {'FY2024': d25[0], 'FY2023': d25[1], 'FY2022': d25[2], 'FY2025': d26[0]}
assert d26[1] == DEP['FY2024'] and d26[2] == DEP['FY2023']
print('gross depreciable', gross, 'depreciation', DEP)
for t, p, q in ((AR25, 48, 'The valuation allowance related to these receivables was not material to our consolidated financial statements at the end of fiscal 2024 or fiscal 2023.'),
                (AR26, 48, 'The valuation allowance related to these receivables was not material to our consolidated financial statements at the end of fiscal 2025 or fiscal 2024.')):
    print('allowance sentence found on p.%d:' % p, q in page(t, p))

M = {}
def put(r, c, v): M.setdefault(r, {})[c] = v
for i, n in enumerate(names):
    e = FY[n]; rev, cogs = ISa[REV][e], ISa[COGS][e]
    put('Gross margin %', n, 100*ISa[GP][e]/rev); put('Operating margin %', n, 100*ISa[OI][e]/rev)
    fcf = CFa[OCF][e] + CFa[CAPEX][e]
    put('FCF', n, fcf); put('FCF / net income', n, fcf/ISa[NI][e]); put('NI - FCF', n, ISa[NI][e]-fcf)
    if e in BSa[AR]:
        put('DSO', n, BSa[AR][e]/rev*days[n]); put('DIO', n, BSa[INV][e]/cogs*days[n]); put('DPO', n, BSa[AP][e]/cogs*days[n])
        put('AR/sales %', n, 100*BSa[AR][e]/rev); put('Inv/sales %', n, 100*BSa[INV][e]/rev); put('AP/sales %', n, 100*BSa[AP][e]/rev)
        p = FY[names[i-1]]
        if p in BSa[AR]:
            rg, cg = rev/ISa[REV][p]-1, cogs/ISa[COGS][p]-1
            put('AR gap pts', n, 100*(BSa[AR][e]/BSa[AR][p]-1-rg)); put('Inv gap pts', n, 100*(BSa[INV][e]/BSa[INV][p]-1-cg))
            put('AP gap pts', n, 100*(BSa[AP][e]/BSa[AP][p]-1-cg))
            put('Accruals %', n, 100*(ISa[NI][e]-CFa[OCF][e])/((BSa[TA][e]+BSa[TA][p])/2))
    if n in DEP: put('Capex / depreciation', n, -CFa[CAPEX][e]/DEP[n])
    if n in gross and names[i-1] in gross:
        put('Avg gross depreciable', n, (gross[n]+gross[names[i-1]])/2); put('Depreciation rate %', n, 100*DEP[n]/M['Avg gross depreciable'][n])

# Beneish M-score, one per pair of years
def beneish(n, pn):
    e, p = FY[n], FY[pn]
    s, sp = ISa[REV][e], ISa[REV][p]
    DSRI = (BSa[AR][e]/s)/(BSa[AR][p]/sp)
    GMI = (ISa[GP][p]/sp)/(ISa[GP][e]/s)
    aq = lambda x: 1-(BSa[CA][x]+BSa[PPE][x])/BSa[TA][x]
    AQI = aq(e)/aq(p); SGI = s/sp
    dr = lambda x, y: DEP[y]/(DEP[y]+BSa[PPE][x])
    DEPI = dr(p, pn)/dr(e, n)
    SGAI = (ISa[SGA][e]/s)/(ISa[SGA][p]/sp)
    lv = lambda x: (BSa[CL][x]+BSa[LTD][x])/BSa[TA][x]
    LVGI = lv(e)/lv(p)
    TATA = (ISa[NI][e]-CFa[OCF][e])/BSa[TA][e]
    m = -4.84+0.920*DSRI+0.528*GMI+0.404*AQI+0.892*SGI+0.115*DEPI-0.172*SGAI+4.679*TATA-0.327*LVGI
    print(n, dict(DSRI=round(DSRI,3), GMI=round(GMI,3), AQI=round(AQI,3), SGI=round(SGI,3), DEPI=round(DEPI,3), SGAI=round(SGAI,3), TATA=round(TATA,4), LVGI=round(LVGI,3)), round(m,2))
    return m
for n, pn in (('FY2023','FY2022'), ('FY2024','FY2023'), ('FY2025','FY2024')): put('M-score', n, beneish(n, pn))

# Quarter: Q4 FY2025 (ended 2026-02-01) beside Q4 FY2024 (ended 2025-02-02); Q4 FY2023 only as the base for growth.
# Later quarters in the workbook (to 2026-08-02) are not covered by an attached report or release.
Q = {'Q4 FY2023': ('2024-01-28', '2023-10-29'), 'Q4 FY2024': ('2025-02-02', '2024-10-27'), 'Q4 FY2025': ('2026-02-01', '2025-11-02')}
qn = list(Q)
for i, n in enumerate(qn):
    e, s0 = Q[n]; qd = (D(e)-D(s0)).days; rev, cogs = ISq[REV][e], ISq[COGS][e]
    put('qdays', n, qd)
    put('Gross margin %', n, 100*ISq[GP][e]/rev); put('Operating margin %', n, 100*ISq[OI][e]/rev)
    put('DSO', n, BSq[AR][e]/rev*qd); put('DIO', n, BSq[INV][e]/cogs*qd); put('DPO', n, BSq[AP][e]/cogs*qd)
    fcf = CFq[OCF][e]+CFq[CAPEX][e]
    put('FCF', n, fcf); put('FCF / net income', n, fcf/ISq[NI][e]); put('NI - FCF', n, ISq[NI][e]-fcf)
    if s0 in BSq[TA]: put('Accruals %', n, 100*(ISq[NI][e]-CFq[OCF][e])/((BSq[TA][e]+BSq[TA][s0])/2))
    if i:
        p = Q[qn[i-1]][0]; rg, cg = rev/ISq[REV][p]-1, cogs/ISq[COGS][p]-1
        put('AR gap pts', n, 100*(BSq[AR][e]/BSq[AR][p]-1-rg)); put('Inv gap pts', n, 100*(BSq[INV][e]/BSq[INV][p]-1-cg))
        put('AP gap pts', n, 100*(BSq[AP][e]/BSq[AP][p]-1-cg))

# check the quarter against the releases
for f, e in (('earnings-release-Q4-FY2025-02-02.txt', '2025-02-02'), ('earnings-release-Q4-FY2026-02-01.txt', '2026-02-01')):
    t = txt(f)
    for lab in (REV, COGS, GP, OI, NI):
        m = re.search(r'^' + re.escape(lab) + r' \| \$? ?([\d,]+)', t, re.M)
        print(f[:28], lab, 'release', m.group(1) if m else None, 'sheet', ISq[lab][e])

for r in M: print(r, {k: round(v, 2) for k, v in M[r].items()})

# Industry
ind = {r[0].value: r[1].value for r in wb['Industry'].iter_rows() if r[1].value is not None}
print(ind)

# Sizes, as shares of the latest year's pre-tax income
L = 'FY2025'; e = FY[L]; PT = ISa[PTI][e]
def at_prior_days(bal, drv, key):
    would = M[key]['FY2024']*ISa[drv][e]/days[L]; amt = BSa[bal][e]-would
    print(key, 'balance', BSa[bal][e], 'at prior days', round(would), 'amount', round(amt), 'share %', round(100*amt/PT, 1)); return amt
at_prior_days(AR, REV, 'DSO'); at_prior_days(INV, COGS, 'DIO'); at_prior_days(AP, COGS, 'DPO')
for k in ('Gross margin %', 'Operating margin %'):
    fall = M[k]['FY2024']-M[k]['FY2025']; amt = fall/100*ISa[REV][e]
    print(k, 'fall pts', round(fall, 3), 'amount', round(amt), 'share %', round(100*amt/PT, 1))
amt = (M['Depreciation rate %']['FY2024']-M['Depreciation rate %']['FY2025'])/100*M['Avg gross depreciable']['FY2025']
print('depreciation rate: charge not taken at prior rate', round(amt), 'share %', round(100*amt/PT, 1))
amt = M['NI - FCF'][L]
print('NI - FCF', ISa[NI][e], M['FCF'][L], amt, 'share %', round(100*amt/PT, 1))
print('pre-tax income', PT)
print('AR/sales', M['AR/sales %'][L], 'vs', ind['Acc Rec/ Sales'], '| Inv/sales', M['Inv/sales %'][L], 'vs', ind['Inventory/Sales'],
      '| AP/sales', M['AP/sales %'][L], 'vs', ind['Acc Pay/ Sales'])
for lab, bal, key, sign in (('AR', AR, 'Acc Rec/ Sales', 1), ('Inventory', INV, 'Inventory/Sales', 1), ('AP', AP, 'Acc Pay/ Sales', -1)):
    lvl = ind[key]/100*ISa[REV][e]; gap = sign*(BSa[bal][e]-lvl)
    print(lab, 'at industry share of sales', round(lvl), 'gap', round(gap), 'share %', round(100*gap/PT, 1))
# acquisitions and what they brought
print('acquisitions', {n: CFa[ACQ][FY[n]] for n in names})
print('goodwill', BSa['Goodwill'], 'intangibles', BSa['Intangible assets, net'], 'total assets', BSa[TA])
print('customer receivables (p.48)', row(AR25, 48, 'Customer receivables'), row(AR26, 48, 'Customer receivables'))
print('SG&A % sales', {n: round(100*ISa[SGA][FY[n]]/ISa[REV][FY[n]], 2) for n in names},
      'D&A % sales', {n: round(100*ISa['Depreciation and amortization'][FY[n]]/ISa[REV][FY[n]], 2) for n in names})
# revisions
for r in wb['Revisions'].iter_rows(min_row=2, values_only=True): print(r[0], '|', r[1], '|', r[3], '|', r[4], '->', r[6], '|', r[8])
print('intangible amortization', CFa['Intangible asset amortization'], 'Q1 FY2024', CFq['Intangible asset amortization']['2024-04-28'],
      'intangibles at 2024-01-28', BSa['Intangible assets, net']['2024-01-28'])

# ---------- the reply ----------
def persist(key, cols):
    v = [M[key][c] for c in cols if c in M[key]]
    d = [b-a for a, b in zip(v, v[1:])]
    k = 0
    for x in reversed(d):
        if x*d[-1] > 0: k += 1
        else: break
    return k
YRS = names; QS = ['Q4 FY2023', 'Q4 FY2024', 'Q4 FY2025']
P = {k: (persist(k, YRS), persist(k, QS)) for k in ('DSO', 'DIO', 'DPO', 'Operating margin %', 'FCF / net income', 'Accruals %', 'Gross margin %')}
print('persistence (years, same-quarter comparisons):', P)

cols = ['FY2021', 'FY2022', 'FY2023', 'FY2024', 'FY2025', 'Q4 FY2024', 'Q4 FY2025']
heads = ['FY2021 (Jan 2022)', 'FY2022 (Jan 2023)', 'FY2023 (Jan 2024)', 'FY2024 (Feb 2025, 53 wk)', 'FY2025 (Feb 2026)', 'Q4 FY2024 (14 wk)', 'Q4 FY2025 (13 wk)', 'Industry']
rows = [('Days sales outstanding', 'DSO', '{:.1f}', None), ('Days inventory', 'DIO', '{:.1f}', None), ('Days payable', 'DPO', '{:.1f}', None),
        ('Receivables growth less sales growth, pts', 'AR gap pts', '{:+.1f}', None), ('Inventory growth less cost-of-sales growth, pts', 'Inv gap pts', '{:+.1f}', None),
        ('Payables growth less cost-of-sales growth, pts', 'AP gap pts', '{:+.1f}', None),
        ('Receivables, % of sales', 'AR/sales %', '{:.1f}', 'Acc Rec/ Sales'), ('Inventory, % of sales', 'Inv/sales %', '{:.1f}', 'Inventory/Sales'),
        ('Payables, % of sales', 'AP/sales %', '{:.1f}', 'Acc Pay/ Sales'),
        ('Gross margin, %', 'Gross margin %', '{:.1f}', 'Gross Margin'), ('Operating margin, %', 'Operating margin %', '{:.1f}', 'Pre-tax Unadjusted Operating Margin'),
        ('Accruals, % of average assets', 'Accruals %', '{:.1f}', None), ('Free cash flow ÷ net income', 'FCF / net income', '{:.2f}', None),
        ('Depreciation rate, %', 'Depreciation rate %', '{:.2f}', None), ('Capital spending ÷ depreciation', 'Capex / depreciation', '{:.2f}', 'Cap Ex/Deprecn'),
        ('Beneish M-score', 'M-score', '{:.2f}', None)]
tb = ['| Measure | ' + ' | '.join(heads) + ' |', '|---|' + '---:|'*len(heads)]
for lab, key, fmt, ik in rows:
    cells = [fmt.format(M[key][c]) if c in M[key] else '' for c in cols] + [str(ind[ik]) if ik else '']
    tb.append('| ' + lab + ' | ' + ' | '.join(cells) + ' |')
table = '\n'.join(tb)

m = lambda k, c, f='{:.1f}': f.format(M[k][c])
n0 = lambda x: '{:,.0f}'.format(x)
pc = lambda x: '{:.1f}%'.format(100*x/PT)
ar_sz = BSa[AR][e]-M['DSO']['FY2024']*ISa[REV][e]/days[L]
inv_sz = BSa[INV][e]-M['DIO']['FY2024']*ISa[COGS][e]/days[L]
ap_sz = M['DPO']['FY2024']*ISa[COGS][e]/days[L]-BSa[AP][e]
om_fall = M['Operating margin %']['FY2024']-M['Operating margin %']['FY2025']; om_sz = om_fall/100*ISa[REV][e]
fcf_sz = M['NI - FCF'][L]
pe = FY['FY2024']
apg = 100*(BSa[AP][e]/BSa[AP][pe]-1); cg = 100*(ISa[COGS][e]/ISa[COGS][pe]-1)
IA = 'Intangible asset amortization'
revs = list(wb['Revisions'].iter_rows(min_row=2, values_only=True))
assert -revs[0][8] == BSa['Intangible assets, net']['2024-01-28']
assert [-r[8] for r in revs[1:]] == [CFa[IA]['2023-01-29'], CFa[IA]['2024-01-28'], CFq[IA]['2024-04-28']]

reply = f"""{table}

Cost line: Cost of sales, as reported. Depreciation throughout is the property note's depreciation and finance lease amortization expense; the M-score's depreciation index sets it against itself plus net property. Q4 is from the spreadsheet (one cash flow line is mixed basis and is left out); quarterly accruals are not annualised.

**Flags**

1. **Receivables.** Days sales outstanding rose from {m('DSO','FY2023')} (FY2023) to {m('DSO','FY2024')} to {m('DSO','FY2025')}; receivables are {m('AR/sales %','FY2025')}% of sales against {ind['Acc Rec/ Sales']}% for the industry. Acquisitions took {n0(-CFa[ACQ][pe])} of cash in FY2024 and {n0(-CFa[ACQ][e])} in FY2025.
   Size: {pc(ar_sz)} of pre-tax income ({n0(ar_sz)}: receivables of {n0(BSa[AR][e])} less {n0(BSa[AR][e]-ar_sz)} at FY2024 days). Up {P['DSO'][0]} years running; {P['DSO'][1]} same-quarter comparisons.
2. **Operating margin.** Fell from {m('Operating margin %','FY2022')}% (FY2022) to {m('Operating margin %','FY2023')}%, {m('Operating margin %','FY2024')}% and {m('Operating margin %','FY2025')}%, while gross margin moved from {m('Gross margin %','FY2022')}% to {m('Gross margin %','FY2025')}%; Q4 fell from {m('Operating margin %','Q4 FY2024')}% to {m('Operating margin %','Q4 FY2025')}%. Intangible amortization rose from {CFa[IA][FY['FY2023']]} to {CFa[IA][e]} over the two acquisition years.
   Size: {pc(om_sz)} of pre-tax income ({n0(om_sz)}: the {om_fall:.2f}-point fall in FY2025 times revenue of {n0(ISa[REV][e])}). Down {P['Operating margin %'][0]} years running; {P['Operating margin %'][1]} same-quarter comparisons.
3. **Cash flow against earnings.** Free cash flow fell from {m('FCF / net income','FY2023','{:.2f}')} times net income (FY2023) to {m('FCF / net income','FY2024','{:.2f}')} to {m('FCF / net income','FY2025','{:.2f}')}; accruals moved from {m('Accruals %','FY2023')}% to {m('Accruals %','FY2025')}% of average assets. Operating cash flow fell {n0(CFa[OCF][pe]-CFa[OCF][e])} in FY2025.
   Size: {pc(fcf_sz)} of pre-tax income ({n0(fcf_sz)}: net income of {n0(ISa[NI][e])} less free cash flow of {n0(M['FCF'][L])}). Down {P['FCF / net income'][0]} years running; {P['FCF / net income'][1]} same-quarter comparisons.
4. **Inventory.** Days inventory rose from {m('DIO','FY2023')} (FY2023) to {m('DIO','FY2024')} to {m('DIO','FY2025')}; Q4 rose from {m('DIO','Q4 FY2024')} to {m('DIO','Q4 FY2025')}. The same acquisitions could explain it; inventory is {m('Inv/sales %','FY2025')}% of sales against {ind['Inventory/Sales']}% for the industry.
   Size: {pc(inv_sz)} of pre-tax income ({n0(inv_sz)}: inventory of {n0(BSa[INV][e])} less {n0(BSa[INV][e]-inv_sz)} at FY2024 days). Up {P['DIO'][0]} years running; {P['DIO'][1]} same-quarter comparisons.
5. **Payables.** Payables fell {abs(apg):.1f}% in FY2025 while cost of sales rose {cg:.1f}%: days payable went from {m('DPO','FY2024')} to {m('DPO','FY2025')}, and payables are {m('AP/sales %','FY2025')}% of sales against {ind['Acc Pay/ Sales']}% for the industry.
   Size: {pc(ap_sz)} of pre-tax income ({n0(ap_sz)}: {n0(BSa[AP][e]+ap_sz)} at FY2024 days less payables of {n0(BSa[AP][e])}). Down {P['DPO'][0]} year; {P['DPO'][1]} same-quarter comparison.

**Revised figures:** four later changes, none of which moved income or equity: other assets at January 2024 went from {n0(revs[0][4])} to {n0(revs[0][6])} when intangible assets of {n0(-revs[0][8])} took their own line, and cash flow depreciation for FY2022, FY2023 and Q1 FY2024 was cut by {-revs[1][8]}, {-revs[2][8]} and {-revs[3][8]}, the intangible amortization now shown separately.

**Not calculated:** the allowance as a share of gross receivables (both annual reports say only that it "was not material"; no amount is given); the depreciation rate before FY2024 and for the quarters (gross property is given at three year-ends only); days, growth gaps and accruals for FY2021 (no balance sheet before January 2023); the two quarters after February 2026 (no attached document covers them).
"""
for bad in ('fraud', 'manipulat', 'red flag', '—'):
    assert bad not in reply.lower(), bad
assert 'was not material' in page(AR25, 48) and 'was not material' in page(AR26, 48)
import os
os.makedirs('runs/v2.12', exist_ok=True)
open('runs/v2.12/numbers-1a.md', 'w', encoding='utf-8').write(reply)
wc = lambda s: len(s.split())
print('words with table', wc(reply), 'without table', wc(reply.replace(table, '')))
print(reply)
