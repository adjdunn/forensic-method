# Scan the numbers: The Home Depot, statements.xlsx + annual report property notes
# Run from the folder that holds the five attached files.
import openpyxl, re, datetime as dt, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
wb = openpyxl.load_workbook('statements.xlsx')


def sheet(name):
    ws = wb[name]
    heads = [c.value for c in ws[1]]
    cols = {}
    for j, h in enumerate(heads[1:-1], start=2):
        m = re.search(r'ended (\d{4}-\d{2}-\d{2})', h)
        cols[m.group(1)] = (j, h)
    rows = {}
    for r in range(2, ws.max_row + 1):
        rows[ws.cell(r, 1).value] = r

    def get(line, end):
        return ws.cell(rows[line], cols[end][0]).value

    def fill(line, end):
        c = ws.cell(rows[line], cols[end][0])
        return c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
    return get, fill, cols


ISa, _, ISa_c = sheet('IS annual')
BSa, _, BSa_c = sheet('BS annual')
CFa, _, CFa_c = sheet('CF annual')
ISq, ISq_fill, ISq_c = sheet('IS quarterly')
BSq, _, BSq_c = sheet('BS quarterly')
CFq, CFq_fill, CFq_c = sheet('CF quarterly')

REV, COGS, GP, SGA, OI, PTI, NI = ('Net sales', 'Cost of sales', 'Gross profit', 'Selling, general and administrative',
                                   'Operating income', 'Earnings before provision for income taxes', 'Net earnings')
AR, INV, AP, TA, CA, PPE, CL, LTD = ('Receivables, net', 'Merchandise inventories', 'Accounts payable', 'Total assets',
                                     'Total current assets', 'Net property and equipment', 'Total current liabilities',
                                     'Long-term debt, excluding current installments')
CFO, CAPEX, ACQ = 'Net cash provided by operating activities', 'Capital expenditures', 'Payments for businesses acquired, net'
DA_CF = 'Depreciation and amortization, excluding amortization of intangible assets'

FY = ['2022-01-30', '2023-01-29', '2024-01-28', '2025-02-02', '2026-02-01']
LBL = {'2022-01-30': 'FY2021', '2023-01-29': 'FY2022', '2024-01-28': 'FY2023', '2025-02-02': 'FY2024', '2026-02-01': 'FY2025'}


def d(s):
    return dt.date.fromisoformat(s)


def days(end, prev):
    return (d(end) - d(prev)).days


# property note (annual reports, p. 58): depreciation and finance lease amortization; cost, land, CIP
A25 = open('annual-report-FY2025-02-02-key-sections.txt', encoding='utf-8').read()
A26 = open('annual-report-FY2026-02-01-key-sections.txt', encoding='utf-8').read()


def page(txt, n):
    i = txt.index('[p. %d]' % n)
    j = txt.index('[p. %d]' % (n + 1))
    return txt[i:j]


def row(txt, label):
    for ln in txt.splitlines():
        if ln.startswith(label + ' |'):
            return [float(x.strip().replace('$', '').replace(',', '')) for x in ln.split('|')[1:]]
    raise KeyError(label)


p25, p26 = page(A25, 58), page(A26, 58)
dep25 = row(p25, 'Depreciation and finance lease amortization expense')  # FY24, FY23, FY22
dep26 = row(p26, 'Depreciation and finance lease amortization expense')  # FY25, FY24, FY23
assert dep25[0] == dep26[1] and dep25[1] == dep26[2]
DEP = {'2023-01-29': dep25[2], '2024-01-28': dep25[1], '2025-02-02': dep25[0], '2026-02-01': dep26[0]}


def gross(p):
    c, l, k = row(p, 'Property and equipment, at cost'), row(p, 'Land'), row(p, 'Construction in progress')
    return [c[i] - l[i] - k[i] for i in range(2)]


g25, g26 = gross(p25), gross(p26)
assert g25[0] == g26[1]
GROSS = {'2024-01-28': g25[1], '2025-02-02': g25[0], '2026-02-01': g26[0]}
# allowance: stated only as "not material" (receivables note, p. 48 of each report)
for t in (A25, A26):
    assert 'The valuation allowance related to these receivables was not material' in page(t, 48)

R = {}


def pct(x):
    return 100 * x


for i, e in enumerate(FY):
    m = {}
    rev, cogs = ISa(REV, e), ISa(COGS, e)
    m['gm'] = pct(ISa(GP, e) / rev)
    m['om'] = pct(ISa(OI, e) / rev)
    m['ni'] = ISa(NI, e)
    m['fcf'] = CFa(CFO, e) + CFa(CAPEX, e)
    m['ni_fcf'] = m['ni'] - m['fcf']
    m['fcf_ni'] = m['fcf'] / m['ni']
    if e in DEP:
        m['capex_dep'] = -CFa(CAPEX, e) / DEP[e]
    if e in BSa_c:
        n = days(e, FY[i - 1])
        m['days'] = n
        m['dso'] = BSa(AR, e) / rev * n
        m['dio'] = BSa(INV, e) / cogs * n
        m['dpo'] = BSa(AP, e) / cogs * n
        m['ar_s'] = pct(BSa(AR, e) / rev)
        m['inv_s'] = pct(BSa(INV, e) / rev)
        m['ap_s'] = pct(BSa(AP, e) / rev)
        p = FY[i - 1]
        if p in BSa_c:
            rg = rev / ISa(REV, p) - 1
            cg = cogs / ISa(COGS, p) - 1
            m['gap_ar'] = pct(BSa(AR, e) / BSa(AR, p) - 1 - rg)
            m['gap_inv'] = pct(BSa(INV, e) / BSa(INV, p) - 1 - cg)
            m['gap_ap'] = pct(BSa(AP, e) / BSa(AP, p) - 1 - cg)
            avg = (BSa(TA, e) + BSa(TA, p)) / 2
            m['accr'] = pct((m['ni'] - CFa(CFO, e)) / avg)
            if p in GROSS:
                m['deprate'] = pct(DEP[e] / ((GROSS[e] + GROSS[p]) / 2))
            # Beneish M-score
            DSRI = (BSa(AR, e) / rev) / (BSa(AR, p) / ISa(REV, p))
            GMI = (ISa(GP, p) / ISa(REV, p)) / (ISa(GP, e) / rev)

            def aq(x):
                return 1 - (BSa(CA, x) + BSa(PPE, x)) / BSa(TA, x)
            AQI = aq(e) / aq(p)
            SGI = rev / ISa(REV, p)

            def dr(x):
                return DEP[x] / (DEP[x] + BSa(PPE, x))
            DEPI = dr(p) / dr(e)
            SGAI = (ISa(SGA, e) / rev) / (ISa(SGA, p) / ISa(REV, p))
            TATA = (m['ni'] - CFa(CFO, e)) / BSa(TA, e)

            def lv(x):
                return (BSa(LTD, x) + BSa(CL, x)) / BSa(TA, x)
            LVGI = lv(e) / lv(p)
            m['M'] = (-4.84 + 0.920 * DSRI + 0.528 * GMI + 0.404 * AQI + 0.892 * SGI + 0.115 * DEPI
                      - 0.172 * SGAI + 4.679 * TATA - 0.327 * LVGI)
            m['Mparts'] = dict(DSRI=DSRI, GMI=GMI, AQI=AQI, SGI=SGI, DEPI=DEPI, SGAI=SGAI, TATA=TATA, LVGI=LVGI)
    R[LBL[e]] = m

# Quarter: the latest quarter the attached documents cover is Q4 ended 2026-02-01.
# (quarter end, previous quarter end, same quarter-end a year earlier)
Q = {'Q4 FY2024': ('2025-02-02', '2024-10-27', '2024-01-28'), 'Q4 FY2025': ('2026-02-01', '2025-11-02', '2025-02-02')}
for e in ('2026-02-01', '2025-02-02'):
    assert 'mixed basis' not in ISq_c[e][1]
    assert '[1 line mixed basis]' in CFq_c[e][1]
ORANGE = CFq_fill(DA_CF, '2026-02-01')
assert ORANGE == CFq_fill(DA_CF, '2025-02-02')
for Ln in (CFO, CAPEX, 'Net earnings'):
    for e in ('2026-02-01', '2025-02-02'):
        assert CFq_fill(Ln, e) != ORANGE  # lines used are not the mixed-basis (orange) line
for q, (e, prevq, yago) in Q.items():
    m = {}
    n = days(e, prevq)
    m['days'] = n
    rev, cogs = ISq(REV, e), ISq(COGS, e)
    m['gm'] = pct(ISq(GP, e) / rev)
    m['om'] = pct(ISq(OI, e) / rev)
    m['dso'] = BSq(AR, e) / rev * n
    m['dio'] = BSq(INV, e) / cogs * n
    m['dpo'] = BSq(AP, e) / cogs * n
    rg = rev / ISq(REV, yago) - 1
    cg = cogs / ISq(COGS, yago) - 1
    m['gap_ar'] = pct(BSq(AR, e) / BSq(AR, yago) - 1 - rg)
    m['gap_inv'] = pct(BSq(INV, e) / BSq(INV, yago) - 1 - cg)
    m['gap_ap'] = pct(BSq(AP, e) / BSq(AP, yago) - 1 - cg)
    m['ni'] = ISq(NI, e)
    m['fcf'] = CFq(CFO, e) + CFq(CAPEX, e)
    m['ni_fcf'] = m['ni'] - m['fcf']
    m['fcf_ni'] = m['fcf'] / m['ni']
    m['accr'] = pct((m['ni'] - CFq(CFO, e)) / ((BSq(TA, e) + BSq(TA, prevq)) / 2))
    R[q] = m
# Q4 FY2023 (ended 2024-01-28) margins and days, for the same-quarter persistence counts only
e0, p0 = '2024-01-28', '2023-10-29'
Q0 = {'om': pct(ISq(OI, e0) / ISq(REV, e0)), 'gm': pct(ISq(GP, e0) / ISq(REV, e0)),
      'dso': BSq(AR, e0) / ISq(REV, e0) * days(e0, p0), 'dio': BSq(INV, e0) / ISq(COGS, e0) * days(e0, p0),
      'dpo': BSq(AP, e0) / ISq(COGS, e0) * days(e0, p0), 'ni_fcf': ISq(NI, e0) - (CFq(CFO, e0) + CFq(CAPEX, e0))}

# tie the spreadsheet quarter to the releases
rel26 = open('earnings-release-Q4-FY2026-02-01.txt', encoding='utf-8').read()
rel25 = open('earnings-release-Q4-FY2025-02-02.txt', encoding='utf-8').read()
TIE = {k: ('{:,}'.format(v) in t) for k, v, t in [
    ('rev26', ISq(REV, '2026-02-01'), rel26), ('ni26', ISq(NI, '2026-02-01'), rel26), ('oi26', ISq(OI, '2026-02-01'), rel26),
    ('rev25', ISq(REV, '2025-02-02'), rel25), ('ni25', ISq(NI, '2025-02-02'), rel25), ('oi25', ISq(OI, '2025-02-02'), rel25)]}

IND = {r[0].value: r[1].value for r in wb['Industry'].iter_rows(min_row=5) if r[1].value is not None}
REVS = [(r[1].value, r[3].value, r[4].value, r[6].value, r[8].value) for r in wb['Revisions'].iter_rows(min_row=2)]

# flag sizes (latest year FY2025), each as a share of FY2025 pre-tax income
L, P = '2026-02-01', '2025-02-02'
PT = ISa(PTI, L)
F = {}
F['ar_excess'] = BSa(AR, L) - R['FY2024']['dso'] / R['FY2025']['days'] * ISa(REV, L)
F['inv_excess'] = BSa(INV, L) - R['FY2024']['dio'] / R['FY2025']['days'] * ISa(COGS, L)
F['ap_excess'] = BSa(AP, L) - R['FY2024']['dpo'] / R['FY2025']['days'] * ISa(COGS, L)
F['om_fall'] = (R['FY2024']['om'] - R['FY2025']['om']) / 100 * ISa(REV, L)
F['gm_fall'] = (R['FY2024']['gm'] - R['FY2025']['gm']) / 100 * ISa(REV, L)
F['dep_short'] = (R['FY2024']['deprate'] - R['FY2025']['deprate']) / 100 * (GROSS[L] + GROSS[P]) / 2
F['ni_fcf'] = R['FY2025']['ni_fcf']
FP = {k: pct(v / PT) for k, v in F.items()}


def run_len(vals, up=True):
    """consecutive latest periods the measure moved the same way"""
    n = 0
    for a, b in zip(vals[-2::-1], vals[::-1]):
        if (b > a) if up else (b < a):
            n += 1
        else:
            break
    return n


Y4 = ['FY2022', 'FY2023', 'FY2024', 'FY2025']
Y5 = ['FY2021'] + Y4
PERSIST = {
    'dso_up': run_len([R[y]['dso'] for y in Y4]),
    'dio_up': run_len([R[y]['dio'] for y in Y4]),
    'dpo_down': run_len([R[y]['dpo'] for y in Y4], up=False),
    'om_down': run_len([R[y]['om'] for y in Y5], up=False),
    'gm_down': run_len([R[y]['gm'] for y in Y5], up=False),
    'ni_fcf_up': run_len([R[y]['ni_fcf'] for y in Y5]),
    'accr_up': run_len([R[y]['accr'] for y in Y4[1:]]),
    'q_om_down': run_len([Q0['om'], R['Q4 FY2024']['om'], R['Q4 FY2025']['om']], up=False),
    'q_dso_up': run_len([Q0['dso'], R['Q4 FY2024']['dso'], R['Q4 FY2025']['dso']]),
    'q_dio_up': run_len([Q0['dio'], R['Q4 FY2024']['dio'], R['Q4 FY2025']['dio']]),
    'q_dpo_down': run_len([Q0['dpo'], R['Q4 FY2024']['dpo'], R['Q4 FY2025']['dpo']], up=False),
    'q_ni_fcf_up': run_len([Q0['ni_fcf'], R['Q4 FY2024']['ni_fcf'], R['Q4 FY2025']['ni_fcf']]),
}

if __name__ == '__main__':
    for k, m in R.items():
        print(k, {a: (round(b, 3) if isinstance(b, (int, float)) else b) for a, b in m.items() if a != 'Mparts'})
        if 'Mparts' in m:
            print('   ', {a: round(b, 4) for a, b in m['Mparts'].items()})
    print('Q4 FY2023', {a: round(b, 2) for a, b in Q0.items()})
    print('TIE', TIE)
    print(IND)
    print(REVS)
    print('F', {k: round(v, 1) for k, v in F.items()})
    print('F%', {k: round(v, 2) for k, v in FP.items()}, 'pretax', PT)
    print('PERSIST', PERSIST)
    print('DEP', DEP, 'GROSS', GROSS)
    print('ACQ', {LBL[e]: CFa(ACQ, e) for e in FY}, 'Q4', CFq(ACQ, '2026-02-01'), CFq(ACQ, '2025-02-02'))
    print('SGA/sales', {LBL[e]: round(100 * ISa(SGA, e) / ISa(REV, e), 2) for e in FY})
    print('AR, INV, AP', {LBL[e]: (BSa(AR, e), BSa(INV, e), BSa(AP, e)) for e in FY[1:]})

# ---------------------------------------------------------------- reply
IA = 'Intangible asset amortization'
assert R['FY2022']['fcf'] < R['FY2022']['ni'] and R['FY2023']['fcf'] > R['FY2023']['ni'] and R['FY2024']['fcf'] > R['FY2024']['ni']
assert all(R[y]['M'] < -1.78 for y in ('FY2023', 'FY2024', 'FY2025'))
assert all(v[4] is not None for v in REVS) and len(REVS) == 4
assert all(TIE.values())
COLS = ['FY2021', 'FY2022', 'FY2023', 'FY2024', 'FY2025', 'Q4 FY2024', 'Q4 FY2025']
MINUS = '−'


def f(v, nd=1, money=False):
    if v is None:
        return ''
    s = '{:,.0f}'.format(abs(v)) if money else '{:.{n}f}'.format(abs(v), n=nd)
    neg = v < 0 and float(s.replace(',', '')) != 0
    return (MINUS if neg else '') + s


ROWS = [
    ('Days sales outstanding', 'dso', 1, False, f(IND['Acc Rec/ Sales'] / 100 * 365, 1)),
    ('Days inventory', 'dio', 1, False, ''),
    ('Days payable', 'dpo', 1, False, ''),
    ('Receivables growth less revenue growth, pts', 'gap_ar', 1, False, ''),
    ('Inventory growth less cost-of-sales growth, pts', 'gap_inv', 1, False, ''),
    ('Payables growth less cost-of-sales growth, pts', 'gap_ap', 1, False, ''),
    ('Receivables, % of sales', 'ar_s', 1, False, f(IND['Acc Rec/ Sales'])),
    ('Inventory, % of sales', 'inv_s', 1, False, f(IND['Inventory/Sales'])),
    ('Payables, % of sales', 'ap_s', 1, False, f(IND['Acc Pay/ Sales'])),
    ('Gross margin, %', 'gm', 1, False, f(IND['Gross Margin'])),
    ('Operating margin, %', 'om', 1, False, f(IND['Pre-tax Unadjusted Operating Margin'])),
    ('Accruals, % of average assets', 'accr', 1, False, ''),
    ('Free cash flow, $m', 'fcf', 0, True, ''),
    ('Net income, $m', 'ni', 0, True, ''),
    ('Depreciation rate, %', 'deprate', 2, False, ''),
    ('Capital spending / depreciation', 'capex_dep', 2, False, f(IND['Cap Ex/Deprecn'], 2)),
    ('Beneish M-score', 'M', 2, False, ''),
]
HEAD = ['Measure', 'FY2021', 'FY2022', 'FY2023', 'FY2024 (53 wks)', 'FY2025', 'Q4 FY2024 (14 wks)', 'Q4 FY2025', 'Industry']
T = ['| ' + ' | '.join(HEAD) + ' |', '|' + '---|' * len(HEAD)]
for name, key, nd, money, ind in ROWS:
    T.append('| ' + ' | '.join([name] + [f(R[c].get(key), nd, money) for c in COLS] + [ind]) + ' |')
TABLE = '\n'.join(T)


def a(y, k, nd=1):
    return f(R[y][k], nd)


def mny(v):
    return f(v, 0, True)


acq24, acq25 = mny(-CFa(ACQ, P)), mny(-CFa(ACQ, L))
reply = f"""Fiscal years are the company's: FY2025 ended February 1, 2026. $ millions.

{TABLE}

Cost line: Cost of sales. Depreciation is the property note's depreciation and finance lease amortization (it includes the part inside cost of sales), used for the depreciation rate, capital spending / depreciation and the M-score; the M-score's accruals term is (net income {MINUS} operating cash flow) ÷ total assets. The quarter is from the spreadsheet: only its cash-flow depreciation line is mixed basis, and no measure uses it.

**Flags**

1. **Cash flow against earnings.** Free cash flow fell from {mny(R['FY2024']['fcf'])} to {mny(R['FY2025']['fcf'])} in FY2025, below net income ({mny(R['FY2025']['ni'])}) for the first time since FY2022; accruals went from {a('FY2024','accr')}% to {a('FY2025','accr')}%. FY2024 had a 53rd week.
   Size: {f(FP['ni_fcf'])}% of pre-tax income (net income less free cash flow, {mny(F['ni_fcf'])}). The gap has widened {PERSIST['ni_fcf_up']} years running, and in {PERSIST['q_ni_fcf_up']} fourth-quarter comparisons.
2. **Operating margin.** {a('FY2022','om')}% in FY2022, then {a('FY2023','om')}%, {a('FY2024','om')}% and {a('FY2025','om')}%; fourth quarter {a('Q4 FY2024','om')}% to {a('Q4 FY2025','om')}%. Gross margin barely moved. Acquisitions took {acq24} of cash in FY2024 and {acq25} in FY2025.
   Size: {f(FP['om_fall'])}% (the {f(R['FY2024']['om'] - R['FY2025']['om'], 2)}-point FY2025 fall × FY2025 revenue, {mny(F['om_fall'])}). Down {PERSIST['om_down']} consecutive years and {PERSIST['q_om_down']} fourth-quarter comparisons.
3. **Inventory.** Days inventory {a('FY2023','dio')}, {a('FY2024','dio')}, {a('FY2025','dio')}; fourth quarter {a('Q4 FY2024','dio')} to {a('Q4 FY2025','dio')}. Inventory grew {a('FY2025','gap_inv')} points faster than cost of sales in FY2025. The FY2025 acquisition ({acq25}) brought balances without a full year of cost of sales.
   Size: {f(FP['inv_excess'])}% (inventory less its level at FY2024 days, {mny(F['inv_excess'])}). Up {PERSIST['dio_up']} consecutive years.
4. **Receivables.** Days sales outstanding {a('FY2023','dso')}, {a('FY2024','dso')}, {a('FY2025','dso')}; receivables are {a('FY2025','ar_s')}% of sales against {f(IND['Acc Rec/ Sales'])}% for the industry. The jump came in the year of the {acq24} acquisition.
   Size: {f(FP['ar_excess'])}% (receivables less their level at FY2024 days, {mny(F['ar_excess'])}). Up {PERSIST['dso_up']} consecutive years.
5. **Payables.** Days payable fell from {a('FY2024','dpo')} to {a('FY2025','dpo')}, with payables growth {f(abs(R['FY2025']['gap_ap']))} points behind cost of sales; payables are {a('FY2025','ap_s')}% of sales against {f(IND['Acc Pay/ Sales'])}% for the industry.
   Size: {f(FP['ap_excess'] * -1)}% (payables are {mny(-F['ap_excess'])} below their level at FY2024 days). Down {PERSIST['dpo_down']} year; FY2024 had risen.

**Revised figures:** none moved income or equity. The {len(REVS)} later changes are reclassifications: {mny(-REVS[0][4])} of intangibles moved out of other assets (January 28, 2024), and intangible amortization split out of cash-flow depreciation ({mny(-REVS[1][4])} in FY2022, {mny(-REVS[2][4])} in FY2023, {mny(-REVS[3][4])} in Q1 FY2024).

**Not calculated:** the allowance percentage (both annual reports say only that it was not material); the quarterly depreciation rate (the property note is annual, and the quarter's depreciation line is mixed basis); balance-sheet measures for FY2021, the FY2022 growth gaps, accruals and M-score, and the depreciation rate before FY2024 (no earlier balance sheet or property note in the documents).
"""
assert '—' not in reply
if __name__ == '__main__':
    import os
    os.makedirs('runs/v2.12', exist_ok=True)
    with open('runs/v2.12/numbers-1b.md', 'w', encoding='utf-8') as fh:
        fh.write(reply)
    print('=====REPLY=====')
    print(reply)
    notable = '\n'.join(ln for ln in reply.splitlines() if not ln.startswith('|'))
    print('WORDS with table', len(reply.split()), 'without table', len(notable.split()))
