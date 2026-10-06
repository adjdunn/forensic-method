# Scan the numbers, step 1: Home Depot FY2025 test data. Every figure in numbers-1.md is computed here.
import openpyxl, datetime as dt, re, os, sys
BASE = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
wb = openpyxl.load_workbook(os.path.join(BASE, 'statements.xlsx'), data_only=True)

def sheet(name):
    rows = list(wb[name].iter_rows(values_only=True)); hdr = rows[0]; d = {}
    for r in rows[1:]:
        if r[0]: d[r[0]] = dict(zip(hdr[1:-1], r[1:-1]))
    return d
IS, BS, CF = sheet('IS annual'), sheet('BS annual'), sheet('CF annual')
QIS, QBS, QCF = sheet('IS quarterly'), sheet('BS quarterly'), sheet('CF quarterly')

FY = ['FY (ended 2022-01-30)','FY (ended 2023-01-29)','FY (ended 2024-01-28)','FY (ended 2025-02-02)','FY (ended 2026-02-01)']
LAB = dict(zip(FY, ['FY Jan-22','FY Jan-23','FY Jan-24','FY Feb-25','FY Feb-26']))
ends = [dt.date(2022,1,30), dt.date(2023,1,29), dt.date(2024,1,28), dt.date(2025,2,2), dt.date(2026,2,1)]
days = {FY[i]: (ends[i]-ends[i-1]).days for i in range(1,5)}; days[FY[0]] = 364
Q = ['Q2 (ended 2024-07-28)','Q2 (ended 2025-08-03)','Q2 (ended 2026-08-02)']
QLAB = {Q[1]:'Q2 Aug-25', Q[2]:'Q2 Aug-26'}
QPREV = {Q[0]:'Q1 (ended 2024-04-28)', Q[1]:'Q1 (ended 2025-05-04)', Q[2]:'Q1 (ended 2026-05-03)'}
qends = {Q[0]:(dt.date(2024,4,28),dt.date(2024,7,28)), Q[1]:(dt.date(2025,5,4),dt.date(2025,8,3)), Q[2]:(dt.date(2026,5,3),dt.date(2026,8,2))}
qdays = {q:(b-a).days for q,(a,b) in qends.items()}

def L(d, line): return {c: d[line].get(c) for c in d[line]}
sales, cogs, gp, sga, opi, pti, ni = (L(IS,'Net sales'), L(IS,'Cost of sales'), L(IS,'Gross profit'),
    L(IS,'Selling, general and administrative'), L(IS,'Operating income'),
    L(IS,'Earnings before provision for income taxes'), L(IS,'Net earnings'))
cfo = L(CF,'Net cash provided by operating activities'); capex = {c:-v for c,v in L(CF,'Capital expenditures').items()}
acq = {c:-v for c,v in L(CF,'Payments for businesses acquired, net').items()}
dacf = L(CF,'Depreciation and amortization, excluding amortization of intangible assets'); iam = L(CF,'Intangible asset amortization')
rec, inv, ap, ta, ca, ppe, cl, ltd = (L(BS,'Receivables, net'), L(BS,'Merchandise inventories'), L(BS,'Accounts payable'),
    L(BS,'Total assets'), L(BS,'Total current assets'), L(BS,'Net property and equipment'),
    L(BS,'Total current liabilities'), L(BS,'Long-term debt, excluding current installments'))

# ---- Annual-report text: PP&E note (p. 58 of each report) and acquisition notes (p. 74 of each) ----
def pages(fname):
    t = open(os.path.join(BASE, fname), encoding='utf-8', errors='replace').read()
    out, cur = {}, None
    for line in t.split('\n'):
        m = re.match(r'\[p\. (\d+)\]', line.strip())
        if m: cur = int(m.group(1)); out.setdefault(cur, [])
        elif cur is not None: out[cur].append(line)
    return {k:'\n'.join(v) for k,v in out.items()}
AR25 = pages('annual-report-FY2025-02-02-key-sections.txt'); AR26 = pages('annual-report-FY2026-02-01-key-sections.txt')
def num(s): return float(s.replace('$','').replace(',',''))
def row(page_text, label, n):
    m = re.search(re.escape(label) + r'\s*\|\s*' + r'\s*\|\s*'.join([r'\$?\(?([\d,]+)\)?']*n), page_text)
    assert m, label
    return [num(x) for x in m.groups()]
# PP&E at cost, land, construction in progress, depreciation (incl. finance lease amortization)
c25 = row(AR25[58], 'Property and equipment, at cost', 2); land25 = row(AR25[58], 'Land', 2); cip25 = row(AR25[58], 'Construction in progress', 2)
c26 = row(AR26[58], 'Property and equipment, at cost', 2); land26 = row(AR26[58], 'Land', 2); cip26 = row(AR26[58], 'Construction in progress', 2)
dep25 = row(AR25[58], 'Depreciation and finance lease amortization expense', 3)  # fiscal 2024, 2023, 2022
dep26 = row(AR26[58], 'Depreciation and finance lease amortization expense', 3)  # fiscal 2025, 2024, 2023
assert c25[0] == c26[1] and dep25[0] == dep26[1]
gross = {FY[2]: c25[1]-land25[1]-cip25[1], FY[3]: c25[0]-land25[0]-cip25[0], FY[4]: c26[0]-land26[0]-cip26[0]}
dep = {FY[1]: dep25[2], FY[2]: dep25[1], FY[3]: dep25[0], FY[4]: dep26[0]}
# Acquisitions: SRS (FY2025 report p. 74), GMS (FY2026 report p. 74)
srs_rec = row(AR25[74], 'Receivables', 1)[0]; srs_inv = row(AR25[74], 'Merchandise inventories', 1)[0]
gms_rec = row(AR26[74], 'Receivables', 1)[0]; gms_inv = row(AR26[74], 'Merchandise inventories', 1)[0]
assert 'June 18, 2024' in AR25[74] and 'September 4, 2025' in AR26[74]
QUOTE = 'The valuation allowance related to these receivables was not material'
assert QUOTE in AR26[48] and QUOTE in AR25[48]
# Revisions sheet: reclassifications only?
REV = list(wb['Revisions'].iter_rows(values_only=True))[1:]
assert len(REV) == 4 and REV[1][6] + iam[FY[1]] == REV[1][4] and REV[2][6] + iam[FY[2]] == REV[2][4]
assert REV[0][8] == -BS['Intangible assets, net'][FY[2]]
IND = {r[0]: r[1] for r in wb['Industry'].iter_rows(values_only=True) if r[1] is not None}

# ---- Annual measures ----
A = {}
for i, c in enumerate(FY):
    p = FY[i-1] if i else None; m = {}
    m['gm'] = gp[c]/sales[c]*100; m['om'] = opi[c]/sales[c]*100
    m['fcf'] = cfo[c]-capex[c]; m['fcf_ni'] = m['fcf']/ni[c]; m['ni_fcf'] = ni[c]-m['fcf']; m['cfo_ni'] = cfo[c]-ni[c]
    if c in rec:
        m['dso'] = rec[c]/sales[c]*days[c]; m['dio'] = inv[c]/cogs[c]*days[c]; m['dpo'] = ap[c]/cogs[c]*days[c]
        m['rec_s'] = rec[c]/sales[c]*100; m['inv_s'] = inv[c]/sales[c]*100; m['ap_s'] = ap[c]/sales[c]*100
        m['cx_dep'] = capex[c]/dacf[c]
    if p and p in rec:
        sg = sales[c]/sales[p]-1; cg = cogs[c]/cogs[p]-1
        m['sg'] = sg*100; m['cg'] = cg*100
        m['rg'] = (rec[c]/rec[p]-1)*100; m['ig'] = (inv[c]/inv[p]-1)*100; m['ag'] = (ap[c]/ap[p]-1)*100
        m['gap_r'] = m['rg']-m['sg']; m['gap_i'] = m['ig']-m['cg']; m['gap_p'] = m['ag']-m['cg']
        m['accr'] = (ni[c]-cfo[c])/((ta[c]+ta[p])/2)*100
        m['xs_rec'] = rec[c]-sales[c]*rec[p]/sales[p]; m['xs_inv'] = inv[c]-cogs[c]*inv[p]/cogs[p]; m['xs_ap'] = ap[c]-cogs[c]*ap[p]/cogs[p]
        DSRI = (rec[c]/sales[c])/(rec[p]/sales[p]); GMI = (gp[p]/sales[p])/(gp[c]/sales[c])
        AQI = (1-(ca[c]+ppe[c])/ta[c])/(1-(ca[p]+ppe[p])/ta[p]); SGI = sales[c]/sales[p]
        DEPI = (dep[p]/(dep[p]+ppe[p]))/(dep[c]/(dep[c]+ppe[c])); SGAI = (sga[c]/sales[c])/(sga[p]/sales[p])
        TATA = (ni[c]-cfo[c])/ta[c]; LVGI = ((ltd[c]+cl[c])/ta[c])/((ltd[p]+cl[p])/ta[p])
        m['M'] = -4.84+0.920*DSRI+0.528*GMI+0.404*AQI+0.892*SGI+0.115*DEPI-0.172*SGAI+4.679*TATA-0.327*LVGI
    if c in gross and p in gross:
        m['deprate'] = dep[c]/((gross[c]+gross[p])/2)*100
    A[c] = m
# ---- Quarterly measures (Q2 against Q2 a year earlier) ----
B = {}
for i, q in enumerate(Q):
    s = QIS['Net sales'][q]; cg_ = QIS['Cost of sales'][q]; r = QBS['Receivables, net'][q]; iv = QBS['Merchandise inventories'][q]
    a = QBS['Accounts payable'][q]; n = QIS['Net earnings'][q]; c_ = QCF['Net cash provided by operating activities'][q]
    cx = -QCF['Capital expenditures'][q]; t = QBS['Total assets'][q]; tp = QBS['Total assets'][QPREV[q]]
    m = dict(gm=QIS['Gross profit'][q]/s*100, om=QIS['Operating income'][q]/s*100, dso=r/s*qdays[q], dio=iv/cg_*qdays[q], dpo=a/cg_*qdays[q],
             accr=(n-c_)/((t+tp)/2)*100, fcf=c_-cx, fcf_ni=(c_-cx)/n, ni_fcf=n-(c_-cx), cx_dep=cx/QCF['Depreciation and amortization, excluding amortization of intangible assets'][q],
             acq=-QCF['Payments for businesses acquired, net'][q])
    if i:
        p = Q[i-1]; m['sg'] = (s/QIS['Net sales'][p]-1)*100; m['cg'] = (cg_/QIS['Cost of sales'][p]-1)*100
        m['rg'] = (r/QBS['Receivables, net'][p]-1)*100; m['ig'] = (iv/QBS['Merchandise inventories'][p]-1)*100; m['ag'] = (a/QBS['Accounts payable'][p]-1)*100
        m['gap_r'] = m['rg']-m['sg']; m['gap_i'] = m['ig']-m['cg']; m['gap_p'] = m['ag']-m['cg']
    B[q] = m

# ---- Sizes and persistence ----
LY, PY = FY[4], FY[3]; PTI = pti[LY]
xs_rec_ly = A[LY]['xs_rec']; xs_rec_2y = rec[LY]-sales[LY]*rec[FY[2]]/sales[FY[2]]
xs_inv_ly = A[LY]['xs_inv']; ap_short = A[LY]['xs_ap']; ni_fcf_ly = A[LY]['ni_fcf']
om_drop = (A[PY]['om']-A[LY]['om'])/100*sales[LY]
dep_not_taken = (A[PY]['deprate']-A[LY]['deprate'])/100*((gross[LY]+gross[PY])/2)
def streak(vals, pred):
    n = 0
    for v in reversed(vals):
        if pred(v): n += 1
        else: break
    return n
yrs = FY[2:]
p_rec = streak([A[c]['gap_r'] for c in yrs], lambda v: v > 0); p_rec_q = streak([B[q]['gap_r'] for q in Q[1:]], lambda v: v > 0)
p_inv = streak([A[c]['gap_i'] for c in yrs], lambda v: v > 0); p_inv_q = streak([B[q]['gap_i'] for q in Q[1:]], lambda v: v > 0)
p_ap = streak([A[c]['gap_p'] for c in yrs], lambda v: v < 0); p_ap_q = streak([B[q]['gap_p'] for q in Q[1:]], lambda v: v < 0)
accr_series = [A[c]['accr'] for c in yrs]; p_accr = streak([accr_series[i]-accr_series[i-1] for i in range(1,len(accr_series))], lambda v: v > 0)
fcfni = [A[c]['fcf_ni'] for c in FY]; p_fcf = streak([fcfni[i]-fcfni[i-1] for i in range(1,5)], lambda v: v < 0)
oms = [A[c]['om'] for c in FY]; p_om = streak([oms[i]-oms[i-1] for i in range(1,5)], lambda v: v < 0)
omq = [B[q]['om'] for q in Q]; p_om_q = streak([omq[i]-omq[i-1] for i in range(1,3)], lambda v: v < 0)

# ---- Table ----
cols = FY + Q[1:]
def hdr(): return '| Measure | ' + ' | '.join([LAB[c] for c in FY] + [QLAB[q] for q in Q[1:]]) + ' | Industry |'
def val(src, key, fmt):
    v = src.get(key); return '' if v is None else fmt.format(v)
def tr(name, key, fmt, ind=''):
    cells = [val(A[c], key, fmt) for c in FY] + [val(B[q], key, fmt) for q in Q[1:]]
    return f'| {name} | ' + ' | '.join(cells) + f' | {ind} |'
f1, s1, f2, m0 = '{:.1f}', '{:+.1f}', '{:.2f}', '{:,.0f}'
table = '\n'.join([hdr(), '|---|' + '---:|'*(len(cols)+1),
    tr('Days sales outstanding', 'dso', f1),
    tr('Days inventory', 'dio', f1),
    tr('Days payable', 'dpo', f1),
    tr('Receivables growth less sales growth, pts', 'gap_r', s1),
    tr('Inventory growth less cost-of-sales growth, pts', 'gap_i', s1),
    tr('Payables growth less cost-of-sales growth, pts', 'gap_p', s1),
    tr('Receivables / sales, %', 'rec_s', f1, f1.format(IND['Acc Rec/ Sales'])),
    tr('Inventory / sales, %', 'inv_s', f1, f1.format(IND['Inventory/Sales'])),
    tr('Payables / sales, %', 'ap_s', f1, f1.format(IND['Acc Pay/ Sales'])),
    tr('Gross margin, %', 'gm', f1, f1.format(IND['Gross Margin'])),
    tr('Operating margin, %', 'om', f1, f1.format(IND['Pre-tax Unadjusted Operating Margin'])),
    tr('Accruals, % of average total assets', 'accr', f1),
    tr('Free cash flow, $M', 'fcf', m0),
    tr('Free cash flow / net income', 'fcf_ni', f2),
    tr('Depreciation rate, % of gross depreciable assets', 'deprate', f2),
    tr('Capex / depreciation', 'cx_dep', f2, f1.format(IND['Cap Ex/Deprecn'])),
    tr('Beneish M-score', 'M', f2),
])
pct = lambda x: f'{x/PTI*100:.1f}%'
d = lambda c: f'${c:,.0f}M'
out = f"""{table}

Cost line: "Cost of sales" (IS annual, IS quarterly); no substitute cost line. Substitutes: the depreciation rate and the M-score's DEPI use "Depreciation and finance lease amortization expense" and gross property and equipment at cost less land and construction in progress from the annual reports (p. 58 of each), so the rate exists for two years only; capex/depreciation uses the cash-flow line "Depreciation and amortization, excluding amortization of intangible assets"; M-score securities set to 0 (none on the balance sheet) and leverage = long-term debt excluding current installments plus total current liabilities; days are actual (364; {days[PY]} for the 53-week FY Feb-25; {qdays[Q[2]]} per quarter); quarterly accruals use the average of the quarter's opening and closing total assets. The latest quarter is Q2 (ended 2026-08-02), so the mixed-basis Q4 cash-flow columns were not used; the results releases carry only full-year cash flows. The company calls FY Feb-26 "fiscal 2025".

**Flags**

1. **Receivables outgrew sales.** Receivables, net (BS annual) rose {A[PY]['rg']:.1f}% in FY Feb-25 against sales +{A[PY]['sg']:.1f}% and {A[LY]['rg']:.1f}% in FY Feb-26 against +{A[LY]['sg']:.1f}%; days sales outstanding {A[FY[2]]['dso']:.1f} to {A[PY]['dso']:.1f} to {A[LY]['dso']:.1f}; Q2 Aug-26 receivables +{B[Q[2]]['rg']:.1f}% against sales +{B[Q[2]]['sg']:.1f}%, with a {d(B[Q[2]]['acq'])} acquisition outflow in the quarter (CF quarterly). Beside it: SRS (June 18, 2024, {d(acq[PY])} paid, receivables acquired {d(srs_rec)}) and GMS (September 4, 2025, {d(acq[LY])} paid, {d(gms_rec)}); receivables/sales {A[LY]['rec_s']:.1f}% against the industry {IND['Acc Rec/ Sales']:.1f}%.
   Size: the FY Feb-26 balance less the balance at the FY Feb-25 ratio, {d(xs_rec_ly)}, {pct(xs_rec_ly)} of pre-tax income ({d(PTI)}); at the FY Jan-24 ratio, before SRS, {d(xs_rec_2y)}, {pct(xs_rec_2y)}. Persistence: {p_rec} consecutive years of receivables growth above sales growth; {p_rec_q} consecutive Q2 comparisons.

2. **Inventory outgrew cost of sales.** Merchandise inventories (BS annual) +{A[PY]['ig']:.1f}% against cost of sales +{A[PY]['cg']:.1f}% in FY Feb-25 and +{A[LY]['ig']:.1f}% against +{A[LY]['cg']:.1f}% in FY Feb-26; days inventory {A[FY[2]]['dio']:.1f} to {A[PY]['dio']:.1f} to {A[LY]['dio']:.1f}; Q2 Aug-26 inventory +{B[Q[2]]['ig']:.1f}% against cost of sales +{B[Q[2]]['cg']:.1f}%. GMS brought {d(gms_inv)} of inventories (SRS {d(srs_inv)} the year before); the level, {A[LY]['inv_s']:.1f}% of sales, sits below the industry {IND['Inventory/Sales']:.1f}%.
   Size: the balance less the balance at the prior-year ratio, {d(xs_inv_ly)}, {pct(xs_inv_ly)} of pre-tax income. Persistence: {p_inv} consecutive years; {p_inv_q} consecutive Q2 comparisons.

3. **Operating cash flow fell faster than earnings.** Net cash provided by operating activities (CF annual) {d(cfo[FY[2]])} to {d(cfo[PY])} to {d(cfo[LY])} while net earnings went {d(ni[FY[2]])} to {d(ni[PY])} to {d(ni[LY])}; free cash flow / net income {A[FY[2]]['fcf_ni']:.2f} to {A[PY]['fcf_ni']:.2f} to {A[LY]['fcf_ni']:.2f}; accruals {A[FY[2]]['accr']:.1f}% to {A[PY]['accr']:.1f}% to {A[LY]['accr']:.1f}% of average assets. The Q2 comparison moved the other way ({B[Q[1]]['fcf_ni']:.2f} to {B[Q[2]]['fcf_ni']:.2f}).
   Size: net earnings less free cash flow, {d(ni_fcf_ly)} in FY Feb-26, {pct(ni_fcf_ly)} of pre-tax income (operating cash flow still exceeds net earnings by {d(A[LY]['cfo_ni'])}). Persistence: {p_accr} consecutive years of rising accruals and {p_fcf} of falling free cash flow / net income; not continued in the latest quarter.

4. **Payables fell while cost of sales rose.** Accounts payable (BS annual) {A[LY]['ag']:.1f}% in FY Feb-26 against cost of sales +{A[LY]['cg']:.1f}%, days payable {A[PY]['dpo']:.1f} to {A[LY]['dpo']:.1f}, after +{A[PY]['ag']:.1f}% against +{A[PY]['cg']:.1f}% in FY Feb-25; Q2 payables growth trailed cost-of-sales growth in both Q2 comparisons ({B[Q[1]]['gap_p']:+.1f} and {B[Q[2]]['gap_p']:+.1f} pts). Payables/sales {A[LY]['ap_s']:.1f}% against the industry {IND['Acc Pay/ Sales']:.1f}%, {IND['Acc Pay/ Sales']-A[LY]['ap_s']:.1f} points below.
   Size: the balance less the balance at the prior-year ratio, {d(-ap_short)} short, {pct(-ap_short)} of pre-tax income. Persistence: {p_ap} year; {p_ap_q} consecutive Q2 comparisons.

5. **Operating margin down three years running.** Operating income / net sales (IS annual) {A[FY[1]]['om']:.1f}% to {A[FY[2]]['om']:.1f}% to {A[PY]['om']:.1f}% to {A[LY]['om']:.1f}% from FY Jan-23 to FY Feb-26, gross margin {A[FY[1]]['gm']:.1f}% to {A[LY]['gm']:.1f}%, Q2 {B[Q[1]]['om']:.1f}% to {B[Q[2]]['om']:.1f}%. The level is {A[LY]['om']-IND['Pre-tax Unadjusted Operating Margin']:.1f} points above the industry {IND['Pre-tax Unadjusted Operating Margin']:.1f}%; gross margin {IND['Gross Margin']-A[LY]['gm']:.1f} points below the industry {IND['Gross Margin']:.1f}%.
   Size: the FY Feb-26 margin decline applied to FY Feb-26 sales, {d(om_drop)}, {pct(om_drop)} of pre-tax income. Persistence: {p_om} consecutive years; {p_om_q} consecutive Q2 comparisons.

Not flagged: the depreciation rate moved {A[PY]['deprate']:.2f}% to {A[LY]['deprate']:.2f}% (charge not taken at the prior rate {d(dep_not_taken)}, {pct(dep_not_taken)} of pre-tax income); the M-score sits below -1.78 in all three years.

**Revised figures:** the Revisions sheet lists four changes, all presentation splits with no effect on net earnings or equity: Other assets at 2024-01-28 cut by {d(-REV[0][8])} when intangible assets were given their own line, and cash-flow "Depreciation and amortization, excluding amortization of intangible assets" cut by {d(-REV[1][8])}, {d(-REV[2][8])} and {d(-REV[3][8])} (FY Jan-23, FY Jan-24, Q1 Apr-24) when intangible amortization was given its own line.

**Not calculated:** allowance for doubtful accounts as a share of gross receivables (neither annual report states the amount; each says only "{QUOTE}", p. 48); depreciation rate for FY Jan-24 and earlier (gross property and equipment at 2023-01-29 not in the documents); days and growth gaps for FY Jan-22 (no balance sheet); quarterly M-score and depreciation rate.
"""
os.makedirs(os.path.join(BASE, 'runs', 'v2.7'), exist_ok=True)
open(os.path.join(BASE, 'runs', 'v2.7', 'numbers-1.md'), 'w', encoding='utf-8').write(out)
print(out)
print('---WORDS---', len(out.split()))
