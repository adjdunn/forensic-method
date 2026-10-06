# Scan the numbers: The Home Depot, statements.xlsx + annual reports (key sections). All figures USD millions.
import openpyxl, re, datetime as dt, os
BASE = os.path.dirname(os.path.abspath(__file__))
COMPANY = os.path.abspath(os.path.join(BASE, '..', '..'))
wb = openpyxl.load_workbook(os.path.join(COMPANY, 'statements.xlsx'), data_only=True)

def sheet(name):
    ws = wb[name]; rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]; d = {}
    for r in rows[1:]:
        if r[0] is None: continue
        d[r[0]] = {hdr[i]: r[i] for i in range(1, len(hdr)-1)}
    return hdr[1:-1], d

IS_cols, IS = sheet('IS annual'); BS_cols, BS = sheet('BS annual'); CF_cols, CF = sheet('CF annual')
ISQ_cols, ISQ = sheet('IS quarterly'); BSQ_cols, BSQ = sheet('BS quarterly'); CFQ_cols, CFQ = sheet('CF quarterly')

def enddate(col):
    return dt.date.fromisoformat(re.search(r'ended (\d{4}-\d{2}-\d{2})', col).group(1))

# Fiscal years with a balance sheet (FY ended Jan-23 .. Feb-26); FY ended Jan-22 is used only as the prior year for growth.
FY = ['FY (ended 2023-01-29)', 'FY (ended 2024-01-28)', 'FY (ended 2025-02-02)', 'FY (ended 2026-02-01)']
FY_ALL = ['FY (ended 2022-01-30)'] + FY
def fydays(col):
    i = FY_ALL.index(col)
    if i == 0: return 364  # 52-week year ended 2022-01-30 (started 2021-02-01)
    return (enddate(col) - enddate(FY_ALL[i-1])).days
DAYS = {c: fydays(c) for c in FY_ALL}

# Latest quarter beside the same quarter a year earlier
QL, QP = 'Q2 (ended 2026-08-02)', 'Q2 (ended 2025-08-03)'
QL_prev, QP_prev = 'Q1 (ended 2026-05-03)', 'Q1 (ended 2025-05-04)'   # quarter before each, for average assets
QDAYS = {QL: (enddate(QL) - enddate(QL_prev)).days, QP: (enddate(QP) - enddate(QP_prev)).days}

# Lines (sheet names kept here)
S  = lambda c: IS['Net sales'][c]
COGS = lambda c: IS['Cost of sales'][c]
GP = lambda c: IS['Gross profit'][c]
SGA = lambda c: IS['Selling, general and administrative'][c]
OI = lambda c: IS['Operating income'][c]
PBT = lambda c: IS['Earnings before provision for income taxes'][c]
NI = lambda c: IS['Net earnings'][c]
AR = lambda c: BS['Receivables, net'][c]
INV = lambda c: BS['Merchandise inventories'][c]
AP = lambda c: BS['Accounts payable'][c]
CA = lambda c: BS['Total current assets'][c]
PPE = lambda c: BS['Net property and equipment'][c]
TA = lambda c: BS['Total assets'][c]
CL = lambda c: BS['Total current liabilities'][c]
LTD = lambda c: BS['Long-term debt, excluding current installments'][c]
CFO = lambda c: CF['Net cash provided by operating activities'][c]
CAPEX = lambda c: -CF['Capital expenditures'][c]
ACQ = lambda c: -CF['Payments for businesses acquired, net'][c]
# Cash-flow D&A excluding intangible amortization; FY Jan-23 and FY Jan-24 taken as revised (Revisions sheet) so the series is on one basis
DA = {c: CF['Depreciation and amortization, excluding amortization of intangible assets'][c] for c in FY_ALL}
DA['FY (ended 2023-01-29)'] = 2796; DA['FY (ended 2024-01-28)'] = 3061

# From the annual reports (note "Net Property and Equipment", p. 58 of each): PP&E at cost, land, construction in progress;
# depreciation and finance lease amortization expense (includes amounts in cost of sales).
PPE_COST = {'FY (ended 2024-01-28)': (53257, 9027, 1192), 'FY (ended 2025-02-02)': (55783, 9060, 1521), 'FY (ended 2026-02-01)': (59467, 9499, 1654)}
GROSS_DEP = {c: v[0]-v[1]-v[2] for c, v in PPE_COST.items()}
DEP_EXP = {'FY (ended 2024-01-28)': 3020, 'FY (ended 2025-02-02)': 3283, 'FY (ended 2026-02-01)': 3452}

pct = lambda a, b: 100.0*a/b
M = {}  # measures[row][col]
def put(row, col, val): M.setdefault(row, {})[col] = val

for i, c in enumerate(FY):
    p = FY_ALL[FY_ALL.index(c)-1]
    d = DAYS[c]
    put('DSO (days)', c, AR(c)/S(c)*d)
    put('Days inventory', c, INV(c)/COGS(c)*d)
    put('Days payable', c, AP(c)/COGS(c)*d)
    if i > 0:
        put('Receivables growth minus sales growth (pts)', c, pct(AR(c)-AR(p), AR(p)) - pct(S(c)-S(p), S(p)))
        put('Inventory growth minus cost-of-sales growth (pts)', c, pct(INV(c)-INV(p), INV(p)) - pct(COGS(c)-COGS(p), COGS(p)))
        put('Payables growth minus cost-of-sales growth (pts)', c, pct(AP(c)-AP(p), AP(p)) - pct(COGS(c)-COGS(p), COGS(p)))
    put('Gross margin (%)', c, pct(GP(c), S(c)))
    put('Operating margin (%)', c, pct(OI(c), S(c)))
    put('SG&A / sales (%)', c, pct(SGA(c), S(c)))
    if i > 0:
        put('Accruals (NI - CFO) / avg assets (%)', c, pct(NI(c)-CFO(c), (TA(c)+TA(p))/2))
    put('Free cash flow / net income (x)', c, (CFO(c)-CAPEX(c))/NI(c))
    put('Free cash flow minus net income', c, CFO(c)-CAPEX(c)-NI(c))
    if c in DEP_EXP and p in GROSS_DEP:
        put('Depreciation rate (%)', c, pct(DEP_EXP[c], (GROSS_DEP[c]+GROSS_DEP[p])/2))
    put('Receivables / sales (%)', c, pct(AR(c), S(c)))
    put('Inventory / sales (%)', c, pct(INV(c), S(c)))
    put('Payables / sales (%)', c, pct(AP(c), S(c)))
    put('Capex / depreciation (x)', c, CAPEX(c)/DA[c])
    put('Acquisitions, cash paid', c, ACQ(c))
    if i > 0:
        DSRI = (AR(c)/S(c))/(AR(p)/S(p))
        GMI = (GP(p)/S(p))/(GP(c)/S(c))
        AQI = (1-(CA(c)+PPE(c))/TA(c))/(1-(CA(p)+PPE(p))/TA(p))
        SGI = S(c)/S(p)
        DEPI = (DA[p]/(DA[p]+PPE(p)))/(DA[c]/(DA[c]+PPE(c)))
        SGAI = (SGA(c)/S(c))/(SGA(p)/S(p))
        TATA = (NI(c)-CFO(c))/TA(c)
        LVGI = ((LTD(c)+CL(c))/TA(c))/((LTD(p)+CL(p))/TA(p))
        m = -4.84 + 0.920*DSRI + 0.528*GMI + 0.404*AQI + 0.892*SGI + 0.115*DEPI - 0.172*SGAI + 4.679*TATA - 0.327*LVGI
        put('Beneish M-score', c, m)
        M.setdefault('_inputs', {})[c] = dict(DSRI=DSRI, GMI=GMI, AQI=AQI, SGI=SGI, DEPI=DEPI, SGAI=SGAI, TATA=TATA, LVGI=LVGI)

# Quarterly (IS Q2 columns are as reported; CF Q2 is derived from two 10-Qs of the same year, not marked mixed basis)
Sq = lambda c: ISQ['Net sales'][c]; COGSq = lambda c: ISQ['Cost of sales'][c]; GPq = lambda c: ISQ['Gross profit'][c]
OIq = lambda c: ISQ['Operating income'][c]; NIq = lambda c: ISQ['Net earnings'][c]; SGAq = lambda c: ISQ['Selling, general and administrative'][c]
ARq = lambda c: BSQ['Receivables, net'][c]; INVq = lambda c: BSQ['Merchandise inventories'][c]; APq = lambda c: BSQ['Accounts payable'][c]; TAq = lambda c: BSQ['Total assets'][c]
CFOq = lambda c: CFQ['Net cash provided by operating activities'][c]; CAPEXq = lambda c: -CFQ['Capital expenditures'][c]; ACQq = lambda c: -CFQ['Payments for businesses acquired, net'][c]
# year-ago quarters for the prior-year quarter's growth gaps
QPP = 'Q2 (ended 2024-07-28)'
for c, prev, yago in [(QP, QP_prev, QPP), (QL, QL_prev, QP)]:
    d = QDAYS[c]
    put('DSO (days)', c, ARq(c)/Sq(c)*d)
    put('Days inventory', c, INVq(c)/COGSq(c)*d)
    put('Days payable', c, APq(c)/COGSq(c)*d)
    put('Receivables growth minus sales growth (pts)', c, pct(ARq(c)-ARq(yago), ARq(yago)) - pct(Sq(c)-Sq(yago), Sq(yago)))
    put('Inventory growth minus cost-of-sales growth (pts)', c, pct(INVq(c)-INVq(yago), INVq(yago)) - pct(COGSq(c)-COGSq(yago), COGSq(yago)))
    put('Payables growth minus cost-of-sales growth (pts)', c, pct(APq(c)-APq(yago), APq(yago)) - pct(COGSq(c)-COGSq(yago), COGSq(yago)))
    put('Gross margin (%)', c, pct(GPq(c), Sq(c)))
    put('Operating margin (%)', c, pct(OIq(c), Sq(c)))
    put('SG&A / sales (%)', c, pct(SGAq(c), Sq(c)))
    put('Accruals (NI - CFO) / avg assets (%)', c, pct(NIq(c)-CFOq(c), (TAq(c)+TAq(prev))/2))
    put('Free cash flow / net income (x)', c, (CFOq(c)-CAPEXq(c))/NIq(c))
    put('Free cash flow minus net income', c, CFOq(c)-CAPEXq(c)-NIq(c))
    put('Receivables / sales (%)', c, pct(ARq(c), Sq(c)))
    put('Inventory / sales (%)', c, pct(INVq(c), Sq(c)))
    put('Payables / sales (%)', c, pct(APq(c), Sq(c)))
    put('Acquisitions, cash paid', c, ACQq(c))

# Industry sheet (Damodaran, Retail (Building Supply))
ind = {r[0]: r[1] for r in wb['Industry'].iter_rows(values_only=True) if r[1] is not None}
IND = 'Industry'
put('Receivables / sales (%)', IND, ind['Acc Rec/ Sales']); put('Inventory / sales (%)', IND, ind['Inventory/Sales']); put('Payables / sales (%)', IND, ind['Acc Pay/ Sales'])
put('Gross margin (%)', IND, ind['Gross Margin']); put('Operating margin (%)', IND, ind['Pre-tax Unadjusted Operating Margin']); put('Capex / depreciation (x)', IND, ind['Cap Ex/Deprecn'])

# ---- Flag sizes, against latest-year pre-tax income
L, P, PP = 'FY (ended 2026-02-01)', 'FY (ended 2025-02-02)', 'FY (ended 2024-01-28)'
pbt = PBT(L)
excess_AR = AR(L) - S(L)*AR(P)/S(P)
excess_INV = INV(L) - COGS(L)*INV(P)/COGS(P)
short_AP = AP(L) - COGS(L)*AP(P)/COGS(P)
fcf_gap = NI(L) - (CFO(L)-CAPEX(L))
oi_short = S(L)*OI(P)/S(P) - OI(L)
dep_short = DEP_EXP[L] - M['Depreciation rate (%)'][P]/100*(GROSS_DEP[L]+GROSS_DEP[P])/2
sizes = {'excess receivables': excess_AR, 'excess inventory': excess_INV, 'payables shortfall': short_AP,
         'NI minus FCF': fcf_gap, 'operating income shortfall at prior margin': oi_short, 'depreciation charge not taken': dep_short}
# growth rates for the flag lines
g = lambda f, c, p: pct(f(c)-f(p), f(p))
growth = dict(AR_L=g(AR, L, P), S_L=g(S, L, P), AR_P=g(AR, P, PP), S_P=g(S, P, PP), INV_L=g(INV, L, P), COGS_L=g(COGS, L, P), INV_P=g(INV, P, PP), COGS_P=g(COGS, P, PP),
              AP_L=g(AP, L, P), CFO_L=g(CFO, L, P), NI_L=g(NI, L, P), ARq=pct(ARq(QL)-ARq(QP), ARq(QP)), Sq=pct(Sq(QL)-Sq(QP), Sq(QP)))
# persistence: consecutive same-direction periods, counted back from the latest year
def run(row, cols, sign):
    n = 0
    for c in reversed(cols):
        v = M[row].get(c)
        if v is None: break
        if (v > 0) == (sign > 0) and v != 0: n += 1
        else: break
    return n
def run_delta(row, cols, sign):
    n = 0
    for a, b in reversed(list(zip(cols, cols[1:]))):
        va, vb = M[row].get(a), M[row].get(b)
        if va is None or vb is None: break
        if ((vb-va) > 0) == (sign > 0) and vb != va: n += 1
        else: break
    return n
persist = {
 'AR gap >0 (years)': run('Receivables growth minus sales growth (pts)', FY, +1),
 'AR gap >0 (Q2 comparisons)': int(M['Receivables growth minus sales growth (pts)'][QL] > 0) + int(M['Receivables growth minus sales growth (pts)'][QP] > 0),
 'INV gap >0 (years)': run('Inventory growth minus cost-of-sales growth (pts)', FY, +1),
 'AP gap <0 (years)': run('Payables growth minus cost-of-sales growth (pts)', FY, -1),
 'Days payable falling (years)': run_delta('Days payable', FY, -1),
 'Accruals rising (years)': run_delta('Accruals (NI - CFO) / avg assets (%)', FY, +1),
 'FCF/NI falling (years)': run_delta('Free cash flow / net income (x)', FY, -1),
 'Operating margin falling (years)': run_delta('Operating margin (%)', FY, -1),
 'SG&A/sales rising (years)': run_delta('SG&A / sales (%)', FY, +1),
 'Dep rate falling (years)': run_delta('Depreciation rate (%)', FY, -1),
}

# ---- Print everything
cols = FY + [QP, QL, IND]
names = {'FY (ended 2023-01-29)': 'FY Jan-23', 'FY (ended 2024-01-28)': 'FY Jan-24', 'FY (ended 2025-02-02)': 'FY Feb-25 (53 wks)', 'FY (ended 2026-02-01)': 'FY Feb-26', QP: 'Q2 Aug-25', QL: 'Q2 Aug-26', IND: 'Industry'}
order = ['DSO (days)', 'Days inventory', 'Days payable', 'Receivables growth minus sales growth (pts)', 'Inventory growth minus cost-of-sales growth (pts)',
         'Payables growth minus cost-of-sales growth (pts)', 'Receivables / sales (%)', 'Inventory / sales (%)', 'Payables / sales (%)', 'Gross margin (%)',
         'Operating margin (%)', 'SG&A / sales (%)', 'Accruals (NI - CFO) / avg assets (%)', 'Free cash flow / net income (x)', 'Free cash flow minus net income',
         'Depreciation rate (%)', 'Capex / depreciation (x)', 'Beneish M-score', 'Acquisitions, cash paid']
def fmt(row, v):
    if v is None: return ''
    if 'M-score' in row: return f'{v:.2f}'
    if 'minus net income' in row or 'Acquisitions' in row: return f'{v:,.0f}'
    if '(x)' in row: return f'{v:.2f}'
    return f'{v:.1f}'
if __name__ == '__main__':
    print('Days:', DAYS, QDAYS)
    print('Gross depreciable assets:', GROSS_DEP)
    hdr = '| Measure | ' + ' | '.join(names[c] for c in cols) + ' |'
    print(hdr); print('|---' * (len(cols)+1) + '|')
    for r in order:
        print('| ' + r + ' | ' + ' | '.join(fmt(r, M[r].get(c)) for c in cols) + ' |')
    print('M inputs:'); [print(c, {k: round(v, 3) for k, v in d.items()}) for c, d in M['_inputs'].items()]
    print('PBT latest:', pbt)
    for k, v in sizes.items(): print(f'{k}: {v:,.0f} = {pct(v, pbt):.1f}% of PBT')
    print('growth:', {k: round(v, 1) for k, v in growth.items()})
    print('persistence:', persist)
    print('Receivable detail FY26 vs FY25 (AR p.48): customer receivables 2588 vs 1896 ->', pct(2588-1896, 1896))
    print('Industry distance: AR/S', M['Receivables / sales (%)'][L]/ind['Acc Rec/ Sales'], 'AP/S', M['Payables / sales (%)'][L]/ind['Acc Pay/ Sales'], 'INV/S', M['Inventory / sales (%)'][L]/ind['Inventory/Sales'], 'capex/dep', M['Capex / depreciation (x)'][L]/ind['Cap Ex/Deprecn'])

# ---- Verify the two phrases taken from the annual report (exact substrings on the cited page)
def page_text(fn, page):
    txt = open(os.path.join(COMPANY, fn), encoding='utf-8', errors='replace').read()
    marks = [(m.start(), int(m.group(1))) for m in re.finditer(r'\[p\. (\d+)\]', txt)]
    for i, (off, pg) in enumerate(marks):
        if pg == page:
            end = marks[i+1][0] if i+1 < len(marks) else len(txt)
            return txt[off:end]
    return ''
AR26 = 'annual-report-FY2026-02-01-key-sections.txt'
assert 'timing of vendor payments and increased inventories' in page_text(AR26, 37)
assert 'was not material' in page_text(AR26, 48) and 'valuation allowance' in page_text(AR26, 48)
assert 'GMS, which was acquired on September 4, 2025' in page_text(AR26, 33)

# ---- Compose the reply
share_rows = ('Receivables / sales (%)', 'Inventory / sales (%)', 'Payables / sales (%)')
table_rows = ['DSO (days)', 'Days inventory', 'Days payable', 'Receivables growth minus sales growth (pts)', 'Inventory growth minus cost-of-sales growth (pts)',
              'Payables growth minus cost-of-sales growth (pts)', 'Receivables / sales (%)', 'Inventory / sales (%)', 'Payables / sales (%)', 'Gross margin (%)',
              'Operating margin (%)', 'Accruals (NI - CFO) / avg assets (%)', 'Free cash flow / net income (x)', 'Depreciation rate (%)', 'Capex / depreciation (x)',
              'Beneish M-score', 'Acquisitions, cash paid']
def cell(r, c):
    if r in share_rows and c in (QL, QP): return ''
    return fmt(r, M[r].get(c))
lines = ['| Measure | ' + ' | '.join(names[c] for c in cols) + ' |', '|---|' + '---:|' * len(cols)]
for r in table_rows:
    lines.append('| ' + r + ' | ' + ' | '.join(cell(r, c) for c in cols) + ' |')
table = '\n'.join(lines)

ms = M; f1 = lambda v: f'{v:.1f}'
reply = f"""{table}

Cost line: "Cost of sales". The depreciation rate uses the annual report's depreciation and finance lease amortization expense over property and equipment at cost excluding land and construction in progress; capex against depreciation uses the cash-flow depreciation line. The latest quarter is Q2 ended Aug-26, taken as reported from the quarterly sheets, so the Q4 release columns were not needed.

**Flags**

1. Receivables outgrew sales: +{f1(growth['AR_P'])}% against +{f1(growth['S_P'])}% in FY Feb-25, +{f1(growth['AR_L'])}% against +{f1(growth['S_L'])}% in FY Feb-26, +{f1(growth['ARq'])}% against +{f1(growth['Sq'])}% in Q2 Aug-26; DSO {f1(ms['DSO (days)'][PP])} to {f1(ms['DSO (days)'][L])} days, and {f1(ms['Receivables / sales (%)'][L])}% of sales against the industry's {f1(ind['Acc Rec/ Sales'])}%. The SRS (${ACQ(P):,.0f}m, FY Feb-25) and GMS (${ACQ(L):,.0f}m, FY Feb-26) acquisitions could explain it.
   Size: excess balance at the prior year's receivables-to-sales ratio, ${excess_AR:,.0f}m, {pct(excess_AR, pbt):.1f}% of pre-tax income. Persistence: {persist['AR gap >0 (years)']} consecutive years and {persist['AR gap >0 (Q2 comparisons)']} consecutive Q2 comparisons.

2. Inventory outgrew cost of sales: +{f1(growth['INV_P'])}% against +{f1(growth['COGS_P'])}% in FY Feb-25, +{f1(growth['INV_L'])}% against +{f1(growth['COGS_L'])}% in FY Feb-26; days inventory {f1(ms['Days inventory'][PP])} to {f1(ms['Days inventory'][L])}. GMS (acquired September 4, 2025) could explain part of the latest year.
   Size: excess balance at the prior year's inventory-to-cost-of-sales ratio, ${excess_INV:,.0f}m, {pct(excess_INV, pbt):.1f}% of pre-tax income. Persistence: {persist['INV gap >0 (years)']} consecutive years, plus a positive gap in both Q2 comparisons.

3. Operating cash flow fell {abs(growth['CFO_L']):.1f}% in FY Feb-26 against a {abs(growth['NI_L']):.1f}% fall in net earnings: free cash flow to net income {ms['Free cash flow / net income (x)'][PP]:.2f}x to {ms['Free cash flow / net income (x)'][L]:.2f}x over two years, accruals {f1(ms['Accruals (NI - CFO) / avg assets (%)'][PP])}% to {f1(ms['Accruals (NI - CFO) / avg assets (%)'][L])}% of average assets. The annual report cites "timing of vendor payments and increased inventories" and a deferred tax payment (p. 37).
   Size: net earnings less free cash flow, ${fcf_gap:,.0f}m, {pct(fcf_gap, pbt):.1f}% of pre-tax income. Persistence: {persist['Accruals rising (years)']} consecutive years of rising accruals and falling cash conversion.

4. Operating margin fell three years running, {f1(ms['Operating margin (%)'][FY[0]])}% to {f1(ms['Operating margin (%)'][L])}%, with gross margin flat and SG&A up from {f1(ms['SG&A / sales (%)'][FY[0]])}% to {f1(ms['SG&A / sales (%)'][L])}% of sales; still above the industry's {f1(ind['Pre-tax Unadjusted Operating Margin'])}%.
   Size: operating income forgone at the prior year's margin, ${oi_short:,.0f}m, {pct(oi_short, pbt):.1f}% of pre-tax income. Persistence: {persist['Operating margin falling (years)']} consecutive years.

5. Payables sit below the industry and fell in the latest year: {f1(ms['Payables / sales (%)'][L])}% of sales against the industry's {f1(ind['Acc Pay/ Sales'])}% ({100-100*ms['Payables / sales (%)'][L]/ind['Acc Pay/ Sales']:.0f}% below); days payable {f1(ms['Days payable'][P])} to {f1(ms['Days payable'][L])} as payables fell {abs(growth['AP_L']):.1f}% while cost of sales rose {f1(growth['COGS_L'])}%.
   Size: shortfall against the prior year's payables-to-cost-of-sales ratio, ${abs(short_AP):,.0f}m, {pct(abs(short_AP), pbt):.1f}% of pre-tax income (part of flag 3's cash gap). Persistence: {persist['AP gap <0 (years)']} year (the prior year rose); negative in both Q2 comparisons.

**Revised figures:** the Revisions sheet lists four later changes, all reclassifications (intangible assets split out of other assets, $3,606m at Jan-24; intangible amortization split out of cash-flow depreciation, $179m FY Jan-23, $186m FY Jan-24, $52m Q1 Apr-24); none moved income or equity.

**Not calculated:** allowance for doubtful accounts (not shown in the spreadsheet or on the balance sheets; both annual reports say the valuation allowance was not material, p. 48); depreciation rate before FY Feb-25 (no property at cost for Jan-23 in the attached reports); growth gaps, accruals and M-score for FY Jan-23 (no Jan-22 balance sheet); quarterly M-score and depreciation rate.
"""
out = os.path.join(BASE, 'numbers-1.md')
with open(out, 'w', encoding='utf-8') as fh: fh.write(reply)
if __name__ == '__main__':
    print('\n' + '='*30 + ' REPLY ' + '='*30 + '\n' + reply)
    body = reply.replace(table, '')
    print('WORDS with table:', len(reply.split()), ' without table:', len(body.split()))
