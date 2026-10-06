# -*- coding: utf-8 -*-
"""Prompt 4 (compare the footnotes), fixed-input run.

Inputs: the two annual-report text files and statements.xlsx, two folders up from this file.
Every number in the reply comes from this script. Figures that exist only in the report
text are pulled through fig(), which asserts the figure is printed on the cited page.
The last section checks every quotation in the reply against the source file and page.
"""
import os, re, sys, io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.abspath(os.path.join(HERE, "..", ".."))
TXT = {y: open(os.path.join(D, f"annual-report-FY{y}-12-31-key-sections.txt"), encoding="utf-8").read()
       for y in (2012, 2013)}


# ---------- page index ----------
def page_at(year, pos):
    last = None
    for m in re.finditer(r"\[p\. (\d+)\]", TXT[year][:pos]):
        last = int(m.group(1))
    return last


def pages_of(year, s):
    t = TXT[year]
    out = []
    i = t.find(s)
    while i != -1:
        out.append(page_at(year, i))
        i = t.find(s, i + 1)
    return out


def fig(year, page, s, value=None):
    """Assert string s is printed on `page` of the FY`year` report; return it as a number."""
    pp = pages_of(year, s)
    assert page in pp, f"FY{year} p.{page}: '{s}' not found there (found on {pp})"
    if value is not None:
        return value
    return float(re.sub(r"[^0-9.\-]", "", s.replace("(", "-").replace(")", "")))


# ---------- workbook ----------
wb = openpyxl.load_workbook(os.path.join(D, "statements.xlsx"), data_only=True)


def sheet(name):
    rows = list(wb[name].iter_rows(values_only=True))
    hdr = rows[0]
    out = {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        out.setdefault(r[0], []).append(dict(zip(hdr[1:], r[1:])))
    return out


IS, BS, CF, BSQ = sheet("IS annual"), sheet("BS annual"), sheet("CF annual"), sheet("BS quarterly")


def a(sh, line, y, k=0):
    return sh[line][k][f"FY (ended {y}-12-31)"]


QEND = {2012: ["Q4 (ended 2011-12-31)", "Q1 (ended 2012-03-31)", "Q2 (ended 2012-06-30)",
               "Q3 (ended 2012-09-30)", "Q4 (ended 2012-12-31)"],
        2013: ["Q4 (ended 2012-12-31)", "Q1 (ended 2013-03-31)", "Q2 (ended 2013-06-30)",
               "Q3 (ended 2013-09-30)", "Q4 (ended 2013-12-31)"]}


def qavg(line, y):
    v = [BSQ[line][0][c] for c in QEND[y]]
    return sum(v) / len(v)


def pct(x, d=1):
    return f"{100 * x:.{d}f}%"


P = print

rev = {y: a(IS, "Total revenues", y) for y in (2011, 2012, 2013)}
carrev = {y: a(IS, "Worldwide car rental", y) for y in (2012, 2013)}
eqrev = {y: a(IS, "Worldwide equipment rental", y) for y in (2012, 2013)}
pti = {y: a(IS, "Income before income taxes", y) for y in (2011, 2012, 2013)}
P("== base (workbook, as first reported) ==")
P("revenue", rev, "growth 2013", pct(rev[2013] / rev[2012] - 1))
P("pre-tax income", pti)

# ---------- 1. depreciation-rate changes (changes in estimate), Note 8 ----------
P("\n== 1. depreciation rate changes ==")
rc12 = -fig(2012, 113, "$130.6")  # net decrease, car rental, 2012 (FY2012 report)
rc12_eq = fig(2012, 113, "an increase of $0.5 million", 0.5)
us13 = -fig(2013, 126, "$44.2 million")
us12 = -fig(2013, 126, "$139.4")
us11 = -fig(2013, 126, "$26.7")
in13 = fig(2013, 126, "net increases of $5.0 million", 5.0)
in12 = fig(2013, 126, "$8.8 million")
eq13 = -fig(2013, 126, "a decrease of $0.4 million", 0.4)
net13 = us13 + in13 + eq13
net12 = rc12 + rc12_eq
P("2012 net effect on depreciation (FY2012 report):", round(net12, 1), "cars only", rc12,
  "| check against FY2013 split:", round(us12 + in12, 1))
P("2013 net effect:", round(net13, 1), "(U.S.", us13, "intl", in13, "equipment", eq13, ")")
P("as % of pre-tax income: 2012", pct(-net12 / pti[2012]), " 2013", pct(-net13 / pti[2013]))
P("cars only 2012 as % pre-tax", pct(-rc12 / pti[2012]), "| U.S. only 2013", pct(-us13 / pti[2013]))
P("U.S. rate cuts, three-year sum:", round(-(us13 + us12 + us11), 1))

# disposal results (Note 8). Negative = gain (reduces depreciation)
adj12 = fig(2012, 113, "(96.8)")
adj13 = fig(2013, 126, "37.2")
P("adjustment on disposal: 2012", adj12, " 2013", adj13, " swing", round(adj13 - adj12, 1))
usg12 = fig(2013, 126, "$100.6")
usl13 = fig(2013, 126, "$48.2 million")
P("U.S. car disposals: 2012 gain", usg12, " 2013 loss", usl13, " swing", round(usg12 + usl13, 1))
P("FY2012 report, all car rental 2012 gain:", fig(2012, 113, "$83.2"))

# depreciation / average gross revenue earning equipment
gross = {y: qavg("Cars", y) + qavg("Other equipment", y) for y in (2012, 2013)}
dep_cf = {y: a(CF, "Depreciation of revenue earning equipment", y) for y in (2012, 2013)}
dep_gross = {2012: fig(2012, 113, "$2,165.2"), 2013: fig(2013, 126, "$2,407.8")}  # before disposal result
P("average gross revenue earning equipment (five quarter-ends):", {y: round(v, 1) for y, v in gross.items()})
for y in (2012, 2013):
    P(y, "depreciation after disposal result / avg gross:", pct(dep_cf[y] / gross[y]),
      "| before disposal result:", pct(dep_gross[y] / gross[y]))
r12 = dep_gross[2012] / gross[2012]
r13 = dep_gross[2013] / gross[2013]
gap = r12 * gross[2013] - dep_gross[2013]
P("rate change:", pct(r13 / r12 - 1), "| 2013 depreciation at the 2012 rate:", round(r12 * gross[2013], 1),
  "vs actual", dep_gross[2013], "gap", round(gap, 1), "=", pct(gap / pti[2013]), "of 2013 pre-tax income")
g2 = {y: (a(BS, "Cars", y - 1) + a(BS, "Other equipment", y - 1) + a(BS, "Cars", y) + a(BS, "Other equipment", y)) / 2
      for y in (2012, 2013)}
P("cross-check, year-end average:", {y: pct(dep_gross[y] / g2[y]) for y in g2})
r12r = fig(2013, 126, "$2,145.9") / gross[2012]  # 2012 depreciation as revised in the FY2013 report
gapr = r12r * gross[2013] - dep_gross[2013]
P("2012 as revised in FY2013 report:", pct(r12r), "| rate change vs 2013:", pct(r13 / r12r - 1),
  "| 2013 depreciation at the revised 2012 rate:", round(r12r * gross[2013], 1), "gap", round(gapr, 1),
  "=", pct(gapr / pti[2013]), "of 2013 pre-tax income")

# U.S. per-vehicle depreciation (FY2013 report segment data)
d13 = fig(2013, 54, "$1,269.3")
d12 = fig(2013, 54, "$940.6")
c13 = fig(2013, 49, "468,500")
c12 = fig(2013, 49, "358,000")
l13 = fig(2013, 49, "21,500")
l12 = fig(2013, 49, "1,100")


def pv(d, c):
    return d * 1e6 / c / 12


P("U.S. depreciation and lease charges per car per month, company-operated cars: 2012", round(pv(d12, c12), 1),
  " 2013", round(pv(d13, c13), 1), pct(pv(d13, c13) / pv(d12, c12) - 1))
P("  including leased cars: 2012", round(pv(d12, c12 + l12), 1), " 2013", round(pv(d13, c13 + l13), 1),
  pct(pv(d13, c13 + l13) / pv(d12, c12 + l12) - 1))

# ---------- 2. holding periods ----------
P("\n== 2. holding periods ==")
lo12, hi12, lo13, hi13 = 4, 28, 4, 36
fig(2012, 81, "Cars | 4 to 28 months", 0)
fig(2013, 91, "Cars | 4 to 36 months", 0)
P("range top change:", hi13 - hi12, "months", pct(hi13 / hi12 - 1), "| midpoint", (lo12 + hi12) / 2, "->", (lo13 + hi13) / 2)
P("stated U.S. average of 18 months as share of range top:", pct(18 / hi12, 0), "->", pct(18 / hi13, 0))
np_us = {2011: 0.83, 2012: 0.95, 2013: 0.91}
np_in = {2011: 0.75, 2012: 0.79, 2013: 0.76}
fig(2013, 42, "91% | 95% | 83%", 0)
fig(2013, 43, "76% | 79% | 75%", 0)
sold_us = {2013: fig(2013, 42, "197,700"), 2012: fig(2013, 42, "136,400")}
sold_in = {2013: fig(2013, 43, "64,500"), 2012: fig(2013, 43, "54,500")}
fleet_us = {2013: c13, 2012: c12}
fleet_in = {2013: fig(2013, 49, "159,700"), 2012: fig(2013, 49, "153,700")}
for y in (2012, 2013):
    s_us = (np_us[y - 1] + np_us[y]) / 2
    s_in = (np_in[y - 1] + np_in[y]) / 2
    P(y, "implied non-program holding period (avg fleet x avg non-program share / non-program cars sold): U.S.",
      round(12 * fleet_us[y] * s_us / sold_us[y], 1), "months; international",
      round(12 * fleet_in[y] * s_in / sold_in[y], 1), "months",
      "| non-program cars sold / avg fleet: U.S.", pct(sold_us[y] / fleet_us[y]), "intl", pct(sold_in[y] / fleet_in[y]))
    need = fleet_us[y] / 1.5 - sold_us[y]
    prog = fleet_us[y] * (1 - s_us)
    P("   U.S. program-car holding period needed to reach an 18-month overall average:", round(12 * prog / need, 1), "months")

# ---------- 3. disposal channels ----------
P("\n== 3. disposal channels ==")
fig(2012, 12, "33% at auction, 47% through dealer direct, 13% through our Rent2Buy program or at retail locations", 0)
fig(2013, 12, "47% at auction, 39% through dealer direct and 14% through our Rent2Buy program or at retail locations", 0)
fig(2013, 12, "83% through dealer direct, 12% at auction and 5% through our Rent2Buy program or at retail locations", 0)
w = sold_us[2013] + sold_in[2013]
P("2013 weighted by non-program cars sold: auction", pct((0.47 * sold_us[2013] + 0.12 * sold_in[2013]) / w),
  "dealer direct", pct((0.39 * sold_us[2013] + 0.83 * sold_in[2013]) / w),
  "retail/Rent2Buy", pct((0.14 * sold_us[2013] + 0.05 * sold_in[2013]) / w),
  "| 2012 worldwide: 33% / 47% / 13% / 7% other")

# ---------- 4. rental equipment ----------
P("\n== 4. rental equipment ==")
fig(2012, 81, "Other equipment | 24 to 108 months", 0)
fig(2013, 91, "Other equipment | 24 to 108 months", 0)
fig(2012, 15, "rental fleet was 43 months", 0)
fig(2013, 15, "rental fleet was 43 months", 0)
for y in (2012, 2013):
    P(y, "accumulated depreciation / gross other equipment:",
      pct(-a(BS, "Less accumulated depreciation", y, 1) / a(BS, "Other equipment", y)))
eqd = {2012: fig(2013, 54, "$272.1"), 2013: fig(2013, 54, "$298.8")}
for y in (2012, 2013):
    P(y, "equipment segment depreciation / avg gross other equipment:", pct(eqd[y] / qavg("Other equipment", y)))

# ---------- 5. property and equipment, intangibles ----------
P("\n== 5. property & equipment, intangibles ==")
for y in (2012, 2013):
    g = (a(BS, "Total property and equipment, at cost", y - 1) + a(BS, "Total property and equipment, at cost", y)) / 2
    d = a(CF, "Depreciation of property and equipment", y)
    P(y, "P&E depreciation / avg gross:", pct(d / g), "implied life", round(g / d, 1), "years")
am = {y: a(CF, "Amortization of other intangible assets", y) for y in (2012, 2013)}
ga = {2011: fig(2012, 86, "747.3"), 2012: fig(2012, 86, "1,154.3"), 2013: fig(2013, 96, "1,167.2")}
ga12r = fig(2013, 96, "1,152.2")
P("amortization / avg gross amortizable intangibles: 2012", pct(am[2012] / ((ga[2011] + ga[2012]) / 2)),
  " 2013", pct(am[2013] / ((ga12r + ga[2013]) / 2)))
fig(2012, 89, "Customer relationships | 16 | $65.0", 0)
fig(2013, 99, "Customer relationships | 16 | $65.0", 0)
fig(2012, 89, "Trademark | 20 | 7.0", 0)
fig(2013, 99, "Trademark | 20 | 7.0", 0)

# ---------- 6. receivables / allowance ----------
P("\n== 6. receivables and allowance ==")
al = {2012: fig(2012, 72, "$25,113") / 1000, 2013: fig(2013, 75, "$30.8"), "2012r": fig(2013, 75, "$29.3")}
REC = "Receivables, less allowance for doubtful accounts of $30.8 and $29.3"
rec = {y: a(BS, REC, y) for y in (2012, 2013)}
rec12r = fig(2013, 75, "1,879.7")
P("allowance / gross receivables: 2012", pct(al[2012] / (rec[2012] + al[2012]), 2),
  "| 2012 revised", pct(al["2012r"] / (rec12r + al["2012r"]), 2),
  "| 2013", pct(al[2013] / (rec[2013] + al[2013]), 2))
prov = {y: a(CF, "Provision for losses on doubtful accounts", y) for y in (2012, 2013)}
P("provision:", prov, "growth", pct(prov[2013] / prov[2012] - 1),
  "| provision / revenue:", {y: pct(prov[y] / rev[y], 2) for y in prov})
stated = {y: 0.003 * carrev[y] + 0.004 * eqrev[y] for y in (2012, 2013)}
P("bad debt implied by the stated 0.3% (car) and 0.4% (equipment):", {y: round(v, 1) for y, v in stated.items()},
  "gap to provision:", {y: round(prov[y] - stated[y], 1) for y in stated})
P("Schedule II charged to expense 2013:", fig(2013, 158, "$41.9"), "vs cash flow statement", prov[2013],
  "difference", round(prov[2013] - 41.9, 1))
P("receivables change:", pct(rec[2013] / rec[2012] - 1), "| vehicle-sale receivables:",
  fig(2012, 78, "617,267") / 1000, "->", fig(2013, 81, "357.3"))

# ---------- 7. public liability and property damage ----------
P("\n== 7. public liability and property damage reserve ==")
pl = {y: a(BS, "Public liability and property damage", y) for y in (2012, 2013)}
pf12 = fig(2013, 99, "10,197.4")
P("reserve", pl, "growth", pct(pl[2013] / pl[2012] - 1), "| / revenue: 2012 reported", pct(pl[2012] / rev[2012], 2),
  "2012 pro forma (with Dollar Thrifty)", pct(pl[2012] / pf12, 2), "2013", pct(pl[2013] / rev[2013], 2))

# ---------- 8. correction of errors ----------
P("\n== 8. correction of errors ==")
tot = fig(2013, 83, "$46.3 million")
prior = fig(2013, 83, "$34.7 million")
q4 = fig(2013, 149, "62.3")
P("errors", tot, "prior-year part", prior, "2013 part", round(tot - prior, 1), "| revised Q4 2013 pre-tax", q4,
  "errors / Q4 pre-tax", pct(tot / q4))
for r in list(wb["Revisions"].iter_rows(values_only=True))[1:]:
    if r[0] == "Income statement" and "12 months" in r[3] and r[1] in (
            "Income before income taxes", "Depreciation of revenue earning equipment and lease charges",
            "Selling, general and administrative"):
        P("  Revisions sheet:", r[1], r[3], round(r[4], 1), "->", r[6], "diff", round(r[8], 1),
          pct(r[8] / r[4]) if r[1] == "Income before income taxes" else "")
fig(2013, 85, "$23.6 million adjustment related to vendor incentives", 0)
fig(2013, 85, "$25.9 million adjustment related to vendor incentives", 0)

# ---------- 9. adjusted measures ----------
P("\n== 9. adjusted measures ==")
seg12 = fig(2012, 50, "1,247.1")
oth12 = fig(2012, 50, "(345.6)")
seg13 = fig(2013, 50, "1,581.7")
oth13 = fig(2013, 50, "(428.5)")
ab12 = seg12 + oth12 - pti[2012]
ab13 = seg13 + oth13 - pti[2013]
P("add-backs to pre-tax income: 2012", round(ab12, 1), pct(ab12 / pti[2012], 0), " 2013", round(ab13, 1),
  pct(ab13 / pti[2013], 0))
new13 = fig(2013, 50, "(40.0)") + fig(2013, 50, "(7.8)") + fig(2013, 50, "(44.0)")
P("2013 add-back lines with no counterpart in the FY2012 table (integration, relocation, impairment):", round(-new13, 1))
rpd12 = fig(2012, 49, "$40.01")
rpd11 = fig(2012, 49, "$41.33")
P("old metric, worldwide RPD 2012 vs 2011:", pct(rpd12 / rpd11 - 1))
t13 = fig(2013, 49, "$47.00")
t12 = fig(2013, 49, "$46.33")
t11 = fig(2013, 49, "$47.67")
P("new metric, U.S. Total RPD: 2012 vs 2011", pct(t12 / t11 - 1), " 2013 vs 2012", pct(t13 / t12 - 1))
P("revenue excluded from the metric: old (2012)", pct(fig(2012, 49, "(1,676.6)") / -fig(2012, 49, "$7,633.0")),
  "of car rental segment revenue; new (2013, U.S.)", pct(fig(2013, 50, "(65.0)") / -fig(2013, 50, "$6,324.4")))

# ---------- 10. other ----------
P("\n== 10. other ==")
fig(2013, 12, "| 18% | 19% | 45%", 0)
fig(2013, 147, "$279 million", 0)
P("Simply Wheelz impairment", fig(2013, 146, "$40.0 million"), "on fair value 279 -> carrying value before",
  279 + 40.0, "write-down", pct(40.0 / 319.0))
fig(2013, 137, "$40.0 million", 0)
P("Sobel possible judgment $40.0m vs 2013 pre-tax income:", pct(40.0 / pti[2013]), "| vs revised Q4 pre-tax", pct(40.0 / q4))
s1, s2, rep = 60050777, 49800405, 23200000
fig(2013, 82, "60,050,777", 0)
fig(2013, 82, "49,800,405", 0)
fig(2013, 82, "23,200,000", 0)
tre = -a(CF, "Purchase of treasury shares", 2013)
P("Sponsor shares sold in 2013:", s1 + s2, "| company bought", pct(rep / s1), "of the March offering",
  "| treasury purchases 2013:", tre, "=", pct(tre / a(IS, "Net income", 2013), 0), "of net income")
P("related-party debt:", fig(2012, 130, "$189.8 million"), "-> none")
P("Sponsor shares sold in 2013, millions:", round((s1 + s2) / 1e6, 1))
P("allowance at end of 2012, revised minus first reported:", round(al["2012r"] - al[2012], 1))
# figures cited straight from the text (asserted on the cited page)
fig(2013, 24, "approximately 20,000 vehicles", 0)
fig(2013, 24, "approximately 9,000 of our vehicles", 0)
fig(2013, 147, "$3.1 million", 0)
fig(2013, 138, "February 25, 2013 through November 4, 2013", 0)
fig(2013, 138, "November 20, 2013", 0)
fig(2013, 136, "approved by the court in June 2013", 0)
fig(2013, 137, "approved by the court in June 2013", 0)
fig(2012, 79, "approximately 26%", 0)
P("text-only figures asserted on their pages: ok")

# ---------- quotation check ----------
P("\n== quotation check ==")
reply_path = os.path.join(HERE, "reply-4-fixed-input.md")
if os.path.exists(reply_path):
    reply = open(reply_path, encoding="utf-8").read()
    cites = re.findall(r'"([^"\n]+)"\s*\((FY2012|FY2013)(?: report)?, p\. (\d+)\)', reply)
    ok = bad = 0
    for q, yr, pg in cites:
        pp = pages_of(int(yr[2:]), q)
        if int(pg) in pp:
            ok += 1
        else:
            bad += 1
            P("FAIL:", yr, "p.", pg, "found on", pp, "|", q[:90])
    allq = re.findall(r'"([^"\n]+)"', reply)
    cited = {q for q, _, _ in cites}
    uncited = [q for q in allq if q not in cited]
    for q in uncited:
        P("UNCITED QUOTE:", q[:90])
    P(f"quotations checked: {ok + bad}; exact and on the cited page: {ok}; failed: {bad}; uncited: {len(uncited)}")
    P("word count:", len(re.sub(r"[|>*#]|(?<!\w)-+(?!\w)", " ", reply).split()))
else:
    P("reply not written yet")
