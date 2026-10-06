# Turn 1: scan the numbers. All figures from statements.xlsx (USD millions),
# except the allowance for doubtful accounts, which is read from the balance-sheet
# captions in the two annual report text files.
import os, re, sys, json
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.abspath(os.path.join(HERE, "..", ".."))
wb = openpyxl.load_workbook(os.path.join(DATA, "statements.xlsx"), data_only=True)

def sheet(name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    cols = [str(h) for h in hdr[1:-1]]
    out = {}
    for r in rows[1:]:
        if r[0] is None: continue
        key = (r[0], r[-1])           # (label, XBRL element): labels repeat
        out[key] = dict(zip(cols, r[1:-1]))
    return cols, out

def get(tab, label, elem_prefix=None):
    cols, d = tab
    hits = [k for k in d if k[0] == label and (elem_prefix is None or str(k[1]).startswith(elem_prefix))]
    assert len(hits) == 1, (label, hits)
    return d[hits[0]]

ISA, ISQ = sheet("IS annual"), sheet("IS quarterly")
BSA, BSQ = sheet("BS annual"), sheet("BS quarterly")
CFA, CFQ = sheet("CF annual"), sheet("CF quarterly")

def fy(y): return f"FY (ended {y}-12-31)"
def q4(y): return f"Q4 (ended {y}-12-31)"
def q3(y): return f"Q3 (ended {y}-09-30)"

def series(tab, label, col, elem=None):
    return get(tab, label, elem).get(col)

L = dict(
 rev="Total revenues", dop="Direct operating",
 dre="Depreciation of revenue earning equipment and lease charges",
 sga="Selling, general and administrative", ni="Net income",
 pti="Income before income taxes", tax="Provision for taxes on income",
 rec="Receivables, less allowance for doubtful accounts of $30.8 and $29.3",
 inv="Inventories, at lower of cost or market", ap="Accounts payable",
 ta="Total assets", tl="Total liabilities",
 cash="Cash and cash equivalents", rcash="Restricted cash and cash equivalents",
 prep="Prepaid expenses and other assets",
 cars="Cars", oeq="Other equipment", ree_net="Total revenue earning equipment",
 ppe_gross="Total property and equipment, at cost", ppe_net="Total property and equipment",
 cfo="Net cash provided by (used in) operating activities",
 ree_capex="Revenue earning equipment expenditures",
 ree_proc="Proceeds from disposal of revenue earning equipment",
 ppe_capex="Property and equipment expenditures",
 ppe_proc="Proceeds from disposal of property and equipment",
 dep_ree="Depreciation of revenue earning equipment",
 dep_ppe="Depreciation of property and equipment",
 acq="Acquisitions, net of cash acquired",
 ree_rec="Sales of revenue earning equipment included in receivables",
)

def period(col_is, col_bs, col_bs_prev, col_cf, days, is_tab, bs_tab, cf_tab):
    g = lambda tab, k, c: series(tab, L[k], c) if c else None
    p = {}
    for k in ("rev","dop","dre","sga","ni","pti","tax"): p[k] = g(is_tab, k, col_is)
    for k in ("rec","inv","ap","ta","tl","cash","rcash","prep","cars","oeq","ree_net","ppe_gross","ppe_net"):
        p[k] = g(bs_tab, k, col_bs); p[k+"_prev"] = g(bs_tab, k, col_bs_prev)
    for k in ("cfo","ree_capex","ree_proc","ppe_capex","ppe_proc","dep_ree","dep_ppe","acq","ree_rec"):
        p[k] = g(cf_tab, k, col_cf)
    p["days"] = days
    return p

P = {}
for y, d in ((2009,365),(2010,365),(2011,365),(2012,366),(2013,365)):
    P[f"FY{y}"] = period(fy(y), fy(y) if y>=2010 else None, fy(y-1) if y>=2011 else None, fy(y), d, ISA, BSA, CFA)
for y in (2012, 2013):
    P[f"Q4 {y}"] = period(q4(y), q4(y), q3(y), q4(y), 92, ISQ, BSQ, CFQ)
P["Q4 2011"] = period(q4(2011), q4(2011), None, q4(2011), 92, ISQ, BSQ, CFQ)

def nz(*a): return all(x is not None for x in a)
def avg(a,b): return (a+b)/2 if nz(a,b) else None

M = {}
for name, p in P.items():
    m = {}
    if nz(p["rec"]): m["DSO"] = p["rec"]/p["rev"]*p["days"]
    if nz(p["inv"]): m["DIO"] = p["inv"]/p["dop"]*p["days"]
    if nz(p["ap"]):  m["DPO"] = p["ap"]/p["dop"]*p["days"]
    m["GM"] = (p["rev"]-p["dop"]-p["dre"])/p["rev"]
    m["OM"] = (p["rev"]-p["dop"]-p["dre"]-p["sga"])/p["rev"]
    m["SGA%"] = p["sga"]/p["rev"]
    m["DOP%"] = p["dop"]/p["rev"]
    ata = avg(p["ta"], p["ta_prev"])
    if ata: m["ACCR"] = (p["ni"]-p["cfo"])/ata
    m["net_fleet"] = -p["ree_capex"]-p["ree_proc"]
    m["net_ppe"]   = -p["ppe_capex"]-p["ppe_proc"]
    m["FCF"] = p["cfo"]-m["net_fleet"]-m["net_ppe"]
    m["NI"] = p["ni"]; m["CFO"] = p["cfo"]
    m["fleet_capex_to_dep"] = m["net_fleet"]/p["dep_ree"]
    m["ppe_capex_to_dep"]   = m["net_ppe"]/p["dep_ppe"]
    if nz(p["cars"], p["cars_prev"]):
        gross_now = p["cars"]+p["oeq"]; gross_prev = p["cars_prev"]+p["oeq_prev"]
        r = p["dep_ree"]/((gross_now+gross_prev)/2)
        m["fleet_dep_rate"] = r*(365/p["days"]) if name.startswith("Q4") else r
        r2 = p["dep_ppe"]/((p["ppe_gross"]+p["ppe_gross_prev"])/2)
        m["ppe_dep_rate"] = r2*(365/p["days"]) if name.startswith("Q4") else r2
    m["acq"] = p["acq"]
    M[name] = m

def growth(a,b): return a/b-1
# growth gaps (annual, and Q4 vs Q4 a year earlier)
for cur, prev in (("FY2011","FY2010"),("FY2012","FY2011"),("FY2013","FY2012"),("Q4 2013","Q4 2012"),("Q4 2012","Q4 2011")):
    c, p = P[cur], P[prev]
    m = M[cur]
    m["rev_g"] = growth(c["rev"],p["rev"]); m["dop_g"] = growth(c["dop"],p["dop"])
    m["rec_g"] = growth(c["rec"],p["rec"]); m["inv_g"] = growth(c["inv"],p["inv"]); m["ap_g"] = growth(c["ap"],p["ap"])
    m["gap_rec"] = m["rec_g"]-m["rev_g"]; m["gap_inv"] = m["inv_g"]-m["dop_g"]; m["gap_ap"] = m["ap_g"]-m["dop_g"]
    m["sga_g"] = growth(c["sga"],p["sga"])
M["FY2010"]["rev_g"] = growth(P["FY2010"]["rev"],P["FY2009"]["rev"])

# allowance: balance-sheet captions in the annual reports (FY2012 in $ thousands, FY2013 in $ millions)
def caption(fn):
    t = open(os.path.join(DATA, fn), encoding="utf-8").read()
    mm = re.search(r"Receivables, less allowance for doubtful accounts of \$([\d,\.]+) and \$([\d,\.]+) \| ([\d,\.]+) \| ([\d,\.]+)", t)
    return [float(x.replace(",","")) for x in mm.groups()]
a12 = caption("annual-report-FY2012-12-31-key-sections.txt")   # thousands: allow 2012, allow 2011, rec 2012, rec 2011
a13 = caption("annual-report-FY2013-12-31-key-sections.txt")   # millions: allow 2013, allow 2012 (as shown later), rec 2013, rec 2012 (later)
ALLOW = {
 "FY2011": (a12[1]/1000, a12[3]/1000),
 "FY2012": (a12[0]/1000, a12[2]/1000),
 "FY2013": (a13[0], a13[2]),
 "FY2012 as shown later": (a13[1], a13[3]),
}
for k,(al,net) in ALLOW.items():
    M.setdefault(k, {})["allow_pct"] = al/(al+net); M[k]["allow"] = al

# Beneish
def ben(cur, prev):
    c, p = P[cur], P[prev]
    gm = lambda x: (x["rev"]-x["dop"]-x["dre"])/x["rev"]
    ca = lambda x: x["cash"]+x["rcash"]+x["rec"]+x["inv"]+x["prep"]
    aq = lambda x: 1-(ca(x)+x["ree_net"]+x["ppe_net"])/x["ta"]
    dr = lambda x: (x["dep_ree"]+x["dep_ppe"])/((x["dep_ree"]+x["dep_ppe"])+x["ree_net"]+x["ppe_net"])
    v = dict(
      DSRI=(c["rec"]/c["rev"])/(p["rec"]/p["rev"]),
      GMI=gm(p)/gm(c),
      AQI=aq(c)/aq(p),
      SGI=c["rev"]/p["rev"],
      DEPI=dr(p)/dr(c),
      SGAI=(c["sga"]/c["rev"])/(p["sga"]/p["rev"]),
      TATA=(c["ni"]-c["cfo"])/c["ta"],
      LVGI=(c["tl"]/c["ta"])/(p["tl"]/p["ta"]),
    )
    co = dict(DSRI=.920,GMI=.528,AQI=.404,SGI=.892,DEPI=.115,SGAI=-.172,TATA=4.679,LVGI=-.327)
    score = -4.84 + sum(co[k]*v[k] for k in co)
    # contribution relative to a neutral input (index = 1, TATA = 0)
    contrib = {k: co[k]*(v[k]-(0 if k=="TATA" else 1)) for k in co}
    neutral = -4.84 + sum(co[k] for k in co if k!="TATA")
    top = sorted(contrib, key=lambda k: abs(contrib[k]), reverse=True)[:2]
    return v, score, contrib, top, neutral
BEN = {}
for cur, prev in (("FY2011","FY2010"),("FY2012","FY2011"),("FY2013","FY2012")):
    BEN[cur] = ben(cur, prev)

# ---- print
def pct(x,d=1): return "n/a" if x is None else f"{x*100:.{d}f}%"
order = ["FY2009","FY2010","FY2011","FY2012","FY2013","Q4 2011","Q4 2012","Q4 2013"]
for name in order:
    m = M[name]; print("\n==", name)
    for k,v in m.items():
        print(f"  {k:22s} {v:.4f}" if isinstance(v,float) else f"  {k:22s} {v}")
for k in ALLOW: print(k, "allowance", ALLOW[k], pct(M[k]["allow_pct"],2))
for k,(v,score,contrib,top,neutral) in BEN.items():
    print("\nBeneish", k, f"M = {score:.2f}", "neutral", round(neutral,3))
    for kk in v: print(f"   {kk:5s} {v[kk]:.3f}  contribution vs neutral {contrib[kk]:+.3f}")
    print("   top two:", top)

# extra lines for the flags
print("\nReceivables", {k: P[k]["rec"] for k in ("FY2010","FY2011","FY2012","FY2013")})
print("Sales of REE included in receivables", {k: P[k]["ree_rec"] for k in ("FY2009","FY2010","FY2011","FY2012","FY2013")})
print("Receivables change 2013", P["FY2013"]["rec"]-P["FY2012"]["rec"], "of which REE-sale receivables change", P["FY2013"]["ree_rec"]-P["FY2012"]["ree_rec"])
print("Prepaid growth 2013", growth(P["FY2013"]["prep"],P["FY2012"]["prep"]), "restricted cash growth", growth(P["FY2013"]["rcash"],P["FY2012"]["rcash"]))
print("Gross REE", {k: P[k]["cars"]+P[k]["oeq"] for k in ("FY2010","FY2011","FY2012","FY2013")})
print("Goodwill", get(BSA,"Goodwill"), "Intangibles", get(BSA,"Other intangible assets, net"))
print("Q4 pre-tax and tax", {k:(P[k]["pti"],P[k]["tax"],P[k]["ni"]) for k in ("Q4 2011","Q4 2012","Q4 2013")})
for y in (2011,2012,2013):
    for q,c in (("Q1",f"Q1 (ended {y}-03-31)"),("Q2",f"Q2 (ended {y}-06-30)"),("Q3",f"Q3 (ended {y}-09-30)"),("Q4",f"Q4 (ended {y}-12-31)")):
        r=series(ISQ,L["rev"],c); d=series(ISQ,L["dop"],c); e=series(ISQ,L["dre"],c); s=series(ISQ,L["sga"],c)
        print(y,q,"GM",pct((r-d-e)/r),"DOP%",pct(d/r),"DRE%",pct(e/r),"SGA%",pct(s/r),"OM",pct((r-d-e-s)/r))

# Revisions sheet: largest later changes
ws = wb["Revisions"]; rows = list(ws.iter_rows(values_only=True))[1:]
rev = [dict(stmt=r[0], line=r[1], period=r[3], first=r[4], first_in=r[5][:40], later=r[6], later_in=r[7][:40], diff=r[8]) for r in rows if r[0]]
print("\nRevisions rows:", len(rev))
for r in sorted(rev, key=lambda r: abs(r["diff"]), reverse=True)[:14]:
    print(f"  {r['stmt']:16s} {r['line'][:50]:50s} {r['period']:28s} {r['first']:>10.1f} -> {r['later']:>10.1f}  diff {r['diff']:+.1f} ({r['diff']/abs(r['first'])*100:+.1f}%)  [{r['first_in']} -> {r['later_in']}]")
print("  -- earnings lines")
KEEP = ("Net income","Income before income taxes","Selling, general and administrative","Direct operating",
        "Depreciation of revenue earning equipment and lease charges","Total revenues",
        "Net cash provided by (used in) operating activities","Total equity","Accumulated deficit","Total assets")
for r in rev:
    if r["line"] in KEEP and r["period"].startswith(("12 months","as at")):
        print(f"  {r['stmt']:16s} {r['line'][:50]:50s} {r['period']:28s} {r['first']:>10.1f} -> {r['later']:>10.1f}  diff {r['diff']:+.1f} ({r['diff']/abs(r['first'])*100:+.1f}%)")

# ---- markdown table for the reply
cols_out = ["FY2009","FY2010","FY2011","FY2012","FY2013","Q4 2013","Q4 2012"]
def f(v, kind):
    if v is None: return ""
    if kind == "d": return f"{v:.1f}"
    if kind == "p": return f"{v*100:.1f}%"
    if kind == "pp": return f"{v*100:+.1f}"
    if kind == "x": return f"{v:.2f}x"
    if kind == "m": return f"{v:,.1f}"
    if kind == "p2": return f"{v*100:.2f}%"
rows_out = [
 ("Days sales outstanding","DSO","d"),("Days inventory","DIO","d"),("Days payable","DPO","d"),
 ("Receivables growth minus revenue growth (pts)","gap_rec","pp"),
 ("Inventory growth minus direct operating growth (pts)","gap_inv","pp"),
 ("Payables growth minus direct operating growth (pts)","gap_ap","pp"),
 ("Allowance, % of gross receivables","allow_pct","p2"),
 ("Gross margin","GM","p"),("Operating margin","OM","p"),
 ("Accruals, % of average total assets","ACCR","p"),
 ("Net income","NI","m"),("Free cash flow","FCF","m"),
 ("Net fleet spending / fleet depreciation","fleet_capex_to_dep","x"),
 ("Fleet depreciation rate","fleet_dep_rate","p"),
 ("Net property spending / property depreciation","ppe_capex_to_dep","x"),
 ("Property depreciation rate","ppe_dep_rate","p"),
]
M["Q4 2013"]["allow_pct"] = M["FY2013"]["allow_pct"]; M["Q4 2012"]["allow_pct"] = M["FY2012"]["allow_pct"]
print()
print("| Measure | " + " | ".join(c.replace("Q4 2013","Q4 FY2013").replace("Q4 2012","Q4 FY2012") for c in cols_out) + " |")
print("|---|" + "---:|"*len(cols_out))
for label, key, kind in rows_out:
    print(f"| {label} | " + " | ".join(f(M[c].get(key), kind) for c in cols_out) + " |")
print("| Beneish M-score | | | " + " | ".join(f"{BEN[k][1]:.2f}" for k in ("FY2011","FY2012","FY2013")) + " | | |")
for k in ("FY2011","FY2012","FY2013"):
    v = BEN[k][0]; print(k, "drivers", [(t, round(v[t],3)) for t in BEN[k][3]])
print("rev growth", {k: f(M[k].get("rev_g"),"p") for k in M if "rev_g" in M[k]})
print("dop growth", {k: f(M[k].get("dop_g"),"p") for k in M if "dop_g" in M[k]})
print("rec growth", {k: f(M[k].get("rec_g"),"p") for k in M if "rec_g" in M[k]})
print("ap growth", {k: f(M[k].get("ap_g"),"p") for k in M if "ap_g" in M[k]})
print("prepaid growth 2012", growth(P["FY2012"]["prep"],P["FY2011"]["prep"]), "2011", growth(P["FY2011"]["prep"],P["FY2010"]["prep"]))
print("gross fleet growth", {y: growth(P[f"FY{y}"]["cars"]+P[f"FY{y}"]["oeq"], P[f"FY{y-1}"]["cars"]+P[f"FY{y-1}"]["oeq"]) for y in (2011,2012,2013)})
print("provision for doubtful accounts", get(CFA,"Provision for losses on doubtful accounts"))
