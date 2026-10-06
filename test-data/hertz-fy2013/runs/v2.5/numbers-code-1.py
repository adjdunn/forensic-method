# Scan the numbers: Hertz Global Holdings, FY2009-FY2013 and Q4 2013 vs Q4 2012.
# Every figure in the reply is produced here. Inputs: statements.xlsx, the two
# annual reports (key sections) and the two fourth-quarter results releases.
import openpyxl, re, os, sys
from datetime import date

BASE = r"C:\Users\aaron\OneDrive\Desktop\Master\WireSift Research\Events\CFA Vancouver Forensic\test-data\hertz-fy2013"
XLSX = os.path.join(BASE, "statements.xlsx")
AR12 = os.path.join(BASE, "annual-report-FY2012-12-31-key-sections.txt")
AR13 = os.path.join(BASE, "annual-report-FY2013-12-31-key-sections.txt")
ER12 = os.path.join(BASE, "earnings-release-Q4-FY2012-12-31.txt")
ER13 = os.path.join(BASE, "earnings-release-Q4-FY2013-12-31.txt")
OUT = os.path.join(BASE, "runs", "v2.5", "numbers-1.md")

# ---------------------------------------------------------------- workbook
wb = openpyxl.load_workbook(XLSX, data_only=True)


def load(name):
    rows = list(wb[name].iter_rows(values_only=True))
    hdr = rows[0]
    out = {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        out[(r[0], r[-1])] = dict(zip(hdr[1:-1], r[1:-1]))
    return list(hdr[1:-1]), out


ISA_H, ISA = load("IS annual")
BSA_H, BSA = load("BS annual")
CFA_H, CFA = load("CF annual")
ISQ_H, ISQ = load("IS quarterly")
BSQ_H, BSQ = load("BS quarterly")


def get(sh, label, col, element=None):
    hits = [v for (l, e), v in sh.items() if l == label and (element is None or e == element)]
    assert len(hits) == 1, (label, element, len(hits))
    v = hits[0][col]
    assert v is not None, (label, col)
    return v


def fy(y):
    return f"FY (ended {y}-12-31)"


def qcol(headers, d):
    hits = [h for h in headers if d in h]
    assert len(hits) == 1, (d, hits)
    return hits[0]


L_REV = "Total revenues"
L_DO = "Direct operating"
L_DEP = "Depreciation of revenue earning equipment and lease charges"
L_SGA = "Selling, general and administrative"
L_NI = "Net income"
L_REC = "Receivables, less allowance for doubtful accounts of $30.8 and $29.3"
L_INV = "Inventories, at lower of cost or market"
L_AP = "Accounts payable"
L_TA = "Total assets"
L_TL = "Total liabilities"
L_CARS = "Cars"
L_OEQ = "Other equipment"
L_REEN = "Total revenue earning equipment"
L_PPEG = "Total property and equipment, at cost"
L_PPEN = "Total property and equipment"
L_CFO = "Net cash provided by (used in) operating activities"
L_REEX = "Revenue earning equipment expenditures"
L_REPR = "Proceeds from disposal of revenue earning equipment"
L_PPEX = "Property and equipment expenditures"
L_PPPR = "Proceeds from disposal of property and equipment"
L_DREE = "Depreciation of revenue earning equipment"
L_DPPE = "Depreciation of property and equipment"
L_ACQ = "Acquisitions, net of cash acquired"
L_SALEREC = "Sales of revenue earning equipment included in receivables"
E_AP = "us-gaap:AccountsPayableCurrentAndNoncurrent"


def days_in_year(y):
    return (date(y, 12, 31) - date(y - 1, 12, 31)).days


Q4_DAYS = (date(2013, 12, 31) - date(2013, 9, 30)).days  # 92 in both years
assert Q4_DAYS == (date(2012, 12, 31) - date(2012, 9, 30)).days


# ---------------------------------------------------------------- text sources
def read(p):
    return open(p, encoding="utf-8").read()


T_AR12, T_AR13, T_ER12, T_ER13 = read(AR12), read(AR13), read(ER12), read(ER13)


def num(s):
    s = s.strip().replace("$", "").replace(",", "")
    neg = s.startswith("(")
    s = s.strip("()")
    return -float(s) if neg else float(s)


def table_start(text, n):
    """Position of the heading line 'Table n' (not the mention in the list of tables)."""
    hits = [m.start() for m in re.finditer(r"^Table %d\s*$" % n, text, re.M)]
    assert len(hits) == 1, (n, hits)
    return hits[0]


def page_of(text, pos):
    m = None
    for m in re.finditer(r"\[p\. ([^\]]+)\]", text[:pos]):
        pass
    return m.group(1) if m else None


NUMCELL = re.compile(r"\$?\(?[\d,]+\.?\d*\)?")


def row(text, label, start=0):
    """First table row that starts with `label` after `start`; returns its numeric cells in order."""
    m = re.compile(r"^" + re.escape(label) + r" \| (.+)$", re.M).search(text, start)
    assert m, label
    vals = []
    for c in m.group(1).split(" | "):
        c = c.strip()
        if NUMCELL.fullmatch(c):
            vals.append(num(c))
    return vals


# Q4 2013 and Q4 2012 (both on the revised basis) from the Q4 2013 results release
t1 = table_start(T_ER13, 1)
R13 = {}
for key, label in [("rev", L_REV), ("do", L_DO), ("dep", L_DEP), ("sga", L_SGA)]:
    v = row(T_ER13, label, t1)
    R13[key] = (v[0], v[1])  # (Q4 2013, Q4 2012)
m = re.compile(r"^Net loss attributable to Hertz Global Holdings[^|]*\| \$\(([\d.]+)\) \| \$\(([\d.]+)\)", re.M).search(T_ER13, t1)
R13["ni"] = (-float(m.group(1)), -float(m.group(2)))
t7 = table_start(T_ER13, 7)
v = row(T_ER13, "Revenue earning equipment expenditures", t7)
assert len(v) == 10
R13["reex"] = (-v[4], -v[9])
v = row(T_ER13, "Proceeds from disposal of revenue earning equipment", t7)
assert len(v) == 10
R13["repr"] = (v[4], v[9])
v = row(T_ER13, "Depreciation of revenue earning equipment", t7)
assert len(v) == 10
R13["dree"] = (v[4], v[9])
v = row(T_ER13, "Property and equipment expenditures, net of disposals", t7)
R13["ppenet"] = (-v[0], -v[1])
v = row(T_ER13, "Depreciation of property and equipment", t7)
R13["dppe"] = (v[0], v[1])
t8 = table_start(T_ER13, 8)
v = row(T_ER13, "Net cash provided by operating activities", t8)
R13["cfo"] = (v[0], v[1])
# the release's full-year column must tie to the workbook's FY2013 column
vfy = row(T_ER13, L_REV, T_ER13.index("Year Ended", t1))
assert abs(vfy[0] - get(ISA, L_REV, fy(2013))) < 0.05

# Q4 2012 and Q4 2011 as first reported, from the Q4 2012 results release (used for the Q4 2012 growth gaps only)
u1 = table_start(T_ER12, 1)
R12 = {}
for key, label in [("rev", L_REV), ("do", L_DO), ("dep", L_DEP)]:
    v = row(T_ER12, label, u1)
    R12[key] = (v[0], v[1])  # (Q4 2012, Q4 2011)

# Allowance for doubtful accounts: balance sheet captions, and Schedule II for 2010
m12 = re.search(r"Receivables, less allowance for doubtful accounts of \$([\d,]+) and \$([\d,]+) \| ([\d,]+) \| ([\d,]+)", T_AR12)
m13 = re.search(r"Receivables, less allowance for doubtful accounts of \$([\d.]+) and \$([\d.]+) \| ([\d,.]+) \| ([\d,.]+)", T_AR13)
ALLOW = {2012: num(m12.group(1)) / 1000, 2011: num(m12.group(2)) / 1000, 2013: num(m13.group(1))}
ALLOW_2012_REV = num(m13.group(2))
REC_2012_REV_AR = num(m13.group(4))
ms = re.search(r"Year ended December 31, 2010 \| ([\d,]+) \| ([\d,]+) \| \(?([\d,]+)\)? \| \(([\d,]+)\) \| \(a\) \| \$?([\d,]+)", T_AR12)
ALLOW[2010] = num(ms.group(5)) / 1000
PAGES = {"allow12": page_of(T_AR12, m12.start()), "allow13": page_of(T_AR13, m13.start()), "sched12": page_of(T_AR12, ms.start())}

# ---------------------------------------------------------------- Revisions sheet
REV = [r for r in wb["Revisions"].iter_rows(min_row=2, values_only=True) if r[0]]


def revised(stmt, label, period, exact=False):
    hits = [r for r in REV if r[0].startswith(stmt) and (r[1] == label if exact else r[1].startswith(label)) and r[3] == period]
    assert len(hits) == 1, (stmt, label, period, len(hits))
    return hits[0]  # [4] as first reported, [6] as shown later, [7] shown later in, [8] difference


REC12R = revised("Balance", "Receivables", "as at 2012-12-31")[6]
AP12R = revised("Balance", "Accounts payable", "as at 2012-12-31")[6]
TA12R = revised("Balance", "Total assets", "as at 2012-12-31")[6]
CARS12R = revised("Balance", "Cars", "as at 2012-12-31")[6]
assert abs(REC12R - REC_2012_REV_AR) < 0.05

# ---------------------------------------------------------------- annual measures (workbook, as originally reported)
YEARS = [2009, 2010, 2011, 2012, 2013]
A = {y: {} for y in YEARS}
for y in YEARS:
    c = fy(y)
    d = A[y]
    d["rev"] = get(ISA, L_REV, c)
    d["do"] = get(ISA, L_DO, c)
    d["dep"] = get(ISA, L_DEP, c)
    d["sga"] = get(ISA, L_SGA, c)
    d["ni"] = get(ISA, L_NI, c)
    d["cost"] = d["do"] + d["dep"]
    d["gm"] = (d["rev"] - d["cost"]) / d["rev"]
    d["om"] = (d["rev"] - d["cost"] - d["sga"]) / d["rev"]
    d["sgapct"] = d["sga"] / d["rev"]
    d["cfo"] = get(CFA, L_CFO, c)
    d["fleetnet"] = -get(CFA, L_REEX, c) - get(CFA, L_REPR, c)
    d["ppenet"] = -get(CFA, L_PPEX, c) - get(CFA, L_PPPR, c)
    d["fcf"] = d["cfo"] - d["fleetnet"] - d["ppenet"]
    d["dree"] = get(CFA, L_DREE, c)
    d["dppe"] = get(CFA, L_DPPE, c)
    d["acq"] = -get(CFA, L_ACQ, c)
    d["salerec"] = get(CFA, L_SALEREC, c)
    d["days"] = days_in_year(y)
    if c in BSA_H:
        d["rec"] = get(BSA, L_REC, c)
        d["inv"] = get(BSA, L_INV, c)
        d["ap"] = get(BSA, L_AP, c, E_AP)
        d["ta"] = get(BSA, L_TA, c)
        d["tl"] = get(BSA, L_TL, c)
        d["reegross"] = get(BSA, L_CARS, c) + get(BSA, L_OEQ, c)
        d["ppegross"] = get(BSA, L_PPEG, c)
        d["ppenetbal"] = get(BSA, L_REEN, c) + get(BSA, L_PPEN, c)
        d["ca"] = sum(get(BSA, l, c) for l in ["Cash and cash equivalents", "Restricted cash and cash equivalents",
                                              L_REC, L_INV, "Prepaid expenses and other assets"])
        d["dso"] = d["rec"] / d["rev"] * d["days"]
        d["dio"] = d["inv"] / d["cost"] * d["days"]
        d["dpo"] = d["ap"] / d["cost"] * d["days"]
        d["allow"] = ALLOW[y]
        d["allowpct"] = d["allow"] / (d["rec"] + d["allow"])
for y in YEARS:
    d, p = A[y], A.get(y - 1)
    if p and "rec" in p and "rec" in d:
        g = lambda k: d[k] / p[k] - 1
        d["g_rev"], d["g_cost"], d["g_rec"], d["g_inv"], d["g_ap"] = g("rev"), g("cost"), g("rec"), g("inv"), g("ap")
        d["g_do"], d["g_dep"] = g("do"), g("dep")
        d["gap_rec"] = d["g_rec"] - d["g_rev"]
        d["gap_inv"] = d["g_inv"] - d["g_cost"]
        d["gap_ap"] = d["g_ap"] - d["g_cost"]
        d["accr"] = (d["ni"] - d["cfo"]) / ((d["ta"] + p["ta"]) / 2)
        d["deprate_ree"] = d["dree"] / ((d["reegross"] + p["reegross"]) / 2)
        d["deprate_ppe"] = d["dppe"] / ((d["ppegross"] + p["ppegross"]) / 2)
        # Beneish M-score, eight variables
        b = {}
        b["DSRI"] = (d["rec"] / d["rev"]) / (p["rec"] / p["rev"])
        b["GMI"] = p["gm"] / d["gm"]
        aq = lambda x: 1 - (x["ca"] + x["ppenetbal"]) / x["ta"]
        b["AQI"] = aq(d) / aq(p)
        b["SGI"] = d["rev"] / p["rev"]
        dr = lambda x: (x["dree"] + x["dppe"]) / (x["dree"] + x["dppe"] + x["ppenetbal"])
        b["DEPI"] = dr(p) / dr(d)
        b["SGAI"] = (d["sga"] / d["rev"]) / (p["sga"] / p["rev"])
        b["TATA"] = (d["ni"] - d["cfo"]) / d["ta"]
        b["LVGI"] = (d["tl"] / d["ta"]) / (p["tl"] / p["ta"])
        b["M"] = (-4.84 + 0.920 * b["DSRI"] + 0.528 * b["GMI"] + 0.404 * b["AQI"] + 0.892 * b["SGI"]
                  + 0.115 * b["DEPI"] - 0.172 * b["SGAI"] + 4.679 * b["TATA"] - 0.327 * b["LVGI"])
        d["beneish"] = b

# ---------------------------------------------------------------- quarter measures (release, revised basis)
Q = {2013: {}, 2012: {}}
c13, c12, c11 = qcol(BSQ_H, "2013-12-31"), qcol(BSQ_H, "2012-12-31"), qcol(BSQ_H, "2011-12-31")
c13q3, c12q3 = qcol(BSQ_H, "2013-09-30"), qcol(BSQ_H, "2012-09-30")
bal = {
    2013: dict(rec=get(BSQ, L_REC, c13), inv=get(BSQ, L_INV, c13), ap=get(BSQ, L_AP, c13, E_AP), ta=get(BSQ, L_TA, c13),
               reegross=get(BSQ, L_CARS, c13) + get(BSQ, L_OEQ, c13), ppegross=get(BSQ, L_PPEG, c13), allow=ALLOW[2013],
               ta_prev=get(BSQ, L_TA, c13q3), reegross_prev=get(BSQ, L_CARS, c13q3) + get(BSQ, L_OEQ, c13q3),
               ppegross_prev=get(BSQ, L_PPEG, c13q3)),
    # December 2012 balances on the revised basis (Revisions sheet), to match the release's revised Q4 2012
    2012: dict(rec=REC12R, inv=get(BSQ, L_INV, c12), ap=AP12R, ta=TA12R,
               reegross=CARS12R + get(BSQ, L_OEQ, c12), ppegross=get(BSQ, L_PPEG, c12), allow=ALLOW_2012_REV,
               ta_prev=get(BSQ, L_TA, c12q3), reegross_prev=get(BSQ, L_CARS, c12q3) + get(BSQ, L_OEQ, c12q3),
               ppegross_prev=get(BSQ, L_PPEG, c12q3)),
}
for i, y in enumerate([2013, 2012]):
    d = Q[y]
    b = bal[y]
    for k in ["rev", "do", "dep", "sga", "ni", "cfo", "reex", "repr", "dree", "ppenet", "dppe"]:
        d[k] = R13[k][i]
    d["cost"] = d["do"] + d["dep"]
    d["gm"] = (d["rev"] - d["cost"]) / d["rev"]
    d["om"] = (d["rev"] - d["cost"] - d["sga"]) / d["rev"]
    d["sgapct"] = d["sga"] / d["rev"]
    d["dso"] = b["rec"] / d["rev"] * Q4_DAYS
    d["dio"] = b["inv"] / d["cost"] * Q4_DAYS
    d["dpo"] = b["ap"] / d["cost"] * Q4_DAYS
    d["allowpct"] = b["allow"] / (b["rec"] + b["allow"])
    d["accr"] = (d["ni"] - d["cfo"]) / ((b["ta"] + b["ta_prev"]) / 2)
    d["fleetnet"] = d["reex"] - d["repr"]
    d["fcf"] = d["cfo"] - d["fleetnet"] - d["ppenet"]
    d["deprate_ree"] = d["dree"] / ((b["reegross"] + b["reegross_prev"]) / 2) * 4
    d["deprate_ppe"] = d["dppe"] / ((b["ppegross"] + b["ppegross_prev"]) / 2) * 4
    d.update({"bal_" + k: v for k, v in b.items()})
# growth gaps, Q4 2013 against Q4 2012 (both revised basis)
d, p = Q[2013], Q[2012]
d["g_rev"] = d["rev"] / p["rev"] - 1
d["g_cost"] = d["cost"] / p["cost"] - 1
d["g_do"] = d["do"] / p["do"] - 1
d["g_dep"] = d["dep"] / p["dep"] - 1
d["g_sga"] = d["sga"] / p["sga"] - 1
d["g_rec"] = bal[2013]["rec"] / bal[2012]["rec"] - 1
d["g_inv"] = bal[2013]["inv"] / bal[2012]["inv"] - 1
d["g_ap"] = bal[2013]["ap"] / bal[2012]["ap"] - 1
d["gap_rec"] = d["g_rec"] - d["g_rev"]
d["gap_inv"] = d["g_inv"] - d["g_cost"]
d["gap_ap"] = d["g_ap"] - d["g_cost"]
# growth gaps, Q4 2012 against Q4 2011, both as first reported (Q4 2012 release; workbook balances as first reported)
d = Q[2012]
rev12, rev11 = R12["rev"]
cost12 = R12["do"][0] + R12["dep"][0]
cost11 = R12["do"][1] + R12["dep"][1]
assert abs(rev11 - get(ISQ, L_REV, qcol(ISQ_H, "2011-12-31"))) < 0.05
d["g_rev"] = rev12 / rev11 - 1
d["g_cost"] = cost12 / cost11 - 1
d["g_rec"] = get(BSQ, L_REC, c12) / get(BSQ, L_REC, c11) - 1
d["g_inv"] = get(BSQ, L_INV, c12) / get(BSQ, L_INV, c11) - 1
d["g_ap"] = get(BSQ, L_AP, c12, E_AP) / get(BSQ, L_AP, c11, E_AP) - 1
d["gap_rec"] = d["g_rec"] - d["g_rev"]
d["gap_inv"] = d["g_inv"] - d["g_cost"]
d["gap_ap"] = d["g_ap"] - d["g_cost"]

# The workbook's own Q4 2013 column (mixed basis), for comparison only
mixcol = qcol(ISQ_H, "2013-12-31")
assert "mixed basis" in mixcol
MIX = dict(rev=get(ISQ, L_REV, mixcol), do=get(ISQ, L_DO, mixcol), dep=get(ISQ, L_DEP, mixcol),
           sga=get(ISQ, L_SGA, mixcol), ni=get(ISQ, L_NI, mixcol))

# ---------------------------------------------------------------- revisions that moved income or equity
RV = {}
for per, tag in [("12 months ended 2011-12-31", 2011), ("12 months ended 2012-12-31", 2012)]:
    pre = revised("Income", "Income before income taxes", per)
    nis = revised("Income", "Net income", per, exact=True)
    RV[tag] = dict(pre_first=pre[4], pre_later=pre[6], pre_diff=pre[8], ni_first=nis[4], ni_later=nis[6], ni_diff=nis[8],
                   sga_diff=revised("Income", "Selling, general", per)[8],
                   dep_diff=revised("Income", "Depreciation of revenue earning equipment and lease", per)[8],
                   do_diff=revised("Income", "Direct operating", per)[8], later_in=pre[7])
eq = revised("Balance", "Total equity", "as at 2012-12-31")
RV["eq"] = dict(first=eq[4], later=eq[6], diff=eq[8])
RV["q4_2013_ni_gap"] = MIX["ni"] - Q[2013]["ni"]  # change to Q1-Q3 2013 absorbed by the derived quarter

# ---------------------------------------------------------------- diagnostics
KEYS = ["rev", "cost", "do", "dep", "sga", "sgapct", "ni", "cfo", "fleetnet", "ppenet", "fcf", "rec", "inv", "ap", "ta", "days",
        "dso", "dio", "dpo", "g_rev", "g_cost", "g_do", "g_dep", "g_sga", "g_rec", "g_inv", "g_ap", "gap_rec", "gap_inv", "gap_ap",
        "allow", "allowpct", "gm", "om", "accr", "dree", "dppe", "reegross", "ppegross", "deprate_ree", "deprate_ppe",
        "acq", "salerec"]


def pr(title, d, keys):
    print(title, {k: (round(d[k], 4) if isinstance(d.get(k), float) else d.get(k)) for k in keys if k in d})


def diagnostics():
    for y in YEARS:
        pr(f"FY{y}", A[y], KEYS)
        if "beneish" in A[y]:
            print("   Beneish", {k: round(v, 3) for k, v in A[y]["beneish"].items()})
    for y in [2012, 2013]:
        pr(f"Q4 {y}", Q[y], KEYS + ["reex", "repr"] + ["bal_" + k for k in bal[y]])
    print("Release Q4 2013 raw", R13)
    print("Release Q4 2012 raw", R12)
    print("Workbook mixed-basis Q4 2013", MIX)
    print("Allowance", ALLOW, "2012 revised", ALLOW_2012_REV, PAGES)
    print("Revisions", RV)


if __name__ == "__main__" and "--diag" in sys.argv:
    diagnostics()

# ---------------------------------------------------------------- extra figures used in the flags
L_PRETAX = "Income before income taxes"
CFQ_H, CFQ = load("CF quarterly")
X = {}
X["pretax13"] = get(ISA, L_PRETAX, fy(2013))
avg_gross13 = (A[2013]["reegross"] + A[2012]["reegross"]) / 2
X["dep_at_2011_rate"] = A[2011]["deprate_ree"] * avg_gross13 - A[2013]["dree"]
X["fleetnet_growth_11_13"] = A[2013]["fleetnet"] / A[2011]["fleetnet"] - 1
X["rev_growth_11_13"] = A[2013]["rev"] / A[2011]["rev"] - 1
cfq12 = qcol(CFQ_H, "2012-12-31")
assert "mixed basis" not in cfq12
X["acq_q4_2012"] = -get(CFQ, L_ACQ, cfq12)
for y in (2009, 2010, 2011):
    assert A[y]["fcf"] > A[y]["ni"]
for y in (2012, 2013):
    assert A[y]["fcf"] < A[y]["ni"]
for y in (2011, 2012, 2013):
    assert A[y]["beneish"]["M"] < -1.78

# ---------------------------------------------------------------- the reply
MINUS = "−"


def f1(x, signed=False):
    s = f"{abs(x):,.1f}"
    if round(x, 1) < 0:
        return MINUS + s
    return ("+" + s) if signed else s


def f2(x):
    s = f"{abs(x):,.2f}"
    return (MINUS + s) if round(x, 2) < 0 else s


def pct(x, signed=False):
    return f1(x * 100, signed)


COLS = [("FY2009", A[2009]), ("FY2010", A[2010]), ("FY2011", A[2011]), ("FY2012", A[2012]), ("FY2013", A[2013]),
        ("Q4 2012", Q[2012]), ("Q4 2013", Q[2013])]
ROWS = [
    ("Days sales outstanding", lambda d: f1(d["dso"]) if "dso" in d else ""),
    ("Days inventory", lambda d: f1(d["dio"]) if "dio" in d else ""),
    ("Days payable", lambda d: f1(d["dpo"]) if "dpo" in d else ""),
    (f"Receivables growth {MINUS} revenue growth (pts)", lambda d: pct(d["gap_rec"], True) if "gap_rec" in d else ""),
    (f"Inventory growth {MINUS} cost growth (pts)", lambda d: pct(d["gap_inv"], True) if "gap_inv" in d else ""),
    (f"Payables growth {MINUS} cost growth (pts)", lambda d: pct(d["gap_ap"], True) if "gap_ap" in d else ""),
    ("Allowance, % of gross receivables", lambda d: pct(d["allowpct"]) if "allowpct" in d else ""),
    ("Gross margin, %", lambda d: pct(d["gm"])),
    ("Operating margin, %", lambda d: pct(d["om"])),
    ("Accruals, % of average assets", lambda d: pct(d["accr"]) if "accr" in d else ""),
    ("Free cash flow ($m)", lambda d: f1(d["fcf"])),
    ("Net income ($m)", lambda d: f1(d["ni"])),
    ("Depreciation rate, rental equipment, %", lambda d: pct(d["deprate_ree"]) if "deprate_ree" in d else ""),
    ("Depreciation rate, other property, %", lambda d: pct(d["deprate_ppe"]) if "deprate_ppe" in d else ""),
    (f"Beneish M-score (look above {MINUS}1.78)", lambda d: f2(d["beneish"]["M"]) if "beneish" in d else ""),
]
table = ["| Measure | " + " | ".join(c for c, _ in COLS) + " |", "|---|" + "---:|" * len(COLS)]
for name, fn in ROWS:
    cells = [fn(d) for _, d in COLS]
    assert any(cells), name
    table.append("| " + name + " | " + " | ".join(cells) + " |")

a11, a12, a13, q12, q13 = A[2011], A[2012], A[2013], Q[2012], Q[2013]
basis = (
    "**Cost line:** *Direct operating* + *Depreciation of revenue earning equipment and lease charges*; operating margin also deducts "
    "*Selling, general and administrative*. **Substitutes:** capital spending is purchases net of disposal proceeds (fleet and property); "
    "depreciation rates use the cash flow depreciation lines over gross *Cars* + *Other equipment* and over *Total property and equipment, "
    "at cost*; the M-score, on an unclassified balance sheet, uses cash through prepaid and other assets as current assets and total "
    f"liabilities for leverage; the FY2010 allowance is from Schedule II (FY2012 annual report, p. {PAGES['sched12']}). "
    "**Quarters:** the workbook's Q4 2013 is marked mixed basis, so both Q4 columns come from the Q4 2013 results release, with December "
    "2012 balances as revised (Revisions sheet) and Q4 2012 growth gaps as first reported (Q4 2012 release); depreciation rates are "
    "×4, accruals are not annualized."
)
flags = [
    f"**Receivables fell {pct(-a13['g_rec'])}% in FY2013 while revenue rose {pct(a13['g_rev'])}%** (gap {pct(a13['gap_rec'], True)} pts, "
    f"after {pct(a11['gap_rec'], True)} and {pct(a12['gap_rec'], True)}): DSO went from {f1(a12['dso'])} to {f1(a13['dso'])} days and the "
    f"allowance from {pct(a12['allowpct'])}% ({pct(q12['allowpct'])}% as later revised) to {pct(a13['allowpct'])}% of gross receivables. "
    f"BS annual, *Receivables, less allowance for doubtful accounts*; FY2012 acquisitions ({f1(a12['acq'])} outflow) may have lifted the "
    f"opening balance.",

    f"**Q4 2013 gross margin fell from {pct(q12['gm'])}% to {pct(q13['gm'])}%** (cost line {pct(q13['g_cost'], True)}% on revenue "
    f"{pct(q13['g_rev'], True)}%) while operating margin rose from {pct(q12['om'])}% to {pct(q13['om'])}%, as SG&A fell "
    f"{pct(-q13['g_sga'])}% ({f1(q12['sga'])} to {f1(q13['sga'])}). Q4 2013 results release, Table 1, the two cost lines and *Selling, "
    f"general and administrative*; Q4 2012 was the acquisition quarter ({f1(X['acq_q4_2012'])} outflow, CF quarterly).",

    f"**The depreciation rate on rental equipment fell from {pct(a11['deprate_ree'])}% (FY2011) to {pct(a12['deprate_ree'])}% to "
    f"{pct(a13['deprate_ree'])}% (FY2013)**, and from {pct(q12['deprate_ree'])}% to {pct(q13['deprate_ree'])}% annualized in Q4; at FY2011's "
    f"rate, FY2013 depreciation would be {f1(X['dep_at_2011_rate'])} higher, against pre-tax income of {f1(X['pretax13'])}. CF annual, "
    f"*Depreciation of revenue earning equipment*, over BS annual, *Cars* + *Other equipment*; fleet acquired late in FY2012 could explain "
    f"part of that year's step.",

    f"**Accounts payable fell {pct(-a13['g_ap'])}% in FY2013 while the cost line rose {pct(a13['g_cost'])}%** (gap {pct(a13['gap_ap'], True)} "
    f"pts, after {pct(a11['gap_ap'], True)} and {pct(a12['gap_ap'], True)}), taking days payable from {f1(a12['dpo'])} to {f1(a13['dpo'])}. "
    f"BS annual, *Accounts payable*.",

    f"**Free cash flow fell below net income** in FY2012 ({f1(a12['fcf'])} against {f1(a12['ni'])}) and FY2013 ({f1(a13['fcf'])} against "
    f"{f1(a13['ni'])}), after exceeding it in FY2009 to FY2011; net fleet spending rose {pct(X['fleetnet_growth_11_13'])}% from FY2011 to "
    f"FY2013 ({f1(a11['fleetnet'])} to {f1(a13['fleetnet'])}) against revenue up {pct(X['rev_growth_11_13'])}%. CF annual, *Revenue earning "
    f"equipment expenditures* less *Proceeds from disposal of revenue earning equipment*.",
]
r11, r12, req = RV[2011], RV[2012], RV["eq"]
filed = re.search(r"filed (\d{4}-\d{2}-\d{2})", r12["later_in"]).group(1)
assert "2013-12-31" in r12["later_in"] and r11["later_in"] == r12["later_in"]
revised_line = (
    f"**Revised figures:** the FY2013 annual report (filed {filed}) cut pre-tax income by {f1(-r11['pre_diff'])} for FY2011 "
    f"({f1(r11['pre_first'])} to {f1(r11['pre_later'])}) and {f1(-r12['pre_diff'])} for FY2012 ({f1(r12['pre_first'])} to "
    f"{f1(r12['pre_later'])}), net income by {f1(-r11['ni_diff'])} and {f1(-r12['ni_diff'])}, and equity at December 2012 by "
    f"{f1(-req['diff'])} ({f1(req['first'])} to {f1(req['later'])}), with SG&A raised by {f1(r11['sga_diff'])} and {f1(r12['sga_diff'])} and "
    f"rental-equipment depreciation lowered by {f1(-r11['dep_diff'])} and {f1(-r12['dep_diff'])}; Q1 to Q3 2013 changed too, since the "
    f"workbook's derived Q4 2013 net income is {f1(MIX['ni'])} against {f1(q13['ni'])} in the release."
)
not_calc = (
    "**Not calculated:** the M-score for the quarters (it is an annual model); balance-sheet measures for FY2009 and year-over-year "
    "measures for FY2010 (the workbook's balance sheets start at December 2010); depreciable property excluding land, and Q4 2011 on the "
    "revised basis, are not in the documents."
)
reply = "\n".join(table) + "\n\n" + basis + "\n\n**Flags**\n\n" + "\n".join(f"{i}. {t}" for i, t in enumerate(flags, 1)) \
        + "\n\n" + revised_line + "\n\n" + not_calc + "\n"

# ---------------------------------------------------------------- checks on the reply text
for w in ["fraud", "manipulat", "red flag", "—"]:
    assert w not in reply.lower(), w
# every line name cited in italics must exist where it is said to be
for n in [L_DO, L_DEP, L_SGA, L_REV]:
    assert any(l == n for (l, e) in ISA), n
    assert re.search(r"^" + re.escape(n) + r" \| ", T_ER13[t1:t7], re.M), n
assert any(l.startswith("Receivables, less allowance for doubtful accounts") for (l, e) in BSA)
for n in [L_CARS, L_OEQ, L_PPEG, L_AP]:
    assert any(l == n for (l, e) in BSA), n
for n in [L_DREE, L_CFO, L_REEX, L_REPR, L_ACQ]:
    assert any(l == n for (l, e) in CFA), n
assert "mixed basis" in mixcol
assert "SCHEDULE II" in T_AR12[T_AR12.rfind("[p. 140]"):T_AR12.rfind("[p. 141]")] and PAGES["sched12"] == "140"

if __name__ == "__main__":
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(reply)
    sys.stdout.reconfigure(encoding="utf-8")
    print(reply)
    body = "\n".join(l for l in reply.splitlines() if not l.startswith("|---"))
    print("WORDS (incl. table cells):", len(re.findall(r"[^\s|*]+", body)))
    prose = "\n".join(l for l in reply.splitlines() if not l.startswith("|"))
    print("WORDS (excl. table):", len(re.findall(r"[^\s*]+", prose)))
