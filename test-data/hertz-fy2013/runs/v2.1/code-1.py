# Turn 1: scan the numbers. Run from anywhere; paths are absolute.
import re, sys, io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
D = r"C:\Users\aaron\OneDrive\Desktop\Master\WireSift Research\Events\CFA Vancouver Forensic\test-data\hertz-fy2013"
wb = openpyxl.load_workbook(D + r"\statements.xlsx", data_only=True)


def sheet(name):
    """Return {line label: {period label: value}}; duplicate labels get a numeric suffix."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    out, seen = {}, {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        lab = r[0]
        seen[lab] = seen.get(lab, 0) + 1
        if seen[lab] > 1:
            lab = f"{lab} #{seen[lab]}"
        out[lab] = {hdr[i]: r[i] for i in range(1, len(hdr) - 1)}
    return out


ISa, ISq, BSa, BSq, CFa, CFq = (sheet(n) for n in
    ["IS annual", "IS quarterly", "BS annual", "BS quarterly", "CF annual", "CF quarterly"])


def fy(y): return f"FY (ended {y}-12-31)"
def q4(y): return f"Q4 (ended {y}-12-31)"
def q3(y): return f"Q3 (ended {y}-09-30)"


REV, DOP, DEPL, SGA = ("Total revenues", "Direct operating",
    "Depreciation of revenue earning equipment and lease charges", "Selling, general and administrative")
REC = "Receivables, less allowance for doubtful accounts of $30.8 and $29.3"
INV, AP, TA = "Inventories, at lower of cost or market", "Accounts payable", "Total assets"
CARS, OEQ, PPEG = "Cars", "Other equipment", "Total property and equipment, at cost"
REEN, PPEN = "Total revenue earning equipment", "Total property and equipment"
OCF = "Net cash provided by (used in) operating activities"
DREE, DPPE = "Depreciation of revenue earning equipment", "Depreciation of property and equipment"
BUY, SELL = "Revenue earning equipment expenditures", "Proceeds from disposal of revenue earning equipment"
PBUY, PSELL = "Property and equipment expenditures", "Proceeds from disposal of property and equipment"
ACQ = "Acquisitions, net of cash acquired"

# Allowance for doubtful accounts, USD millions: balance-sheet captions and Schedule II of the two annual reports.
# FY2012 report p. 72 and p. 140: 25.113 (2012), 20.282 (2011), 19.708 (2010). FY2013 report p. 75: 30.8 (2013), 29.3 (2012 as revised).
ALLOW = {2010: 19.708, 2011: 20.282, 2012: 25.113, 2013: 30.8}
ALLOW_2012_REVISED = 29.3
# check these against the text
t12 = open(D + r"\annual-report-FY2012-12-31-key-sections.txt", encoding="utf-8").read()
t13 = open(D + r"\annual-report-FY2013-12-31-key-sections.txt", encoding="utf-8").read()
assert "allowance for doubtful accounts of $25,113 and $20,282" in t12
assert "Year ended December 31, 2010 | 21,268 | 19,667 | (695) | (20,532) | (a) | 19,708" in t12
assert "allowance for doubtful accounts of $30.8 and $29.3" in t13

days = {2009: 365, 2010: 365, 2011: 365, 2012: 366, 2013: 365}
cos = lambda S, p: S[DOP][p] + S[DEPL][p]          # cost-of-sales substitute
gm = lambda S, p: (S[REV][p] - cos(S, p)) / S[REV][p]
om = lambda S, p: (S[REV][p] - cos(S, p) - S[SGA][p]) / S[REV][p]
g = lambda a, b: a / b - 1

T = {}  # measure -> {column: value}
def put(m, c, v): T.setdefault(m, {})[c] = v

for y in range(2009, 2014):
    c = f"FY{y}"; p = fy(y)
    put("Gross margin %", c, 100 * gm(ISa, p))
    put("Operating margin %", c, 100 * om(ISa, p))
    put("Direct operating % revenue", c, 100 * ISa[DOP][p] / ISa[REV][p])
    put("Fleet depreciation and lease charges % revenue", c, 100 * ISa[DEPL][p] / ISa[REV][p])
    put("SG&A % revenue", c, 100 * ISa[SGA][p] / ISa[REV][p])
    netfleet = -CFa[BUY][p] - CFa[SELL][p]
    netppe = -CFa[PBUY][p] - CFa[PSELL][p]
    put("Net fleet spending", c, netfleet)
    put("Free cash flow", c, CFa[OCF][p] - netfleet - netppe)
    put("Net income", c, CFa["Net income"][p])
    put("Operating cash flow", c, CFa[OCF][p])
    put("Acquisitions", c, CFa[ACQ][p])
    if y >= 2010:
        put("DSO", c, BSa[REC][p] / ISa[REV][p] * days[y])
        put("Days inventory", c, BSa[INV][p] / cos(ISa, p) * days[y])
        put("Days payable", c, BSa[AP][p] / cos(ISa, p) * days[y])
        put("Allowance % gross receivables", c, 100 * ALLOW[y] / (BSa[REC][p] + ALLOW[y]))
    if y >= 2011:
        pp = fy(y - 1)
        rg, cg = g(ISa[REV][p], ISa[REV][pp]), g(cos(ISa, p), cos(ISa, pp))
        put("Revenue growth %", c, 100 * rg)
        put("Receivables growth - revenue growth, pts", c, 100 * (g(BSa[REC][p], BSa[REC][pp]) - rg))
        put("Inventory growth - cost growth, pts", c, 100 * (g(BSa[INV][p], BSa[INV][pp]) - cg))
        put("Payables growth - cost growth, pts", c, 100 * (g(BSa[AP][p], BSa[AP][pp]) - cg))
        avgTA = (BSa[TA][p] + BSa[TA][pp]) / 2
        put("Accruals % avg assets", c, 100 * (CFa["Net income"][p] - CFa[OCF][p]) / avgTA)
        gross_fleet = lambda q: BSa[CARS][q] + BSa[OEQ][q]
        put("Fleet depreciation rate %", c, 100 * CFa[DREE][p] / ((gross_fleet(p) + gross_fleet(pp)) / 2))
        put("Cars: gross", c, BSa[CARS][p])
        put("Property depreciation rate %", c, 100 * CFa[DPPE][p] / ((BSa[PPEG][p] + BSa[PPEG][pp]) / 2))
        # Beneish
        def ca(q): return sum(BSa[k][q] for k in ["Cash and cash equivalents", "Restricted cash and cash equivalents", REC, INV, "Prepaid expenses and other assets"])
        def ppe(q): return BSa[REEN][q] + BSa[PPEN][q]
        def dep(q): return CFa[DREE][q] + CFa[DPPE][q]
        def lev(q): return (BSa["Total liabilities"][q] - BSa["Public liability and property damage"][q] -
                            (BSa["Deferred taxes on income"][q] if BSa["Deferred taxes on income"][q] is not None else BSa["Deferred Tax Liabilities, Net"][q])) / BSa[TA][q]
        DSRI = (BSa[REC][p] / ISa[REV][p]) / (BSa[REC][pp] / ISa[REV][pp])
        GMI = gm(ISa, pp) / gm(ISa, p)
        AQI = (1 - (ca(p) + ppe(p)) / BSa[TA][p]) / (1 - (ca(pp) + ppe(pp)) / BSa[TA][pp])
        SGI = ISa[REV][p] / ISa[REV][pp]
        DEPI = (dep(pp) / (dep(pp) + ppe(pp))) / (dep(p) / (dep(p) + ppe(p)))
        SGAI = (ISa[SGA][p] / ISa[REV][p]) / (ISa[SGA][pp] / ISa[REV][pp])
        TATA = (CFa["Net income"][p] - CFa[OCF][p]) / BSa[TA][p]
        LVGI = lev(p) / lev(pp)
        M = (-4.84 + 0.920 * DSRI + 0.528 * GMI + 0.404 * AQI + 0.892 * SGI + 0.115 * DEPI
             - 0.172 * SGAI + 4.679 * TATA - 0.327 * LVGI)
        put("Beneish M-score", c, M)
        print(f"Beneish {y}: DSRI {DSRI:.3f} GMI {GMI:.3f} AQI {AQI:.3f} SGI {SGI:.3f} DEPI {DEPI:.3f} SGAI {SGAI:.3f} TATA {TATA:.3f} LVGI {LVGI:.3f} M {M:.2f}")

# Latest quarter beside the same quarter a year earlier (92 days). Year-over-year growth uses the same quarter.
for y in (2012, 2013):
    c = f"Q4 {y}"; p = q4(y); pp = q4(y - 1)
    put("Gross margin %", c, 100 * gm(ISq, p))
    put("Operating margin %", c, 100 * om(ISq, p))
    put("Direct operating % revenue", c, 100 * ISq[DOP][p] / ISq[REV][p])
    put("Fleet depreciation and lease charges % revenue", c, 100 * ISq[DEPL][p] / ISq[REV][p])
    put("SG&A % revenue", c, 100 * ISq[SGA][p] / ISq[REV][p])
    netfleet = -CFq[BUY][p] - CFq[SELL][p]
    netppe = -CFq[PBUY][p] - CFq[PSELL][p]
    put("Net fleet spending", c, netfleet)
    put("Free cash flow", c, CFq[OCF][p] - netfleet - netppe)
    put("Net income", c, CFq["Net income"][p])
    put("Operating cash flow", c, CFq[OCF][p])
    put("Acquisitions", c, CFq[ACQ][p])
    put("DSO", c, BSq[REC][p] / ISq[REV][p] * 92)
    put("Days inventory", c, BSq[INV][p] / cos(ISq, p) * 92)
    put("Days payable", c, BSq[AP][p] / cos(ISq, p) * 92)
    put("Allowance % gross receivables", c, 100 * ALLOW[y] / (BSq[REC][p] + ALLOW[y]))
    rg, cg = g(ISq[REV][p], ISq[REV][pp]), g(cos(ISq, p), cos(ISq, pp))
    put("Revenue growth %", c, 100 * rg)
    put("Receivables growth - revenue growth, pts", c, 100 * (g(BSq[REC][p], BSq[REC][pp]) - rg))
    put("Inventory growth - cost growth, pts", c, 100 * (g(BSq[INV][p], BSq[INV][pp]) - cg))
    put("Payables growth - cost growth, pts", c, 100 * (g(BSq[AP][p], BSq[AP][pp]) - cg))
    pq = q3(y)
    put("Accruals % avg assets", c, 100 * (CFq["Net income"][p] - CFq[OCF][p]) / ((BSq[TA][p] + BSq[TA][pq]) / 2))
    gf = lambda q: BSq[CARS][q] + BSq[OEQ][q]
    put("Fleet depreciation rate %", c, 100 * 4 * CFq[DREE][p] / ((gf(p) + gf(pq)) / 2))
    put("Property depreciation rate %", c, 100 * 4 * CFq[DPPE][p] / ((BSq[PPEG][p] + BSq[PPEG][pq]) / 2))
    put("Fleet purchases", c, -CFq[BUY][p])

cols = [f"FY{y}" for y in range(2009, 2014)] + ["Q4 2012", "Q4 2013"]
print("\n| Measure | " + " | ".join(cols) + " |")
for m, v in T.items():
    print(f"| {m} | " + " | ".join("" if c not in v else f"{v[c]:,.1f}" if abs(v[c]) < 1000 or True else "" for c in cols) + " |")
print()
for y in range(2010, 2014):
    print("M", y, round(T["Beneish M-score"].get(f"FY{y}", float("nan")), 2))

# Supporting figures for the flags
print("\n-- supporting --")
print("Receivables", [BSa[REC][fy(y)] for y in range(2010, 2014)], "change 2013 %", 100 * g(BSa[REC][fy(2013)], BSa[REC][fy(2012)]))
print("Fleet sales in receivables", [CFa["Sales of revenue earning equipment included in receivables"][fy(y)] for y in range(2009, 2014)])
print("Receivables less fleet sales in receivables", [round(BSa[REC][fy(y)] - CFa["Sales of revenue earning equipment included in receivables"][fy(y)], 1) for y in range(2010, 2014)])
print("Prepaid and other", [BSa["Prepaid expenses and other assets"][fy(y)] for y in range(2010, 2014)], "growth 2013 %", 100 * g(BSa["Prepaid expenses and other assets"][fy(2013)], BSa["Prepaid expenses and other assets"][fy(2012)]))
print("Restricted cash", [BSa["Restricted cash and cash equivalents"][fy(y)] for y in range(2010, 2014)])
print("AP", [BSa[AP][fy(y)] for y in range(2010, 2014)], "growth 2013 %", 100 * g(BSa[AP][fy(2013)], BSa[AP][fy(2012)]))
print("Fleet purchases Q4", CFq[BUY][q4(2011)], CFq[BUY][q4(2012)], CFq[BUY][q4(2013)], "chg %", 100 * g(CFq[BUY][q4(2013)], CFq[BUY][q4(2012)]))
print("Fleet purchases FY", [CFa[BUY][fy(y)] for y in range(2009, 2014)])
print("Cars gross Q3 2013 -> Q4 2013", BSq[CARS][q3(2013)], BSq[CARS][q4(2013)])
print("Dep REE FY", [CFa[DREE][fy(y)] for y in range(2009, 2014)], "growth 2013 %", 100 * g(CFa[DREE][fy(2013)], CFa[DREE][fy(2012)]))
gfa = lambda y: BSa[CARS][fy(y)] + BSa[OEQ][fy(y)]
print("Gross fleet FY", [round(gfa(y), 1) for y in range(2010, 2014)], "growth 2013 %", 100 * g(gfa(2013), gfa(2012)), "2012 %", 100 * g(gfa(2012), gfa(2011)))
print("Dep REE Q4", CFq[DREE][q4(2011)], CFq[DREE][q4(2012)], CFq[DREE][q4(2013)], "growth %", 100 * g(CFq[DREE][q4(2013)], CFq[DREE][q4(2012)]))
print("Q4 revenue growth %", 100 * g(ISq[REV][q4(2013)], ISq[REV][q4(2012)]), "Q4 direct operating growth %", 100 * g(ISq[DOP][q4(2013)], ISq[DOP][q4(2012)]))
print("Provision for doubtful accounts FY", [CFa["Provision for losses on doubtful accounts"][fy(y)] for y in range(2009, 2014)])
print("Allowance 2012 revised %", 100 * ALLOW_2012_REVISED / (1879.7 + ALLOW_2012_REVISED))
print("OCF growth 2013 %", 100 * g(CFa[OCF][fy(2013)], CFa[OCF][fy(2012)]), "NI growth %", 100 * g(346.2, 243.1))
print("Deferred taxes on income (CF)", [CFa["Deferred taxes on income"][fy(y)] for y in range(2009, 2014)])
print("Goodwill+intangibles", [(BSa["Goodwill"][fy(y)], BSa["Other intangible assets, net"][fy(y)]) for y in range(2010, 2014)])
print("Accrued liabilities Q", {k: v for k, v in BSq["Accrued liabilities"].items()})

# Q4 2013 on the revised nine months (FY2013 annual report, Note 2, p. 87): full year less revised nine months
rev9 = {"rev": 8215.6, "dop": 4312.7, "depl": 1882.5, "sga": 800.0, "pti": 600.8, "ni": 346.8}
assert "Total revenues | $3,069.4 | $— | $3,069.4 | $8,220.6 | $(5.0) | $8,215.6" in t13
assert "Direct operating | 1,525.4 | 14.1 | 1,539.5 | 4,282.6 | (g) | 30.1 | 4,312.7" in t13
assert "Depreciation of revenue earning equipment and lease charges | 676.7 | (1.8) | 674.9 | 1,904.8 | (h) | (22.3) | 1,882.5" in t13
assert "Selling, general and administrative | 276.8 | (5.3) | 271.5 | 803.5 | (i) | (3.5) | 800.0" in t13
p = fy(2013)
r4 = ISa[REV][p] - rev9["rev"]; d4 = ISa[DOP][p] - rev9["dop"]; l4 = ISa[DEPL][p] - rev9["depl"]; s4 = ISa[SGA][p] - rev9["sga"]
print(f"\nQ4 2013 on revised nine months: revenue {r4:.1f}, direct operating {d4:.1f} ({100*d4/r4:.1f}% of revenue), fleet dep {l4:.1f}, SG&A {s4:.1f}, "
      f"gross margin {100*(r4-d4-l4)/r4:.1f}%, operating margin {100*(r4-d4-l4-s4)/r4:.1f}%, pre-tax {ISa['Income before income taxes'][p]-rev9['pti']:.1f}, net income {ISa['Net income'][p]-rev9['ni']:.1f}")
# Q4 2012 as revised (p. 89)
assert "Total revenues | $2,318.5 | $1.2 | $2,319.7" in t13
print(f"Q4 2012 as revised: direct operating % revenue {100*1257.7/2319.7:.1f}, gross margin {100*(2319.7-1257.7-550.0)/2319.7:.1f}%, operating margin {100*(2319.7-1257.7-550.0-331.1)/2319.7:.1f}%")

# Revisions that moved income or equity
ws = wb["Revisions"]
rows = list(ws.iter_rows(values_only=True))
print("\n-- revisions touching income or equity --")
for r in rows[1:]:
    if r[1] and any(k in r[1] for k in ["Net income", "Income before", "Total equity", "Accumulated deficit", "Total revenues", "stockholders"]) and r[3] and ("12 months" in r[3] or "as at" in r[3]):
        print(r[0], "|", r[1][:60], "|", r[3], "|", r[4], "->", r[6], "|", r[8], "|", r[7][:40])
print("rows on Revisions sheet:", len(rows) - 1)

# ---- reply table ----
def f(m, c, nd=1, plus=False):
    v = T.get(m, {}).get(c)
    if v is None: return ""
    s = f"{v:,.{nd}f}"
    return ("+" + s) if plus and v > 0 else s

layout = [
    ("Days sales outstanding", "DSO", 0, False),
    ("Days inventory", "Days inventory", 1, False),
    ("Days payable", "Days payable", 0, False),
    ("Receivables growth less revenue growth (pts)", "Receivables growth - revenue growth, pts", 1, True),
    ("Inventory growth less cost growth (pts)", "Inventory growth - cost growth, pts", 1, True),
    ("Payables growth less cost growth (pts)", "Payables growth - cost growth, pts", 1, True),
    ("Allowance, % of gross receivables", "Allowance % gross receivables", 1, False),
    ("Gross margin, %", "Gross margin %", 1, False),
    ("Operating margin, %", "Operating margin %", 1, False),
    ("Accruals, % of average assets", "Accruals % avg assets", 1, False),
    ("Free cash flow, $m", "Free cash flow", 1, False),
    ("Net income, $m", "Net income", 1, False),
    ("Depreciation rate, rental equipment, %", "Fleet depreciation rate %", 1, False),
    ("Depreciation rate, other property, %", "Property depreciation rate %", 1, False),
    ("Beneish M-score", "Beneish M-score", 2, False),
]
print("\n| Measure | " + " | ".join(cols) + " |")
print("|---|" + "---:|" * len(cols))
for lab, key, nd, plus in layout:
    print(f"| {lab} | " + " | ".join(f(key, c, nd, plus) for c in cols) + " |")

# extra supporting numbers used in the flags
gfq = lambda q: BSq[CARS][q] + BSq[OEQ][q]
a13 = (gfq(q4(2013)) + gfq(q3(2013))) / 2; a12 = (gfq(q4(2012)) + gfq(q3(2012))) / 2
print(f"\nAverage gross rental equipment Q4 2013 {a13:,.1f} vs Q4 2012 {a12:,.1f}: growth {100*g(a13,a12):.1f}%")
print(f"Receivables decline 2013: {BSa[REC][fy(2012)]-BSa[REC][fy(2013)]:.1f}; decline in fleet sales included in receivables: {617.3-357.3:.1f}")
print(f"Cost growth 2013 {100*g(cos(ISa,fy(2013)),cos(ISa,fy(2012))):.1f}%")
print(f"Direct operating Q4 2013 {ISq[DOP][q4(2013)]} / rev {ISq[REV][q4(2013)]}; Q4 2012 {ISq[DOP][q4(2012)]} / {ISq[REV][q4(2012)]}")
print(f"Nine-month 2013 direct operating: first reported 4282.6, revised 4312.7, difference {4312.7-4282.6:.1f}")
print(f"Acquisition outflow FY2012 {CFa[ACQ][fy(2012)]}, Q4 2012 {CFq[ACQ][q4(2012)]}")
