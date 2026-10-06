# Scan the numbers: Hertz Global Holdings, fiscal years 2009-2013 and Q4 2013 vs Q4 2012.
# Inputs (attached files only): statements.xlsx, the two annual reports (key sections),
# the two fourth-quarter results releases. USD millions.
import io
import os
import re
import sys

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
P = lambda name: os.path.join(BASE, name)

wb = openpyxl.load_workbook(P("statements.xlsx"), data_only=True)


def sheet(name):
    """Return {xbrl_element: {column_heading: value}} and the list of headings."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    out = {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        key = r[-1]
        if key in out:  # duplicate element (cash at beginning / end of period): keep label too
            key = f"{r[-1]}|{r[0]}"
        out[key] = {head[i]: r[i] for i in range(1, len(head) - 1)}
        out[key]["_label"] = r[0]
    return out, head[1:-1]


ISA, isa_cols = sheet("IS annual")
BSA, bsa_cols = sheet("BS annual")
CFA, cfa_cols = sheet("CF annual")
ISQ, isq_cols = sheet("IS quarterly")
BSQ, bsq_cols = sheet("BS quarterly")
CFQ, cfq_cols = sheet("CF quarterly")

print("About sheet:")
for r in wb["About"].iter_rows(values_only=True):
    print("  ", r[0])
print("IS quarterly headings:", isq_cols)
print("CF quarterly headings:", cfq_cols)
mixed = [c for c in isq_cols + cfq_cols if "mixed basis" in c]
print("Columns marked mixed basis:", mixed)

FY = lambda y: f"FY (ended {y}-12-31)"
QE = {"Q3 2012": "Q3 (ended 2012-09-30)", "Q4 2012": "Q4 (ended 2012-12-31)",
      "Q3 2013": "Q3 (ended 2013-09-30)", "Q4 2013": "Q4 (ended 2013-12-31)"}

# XBRL elements used
E = dict(
    rev="us-gaap:Revenues",
    dirop="us-gaap:DirectOperatingCosts",
    depree_is="htz:CostOfServicesDepreciationAndLeaseCharges",
    sga="us-gaap:SellingGeneralAndAdministrativeExpense",
    ni="us-gaap:ProfitLoss",
    rec="us-gaap:AccountsReceivableNet",
    inv="us-gaap:InventoryNet",
    prepaid="us-gaap:PrepaidExpenseAndOtherAssets",
    cash="us-gaap:CashAndCashEquivalentsAtCarryingValue",
    rcash="us-gaap:RestrictedCashAndCashEquivalents",
    cars="us-gaap:PropertySubjectToOrAvailableForOperatingLeaseGross",
    cars_ad="us-gaap:PropertySubjectToOrAvailableForOperatingLeaseAccumulatedDepreciation",
    oeq="htz:OtherEquipmentGross",
    ree_net="us-gaap:PropertySubjectToOrAvailableForOperatingLeaseNet",
    ppe_gross="us-gaap:PropertyPlantAndEquipmentGross",
    ppe_net="us-gaap:PropertyPlantAndEquipmentNet",
    ta="us-gaap:Assets",
    ap="us-gaap:AccountsPayableCurrentAndNoncurrent",
    tl="us-gaap:Liabilities",
    ocf="us-gaap:NetCashProvidedByUsedInOperatingActivitiesContinuingOperations",
    dep_ree="htz:CostOfServicesDepreciationAsAdjustedForGainLossUponDisposal",
    dep_ppe="us-gaap:DepreciationNonproduction",
    ree_exp="htz:PaymentsToAcquireRevenueEarningEquipment",
    ree_proc="htz:ProceedsFromSaleOfRevenueEarningEquipment",
    ppe_exp="us-gaap:PaymentsToAcquireOtherPropertyPlantAndEquipment",
    ppe_proc="us-gaap:ProceedsFromSaleOfOtherPropertyPlantAndEquipment",
    acq="us-gaap:PaymentsToAcquireBusinessesNetOfCashAcquired",
    prov="us-gaap:ProvisionForDoubtfulAccounts",
    veh_rec="htz:SalesOfRevenueEarningEquipmentEarnedButNotReceived",
)


def a(sh, key, y):
    return sh[E[key]].get(FY(y))


def bq(key, q):
    return BSQ[E[key]].get(QE[q])


# ---------------------------------------------------------------- results release (Q4)
rel13 = open(P("earnings-release-Q4-FY2013-12-31.txt"), encoding="utf-8").read()
rel12 = open(P("earnings-release-Q4-FY2012-12-31.txt"), encoding="utf-8").read()


def num(s):
    s = s.strip().replace("$", "").replace(",", "")
    neg = s.startswith("(")
    s = s.strip("()")
    if s in ("—", "-", ""):
        return 0.0
    return -float(s) if neg else float(s)


def line_after(text, anchor, label, occurrence=1):
    """Values on the n-th line that starts with `label`, searching from `anchor`."""
    i = text.find(anchor)
    assert i >= 0, anchor
    pos = i
    for _ in range(occurrence):
        m = re.compile(r"^" + re.escape(label) + r" \| (.*)$", re.M).search(text, pos)
        assert m, label
        pos = m.end()
    return [num(x) for x in m.group(1).split(" | ") if re.search(r"\d|—", x) and "%" not in x], m.group(0)


T1 = "Table 1\n\nHERTZ GLOBAL HOLDINGS, INC.\n\nCONDENSED CONSOLIDATED STATEMENTS OF OPERATIONS"
assert T1 in rel13
R = {}
src_lines = {}
for key, label in [("rev", "Total revenues"), ("dirop", "Direct operating"),
                   ("depree_is", "Depreciation of revenue earning equipment and lease charges"),
                   ("sga", "Selling, general and administrative"),
                   ("pretax", "Income (loss) before income taxes"),
                   ("ni", "Net loss attributable to Hertz Global Holdings, Inc. and Subsidiaries’ common stockholders")]:
    v, ln = line_after(rel13, T1, label)
    R[key] = {"Q4 2013": v[0], "Q4 2012": v[1]}
    src_lines[key] = ln
T8 = "RECONCILIATION FROM OPERATING"
v, ln = line_after(rel13, T8, "Net cash provided by operating activities")
R["ocf"] = {"Q4 2013": v[0], "Q4 2012": v[1]}; src_lines["ocf"] = ln
T7 = "FLEET GROWTH | Rental"
v, ln = line_after(rel13, T7, "Revenue earning equipment expenditures")
R["ree_exp"] = {"Q4 2013": -v[4], "Q4 2012": -v[9]}; src_lines["ree_exp"] = ln
v, ln = line_after(rel13, T7, "Proceeds from disposal of revenue earning equipment")
R["ree_proc"] = {"Q4 2013": v[4], "Q4 2012": v[9]}; src_lines["ree_proc"] = ln
v, ln = line_after(rel13, T7, "Depreciation of revenue earning equipment")
R["dep_ree"] = {"Q4 2013": v[4], "Q4 2012": v[9]}; src_lines["dep_ree"] = ln
v, ln = line_after(rel13, "FREE CASH FLOW (1)", "Property and equipment expenditures, net of disposals")
R["ppe_net_capex"] = {"Q4 2013": -v[0], "Q4 2012": -v[1]}; src_lines["ppe_net_capex"] = ln
i = rel13.find("Depreciation of property and equipment:\nU.S. Car Rental")
m = re.compile(r"^\$([\d.,]+) \| \$([\d.,]+) \| \$([\d.,]+) \| \$([\d.,]+)$", re.M).search(rel13, i)
R["dep_ppe"] = {"Q4 2013": num(m.group(1)), "Q4 2012": num(m.group(2))}; src_lines["dep_ppe"] = m.group(0)

print("\nRelease (filed 2014-03-18) fourth-quarter figures used:")
for k, v in R.items():
    print(f"  {k:14s} Q4 2013 {v['Q4 2013']:>9.1f}   Q4 2012 {v['Q4 2012']:>9.1f}   <- {src_lines[k][:110]}")

print("\nWorkbook derived Q4 2013 [mixed basis] vs release Q4 2013:")
for key, sh in [("rev", ISQ), ("dirop", ISQ), ("depree_is", ISQ), ("sga", ISQ), ("ni", ISQ), ("ocf", CFQ)]:
    w = sh[E[key]][[c for c in (isq_cols if sh is ISQ else cfq_cols) if c.startswith("Q4 (ended 2013-12-31)")][0]]
    print(f"  {key:10s} workbook {w:>8.1f}  release {R[key]['Q4 2013']:>8.1f}  diff {R[key]['Q4 2013'] - w:>+7.1f}")
w12 = {k: (ISQ if k != "ocf" else CFQ)[E[k]]["Q4 (ended 2012-12-31)"] for k in ["rev", "dirop", "depree_is", "sga", "ni", "ocf"]}
print("Workbook derived Q4 2012 (first reported) vs FY2013 release Q4 2012 (revised):")
for k, w in w12.items():
    print(f"  {k:10s} workbook {w:>8.1f}  release {R[k]['Q4 2012']:>8.1f}  diff {R[k]['Q4 2012'] - w:>+7.1f}")

# quotations used: verify exact substrings and the page they sit on
ar12 = open(P("annual-report-FY2012-12-31-key-sections.txt"), encoding="utf-8").read()
ar13 = open(P("annual-report-FY2013-12-31-key-sections.txt"), encoding="utf-8").read()


def page_of(text, snippet):
    i = text.find(snippet)
    assert i >= 0, snippet
    return re.findall(r"\[p\. ([^\]]+)\]", text[:i])[-1]


# ---------------------------------------------------------------- allowance (annual reports)
ALLOW_SRC = [
    (2013, ar13, "Receivables, less allowance for doubtful accounts of $30.8 and $29.3 | 1,512.6 | 1,879.7"),
    (2012, ar12, "Receivables, less allowance for doubtful accounts of $25,113 and $20,282 | 1,886,596 | 1,616,382"),
    (2010, ar12, "Year ended December 31, 2011 | 19,708 | 28,164 | 68 | (27,658) | (a) | 20,282"),
]
for y, txt, snip in ALLOW_SRC:
    print(f"allowance source FY{y}: p. {page_of(txt, snip)}: {snip}")
allow = {2010: 19.708, 2011: 20.282, 2012: 25.113, 2013: 30.8}   # 2012 as first reported
allow_2012_revised = 29.3
rec_2012_revised = 1879.7

# ---------------------------------------------------------------- measures
YEARS = [2009, 2010, 2011, 2012, 2013]
DAYS = {2009: 365, 2010: 365, 2011: 365, 2012: 366, 2013: 365, "Q4 2012": 92, "Q4 2013": 92}
M = {}  # measure -> period -> value


def put(name, period, val):
    M.setdefault(name, {})[period] = val


def growth(new, old):
    return new / old - 1


for y in YEARS:
    rev = a(ISA, "rev", y)
    cost = a(ISA, "dirop", y) + a(ISA, "depree_is", y)
    sga = a(ISA, "sga", y)
    ni = a(ISA, "ni", y)
    ocf = a(CFA, "ocf", y)
    put("Gross margin", y, (rev - cost) / rev)
    put("Operating margin", y, (rev - cost - sga) / rev)
    fleet_net = -a(CFA, "ree_exp", y) - a(CFA, "ree_proc", y)
    ppe_net = -a(CFA, "ppe_exp", y) - a(CFA, "ppe_proc", y)
    put("Net fleet capex", y, fleet_net)
    put("Net P&E capex", y, ppe_net)
    put("Free cash flow", y, ocf - fleet_net - ppe_net)
    put("Net income", y, ni)
    put("OCF", y, ocf)
    put("Acquisitions", y, -a(CFA, "acq", y))
    if y >= 2010:
        rec, inv, ap = a(BSA, "rec", y), a(BSA, "inv", y), a(BSA, "ap", y)
        put("DSO", y, rec / rev * DAYS[y])
        put("Days inventory", y, inv / cost * DAYS[y])
        put("Days payable", y, ap / cost * DAYS[y])
        put("Allowance %", y, allow[y] / (rec + allow[y]))
    if y >= 2011:
        p = y - 1
        rev0 = a(ISA, "rev", p)
        cost0 = a(ISA, "dirop", p) + a(ISA, "depree_is", p)
        put("Rec gap", y, growth(a(BSA, "rec", y), a(BSA, "rec", p)) - growth(rev, rev0))
        put("Inv gap", y, growth(a(BSA, "inv", y), a(BSA, "inv", p)) - growth(cost, cost0))
        put("AP gap", y, growth(a(BSA, "ap", y), a(BSA, "ap", p)) - growth(cost, cost0))
        put("Rev growth", y, growth(rev, rev0))
        put("Cost growth", y, growth(cost, cost0))
        put("Rec growth", y, growth(a(BSA, "rec", y), a(BSA, "rec", p)))
        avg_ta = (a(BSA, "ta", y) + a(BSA, "ta", p)) / 2
        put("Accruals", y, (ni - ocf) / avg_ta)
        g_ree = lambda yy: a(BSA, "cars", yy) + a(BSA, "oeq", yy)
        put("Dep rate fleet", y, a(CFA, "dep_ree", y) / ((g_ree(y) + g_ree(p)) / 2))
        put("Dep rate P&E", y, a(CFA, "dep_ppe", y) / ((a(BSA, "ppe_gross", y) + a(BSA, "ppe_gross", p)) / 2))

# Beneish M-score per pair of years (workbook annual sheets, as first reported)


def beneish(y):
    p = y - 1
    d = {}
    S, S0 = a(ISA, "rev", y), a(ISA, "rev", p)
    cost = lambda yy: a(ISA, "dirop", yy) + a(ISA, "depree_is", yy)
    gm = lambda yy: (a(ISA, "rev", yy) - cost(yy)) / a(ISA, "rev", yy)
    d["DSRI"] = (a(BSA, "rec", y) / S) / (a(BSA, "rec", p) / S0)
    d["GMI"] = gm(p) / gm(y)
    # unclassified balance sheet: current assets = cash + restricted cash + receivables + inventories + prepaid/other
    ca = lambda yy: sum(a(BSA, k, yy) for k in ["cash", "rcash", "rec", "inv", "prepaid"])
    ppe = lambda yy: a(BSA, "ree_net", yy) + a(BSA, "ppe_net", yy)
    aq = lambda yy: 1 - (ca(yy) + ppe(yy)) / a(BSA, "ta", yy)
    d["AQI"] = aq(y) / aq(p)
    d["SGI"] = S / S0
    dep = lambda yy: a(CFA, "dep_ree", yy) + a(CFA, "dep_ppe", yy)
    dr = lambda yy: dep(yy) / (dep(yy) + ppe(yy))
    d["DEPI"] = dr(p) / dr(y)
    d["SGAI"] = (a(ISA, "sga", y) / S) / (a(ISA, "sga", p) / S0)
    d["TATA"] = (a(ISA, "ni", y) - a(CFA, "ocf", y)) / a(BSA, "ta", y)
    lv = lambda yy: a(BSA, "tl", yy) / a(BSA, "ta", yy)   # no current/long-term split: total liabilities
    d["LVGI"] = lv(y) / lv(p)
    d["M"] = (-4.84 + 0.920 * d["DSRI"] + 0.528 * d["GMI"] + 0.404 * d["AQI"] + 0.892 * d["SGI"]
              + 0.115 * d["DEPI"] - 0.172 * d["SGAI"] + 4.679 * d["TATA"] - 0.327 * d["LVGI"])
    return d


BEN = {y: beneish(y) for y in [2011, 2012, 2013]}
for y in BEN:
    put("M-score", y, BEN[y]["M"])

# Fourth quarter: flows from the FY2013 results release; balances from the BS quarterly sheet
for q, q_prev_end, yq in [("Q4 2012", "Q3 2012", 2012), ("Q4 2013", "Q3 2013", 2013)]:
    rev = R["rev"][q]
    cost = R["dirop"][q] + R["depree_is"][q]
    put("Gross margin", q, (rev - cost) / rev)
    put("Operating margin", q, (rev - cost - R["sga"][q]) / rev)
    rec, inv, ap = bq("rec", q), bq("inv", q), bq("ap", q)
    put("DSO", q, rec / rev * 92)
    put("Days inventory", q, inv / cost * 92)
    put("Days payable", q, ap / cost * 92)
    put("Allowance %", q, allow[yq] / (rec + allow[yq]))
    avg_ta = (bq("ta", q) + bq("ta", q_prev_end)) / 2
    put("Accruals", q, (R["ni"][q] - R["ocf"][q]) / avg_ta)
    fleet_net = R["ree_exp"][q] - R["ree_proc"][q]
    put("Net fleet capex", q, fleet_net)
    put("Net P&E capex", q, R["ppe_net_capex"][q])
    put("Free cash flow", q, R["ocf"][q] - fleet_net - R["ppe_net_capex"][q])
    put("Net income", q, R["ni"][q])
    put("OCF", q, R["ocf"][q])
    g = lambda qq: bq("cars", qq) + bq("oeq", qq)
    put("Dep rate fleet", q, R["dep_ree"][q] * 4 / ((g(q) + g(q_prev_end)) / 2))
    put("Dep rate P&E", q, R["dep_ppe"][q] * 4 / ((bq("ppe_gross", q) + bq("ppe_gross", q_prev_end)) / 2))
q, q0 = "Q4 2013", "Q4 2012"
rg = growth(R["rev"][q], R["rev"][q0])
cg = growth(R["dirop"][q] + R["depree_is"][q], R["dirop"][q0] + R["depree_is"][q0])
put("Rev growth", q, rg); put("Cost growth", q, cg)
put("Rec gap", q, growth(bq("rec", q), bq("rec", q0)) - rg)
put("Inv gap", q, growth(bq("inv", q), bq("inv", q0)) - cg)
put("AP gap", q, growth(bq("ap", q), bq("ap", q0)) - cg)

PERIODS = [2009, 2010, 2011, 2012, 2013, "Q4 2012", "Q4 2013"]
print("\n================ MEASURES ================")
for name, vals in M.items():
    cells = []
    for p in PERIODS:
        v = vals.get(p)
        if v is None:
            cells.append(f"{'':>10s}")
        elif name in ("DSO", "Days inventory", "Days payable"):
            cells.append(f"{v:>10.1f}")
        elif name in ("Net fleet capex", "Net P&E capex", "Free cash flow", "Net income", "OCF", "Acquisitions"):
            cells.append(f"{v:>10.1f}")
        elif name == "M-score":
            cells.append(f"{v:>10.2f}")
        else:
            cells.append(f"{v * 100:>9.2f}%")
    print(f"{name:18s}" + "".join(cells))
print("periods:           " + "".join(f"{str(p):>10s}" for p in PERIODS))
print("\nBeneish inputs:")
for y, d in BEN.items():
    print(f"  {y}/{y - 1}: " + "  ".join(f"{k} {v:.3f}" for k, v in d.items()))

# ---------------------------------------------------------------- supporting detail for the flags
print("\n================ SUPPORT ================")
for y in [2010, 2011, 2012, 2013]:
    print(f"FY{y}: receivables {a(BSA, 'rec', y):.1f}  revenue {a(ISA, 'rev', y):.1f}  inventories {a(BSA, 'inv', y):.1f}  "
          f"AP {a(BSA, 'ap', y):.1f}  allowance {allow[y]:.1f}  provision {a(CFA, 'prov', y):.1f}  "
          f"vehicle-sale receivables (CF supplemental) {a(CFA, 'veh_rec', y):.1f}  prepaid/other {a(BSA, 'prepaid', y):.1f}  "
          f"restricted cash {a(BSA, 'rcash', y):.1f}  total assets {a(BSA, 'ta', y):.1f}")
print(f"receivables change FY2013: {a(BSA, 'rec', 2013) - a(BSA, 'rec', 2012):+.1f} ({growth(a(BSA, 'rec', 2013), a(BSA, 'rec', 2012)) * 100:+.1f}%)"
      f" ; revenue {growth(a(ISA, 'rev', 2013), a(ISA, 'rev', 2012)) * 100:+.1f}%")
print(f"receivables change FY2012: {growth(a(BSA, 'rec', 2012), a(BSA, 'rec', 2011)) * 100:+.1f}% ; revenue {growth(a(ISA, 'rev', 2012), a(ISA, 'rev', 2011)) * 100:+.1f}%")
print(f"allowance % FY2012 on revised basis (FY2013 annual report): {allow_2012_revised / (rec_2012_revised + allow_2012_revised) * 100:.2f}%")
print(f"allowance change FY2013 vs first-reported FY2012: {growth(allow[2013], allow[2012]) * 100:+.1f}% ; provision {growth(a(CFA, 'prov', 2013), a(CFA, 'prov', 2012)) * 100:+.1f}%")
print(f"prepaid/other growth FY2013: {growth(a(BSA, 'prepaid', 2013), a(BSA, 'prepaid', 2012)) * 100:+.1f}%")
for y in [2010, 2011, 2012, 2013]:
    print(f"FY{y}: cars accumulated depreciation / gross cars {-a(BSA, 'cars_ad', y) / a(BSA, 'cars', y) * 100:.1f}%  "
          f"gross fleet {a(BSA, 'cars', y) + a(BSA, 'oeq', y):.1f}  fleet depreciation {a(CFA, 'dep_ree', y):.1f}  "
          f"IS dep+lease/revenue {a(ISA, 'depree_is', y) / a(ISA, 'rev', y) * 100:.1f}%  direct operating/revenue {a(ISA, 'dirop', y) / a(ISA, 'rev', y) * 100:.1f}%  "
          f"SG&A/revenue {a(ISA, 'sga', y) / a(ISA, 'rev', y) * 100:.1f}%")
for q in ["Q4 2012", "Q4 2013"]:
    print(f"{q}: revenue {R['rev'][q]:.1f}  direct operating {R['dirop'][q]:.1f} ({R['dirop'][q] / R['rev'][q] * 100:.1f}%)  "
          f"dep+lease {R['depree_is'][q]:.1f} ({R['depree_is'][q] / R['rev'][q] * 100:.1f}%)  SG&A {R['sga'][q]:.1f} ({R['sga'][q] / R['rev'][q] * 100:.1f}%)  "
          f"pretax {R['pretax'][q]:.1f}  NI {R['ni'][q]:.1f}  OCF {R['ocf'][q]:.1f}")
print(f"Q4 growth: revenue {rg * 100:+.1f}%  direct operating {growth(R['dirop'][q], R['dirop'][q0]) * 100:+.1f}%  "
      f"dep+lease {growth(R['depree_is'][q], R['depree_is'][q0]) * 100:+.1f}%  fleet depreciation {growth(R['dep_ree'][q], R['dep_ree'][q0]) * 100:+.1f}%")
print(f"Gross margin change Q4: {(M['Gross margin']['Q4 2013'] - M['Gross margin']['Q4 2012']) * 100:+.2f} pts ; FY2013 vs FY2012 {(M['Gross margin'][2013] - M['Gross margin'][2012]) * 100:+.2f} pts")
print(f"Operating margin change Q4: {(M['Operating margin']['Q4 2013'] - M['Operating margin']['Q4 2012']) * 100:+.2f} pts")
for y in YEARS:
    print(f"FY{y}: OCF {M['OCF'][y]:.1f}  net fleet capex {M['Net fleet capex'][y]:.1f}  net P&E capex {M['Net P&E capex'][y]:.1f}  "
          f"FCF {M['Free cash flow'][y]:.1f}  NI {M['Net income'][y]:.1f}  FCF-NI {M['Free cash flow'][y] - M['Net income'][y]:+.1f}  acquisitions {M['Acquisitions'][y]:.1f}")
print(f"cumulative FY2011-13: FCF {sum(M['Free cash flow'][y] for y in [2011, 2012, 2013]):.1f}  NI {sum(M['Net income'][y] for y in [2011, 2012, 2013]):.1f}")
print(f"acquisitions Q4 2012 (CF quarterly): {-CFQ[E['acq']]['Q4 (ended 2012-12-31)']:.1f}")
print(f"total assets growth FY2012 {growth(a(BSA, 'ta', 2012), a(BSA, 'ta', 2011)) * 100:+.1f}%")

# ---------------------------------------------------------------- Revisions sheet
print("\n================ REVISIONS (income or equity) ================")
ws = wb["Revisions"]
rows = [r for r in ws.iter_rows(values_only=True)][1:]
keys = ("Net income", "Income before income taxes", "Total equity", "Accumulated deficit", "Total revenues",
        "Total expenses", "Net cash provided by (used in) operating activities", "Total assets",
        "Selling, general and administrative", "Depreciation of revenue earning equipment and lease charges", "Direct operating")
for r in rows:
    if r[0] and r[1] in keys and ("12 months" in r[3] or "as at" in r[3]) and r[0] != "Cash flow statement" or (r[0] == "Cash flow statement" and r[1].startswith("Net cash provided by (used in) operating") and "12 months" in r[3]):
        later = re.search(r"10-[KQ] for the period ended \d{4}-\d\d-\d\d", r[7]).group(0)
        print(f"  {r[0]:17s} {r[1][:58]:58s} {r[3]:27s} {r[4]:>10.1f} -> {r[6]:>10.1f}  ({r[8]:+.1f}, {r[8] / abs(r[4]) * 100:+.1f}%)  in {later}")
print(f"  rows on sheet: {len(rows)}; statements: {sorted(set(r[0] for r in rows if r[0]))}")
print(f"  'shown later in' filings: {sorted(set(re.search(r'10-[KQ] for the period ended [0-9-]+', r[7]).group(0) for r in rows if r[7]))}")

# quotations and page markers relied on
for txt, snip, name in [
    (rel13, "During the fourth quarter of 2013, management identified errors and revised previously issued financial statements as of and for the years ended December 31, 2012 and 2011 and the related interim periods.", "release FY2013"),
]:
    print(f"quote check [{name}] p. {page_of(txt, snip)}: found = {snip in txt}")

# ---------------------------------------------------------------- Q4 2012 growth gaps (vs Q4 2011) from the FY2012 results release
T1b = "CONDENSED CONSOLIDATED STATEMENTS OF OPERATIONS"
R12 = {}
for key, label in [("rev", "Total revenues"), ("dirop", "Direct operating"),
                   ("depree_is", "Depreciation of revenue earning equipment and lease charges")]:
    v, ln = line_after(rel12, T1b, label)
    R12[key] = {"Q4 2012": v[0], "Q4 2011": v[1]}
    print(f"FY2012 release: {ln[:90]}")
rg12 = growth(R12["rev"]["Q4 2012"], R12["rev"]["Q4 2011"])
cg12 = growth(R12["dirop"]["Q4 2012"] + R12["depree_is"]["Q4 2012"], R12["dirop"]["Q4 2011"] + R12["depree_is"]["Q4 2011"])
q11 = "Q4 (ended 2011-12-31)"
bq11 = lambda key: BSQ[E[key]][q11]
put("Rec gap", "Q4 2012", growth(bq("rec", "Q4 2012"), bq11("rec")) - rg12)
put("Inv gap", "Q4 2012", growth(bq("inv", "Q4 2012"), bq11("inv")) - cg12)
put("AP gap", "Q4 2012", growth(bq("ap", "Q4 2012"), bq11("ap")) - cg12)
print(f"Q4 2012 vs Q4 2011 (FY2012 release, as first reported): revenue {rg12 * 100:+.1f}%  cost {cg12 * 100:+.1f}%")

# extra support
for y in [2011, 2012, 2013]:
    g1 = a(BSA, "cars", y) + a(BSA, "oeq", y); g0 = a(BSA, "cars", y - 1) + a(BSA, "oeq", y - 1)
    print(f"FY{y}: gross fleet growth {growth(g1, g0) * 100:+.1f}%  AP growth {growth(a(BSA, 'ap', y), a(BSA, 'ap', y - 1)) * 100:+.1f}%  "
          f"inventory growth {growth(a(BSA, 'inv', y), a(BSA, 'inv', y - 1)) * 100:+.1f}%")
for y in [2009, 2010, 2011, 2012, 2013]:
    print(f"FY{y}: net fleet capex less fleet depreciation {M['Net fleet capex'][y] - a(CFA, 'dep_ree', y):+.1f}")
print("Dollar Thrifty date in FY2013 release p.", page_of(rel13, "Dollar Thrifty, which was acquired on November 19, 2012"))

# labels cited in the reply must exist in the workbook
for sh, lab in [(BSA, "Receivables, less allowance for doubtful accounts of $30.8 and $29.3"), (ISA, "Total revenues"),
                (CFA, "Sales of revenue earning equipment included in receivables"), (CFA, "Provision for losses on doubtful accounts"),
                (ISA, "Direct operating"), (ISA, "Depreciation of revenue earning equipment and lease charges"),
                (CFA, "Depreciation of revenue earning equipment"), (BSA, "Cars"), (BSA, "Other equipment"), (BSA, "Less accumulated depreciation"),
                (CFA, "Revenue earning equipment expenditures"), (CFA, "Proceeds from disposal of revenue earning equipment"),
                (CFA, "Acquisitions, net of cash acquired"), (CFA, "Net cash provided by (used in) operating activities"),
                (ISA, "Selling, general and administrative"), (BSA, "Accounts payable")]:
    assert any(v["_label"] == lab for v in sh.values()), lab
print("all cited labels found in workbook")

# ---------------------------------------------------------------- markdown table for the reply
def fmt(name, v):
    if v is None:
        return ""
    if name in ("DSO", "Days inventory", "Days payable"):
        return f"{v:.1f}"
    if name in ("Free cash flow", "Net income"):
        return f"{v:,.1f}".replace("-", "−")
    if name == "M-score":
        return f"{v:.2f}".replace("-", "−")
    if name == "Allowance %":
        return f"{v * 100:.2f}%"
    if name in ("Rec gap", "Inv gap", "AP gap"):
        return f"{v * 100:+.1f}".replace("-", "−")
    return f"{v * 100:.1f}%".replace("-", "−")

ROWS = [("Days sales outstanding", "DSO"), ("Days inventory", "Days inventory"), ("Days payable", "Days payable"),
        ("Receivables growth less revenue growth (pts)", "Rec gap"), ("Inventory growth less cost growth (pts)", "Inv gap"),
        ("Payables growth less cost growth (pts)", "AP gap"), ("Allowance, % of gross receivables", "Allowance %"),
        ("Gross margin", "Gross margin"), ("Operating margin", "Operating margin"),
        ("Accruals ÷ average total assets", "Accruals"), ("Free cash flow ($m)", "Free cash flow"), ("Net income ($m)", "Net income"),
        ("Depreciation rate, rental fleet", "Dep rate fleet"), ("Depreciation rate, property and equipment", "Dep rate P&E"),
        ("Beneish M-score (pair ending that year)", "M-score")]
hdr = ["Measure", "FY2009", "FY2010", "FY2011", "FY2012", "FY2013", "Q4 2012", "Q4 2013"]
print("\n| " + " | ".join(hdr) + " |")
print("|" + "---|" + "--:|" * (len(hdr) - 1))
for label, key in ROWS:
    print("| " + label + " | " + " | ".join(fmt(key, M[key].get(p)) for p in PERIODS) + " |")

# days of the acquired business inside Q4 2012 (acquisition date from the FY2013 release, p. 1)
from datetime import date
d = (date(2012, 12, 31) - date(2012, 11, 19)).days
print(f"\nacquired business in Q4 2012: {d} days = {d / 7:.1f} weeks of a 92-day quarter")
print(f"SG&A / revenue Q4 2012 {R['sga']['Q4 2012'] / R['rev']['Q4 2012'] * 100:.1f}%  Q4 2013 {R['sga']['Q4 2013'] / R['rev']['Q4 2013'] * 100:.1f}%")
print(f"M-scores all below -1.78: {all(BEN[y]['M'] < -1.78 for y in BEN)}")
