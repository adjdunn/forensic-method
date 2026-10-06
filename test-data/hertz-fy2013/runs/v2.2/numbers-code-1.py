# Turn 1: scan the numbers. Hertz Global Holdings, statements.xlsx (as originally reported).
# Run from the hertz-fy2013 folder:  python -X utf8 runs/v2.2/numbers-code-1.py
import re
import openpyxl

XLSX = "statements.xlsx"
AR12 = "annual-report-FY2012-12-31-key-sections.txt"
AR13 = "annual-report-FY2013-12-31-key-sections.txt"
ER13 = "earnings-release-Q4-FY2013-12-31.txt"

wb = openpyxl.load_workbook(XLSX, data_only=True)


def sheet(name):
    """Return {xbrl_element: {period_label: value}} and {xbrl: line label}. Keys use the XBRL element
    because some labels repeat (three 'Less accumulated depreciation' lines)."""
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    periods = [re.search(r"ended (\d{4}-\d{2}-\d{2})", str(h)).group(1) for h in hdr[1:-1]]
    data, labels = {}, {}
    for r in rows[1:]:
        if r[0] is None:
            continue
        key = r[-1]
        if key in data:  # same element used twice (cash at beginning / end): keep both by label
            key = f"{r[-1]}|{r[0]}"
        data[key] = {p: v for p, v in zip(periods, r[1:-1])}
        labels[key] = r[0]
    return data, labels


ISa, _ = sheet("IS annual")
ISq, _ = sheet("IS quarterly")
BSa, _ = sheet("BS annual")
BSq, _ = sheet("BS quarterly")
CFa, _ = sheet("CF annual")
CFq, _ = sheet("CF quarterly")

REV = "us-gaap:Revenues"
DO = "us-gaap:DirectOperatingCosts"
DEPLC = "htz:CostOfServicesDepreciationAndLeaseCharges"
SGA = "us-gaap:SellingGeneralAndAdministrativeExpense"
NI = "us-gaap:ProfitLoss"
AR = "us-gaap:AccountsReceivableNet"
INV = "us-gaap:InventoryNet"
AP = "us-gaap:AccountsPayableCurrentAndNoncurrent"
TA = "us-gaap:Assets"
TL = "us-gaap:Liabilities"
CARS = "us-gaap:PropertySubjectToOrAvailableForOperatingLeaseGross"
OEQ = "htz:OtherEquipmentGross"
CARS_AD = "us-gaap:PropertySubjectToOrAvailableForOperatingLeaseAccumulatedDepreciation"
REE_NET = "us-gaap:PropertySubjectToOrAvailableForOperatingLeaseNet"
PPE_G = "us-gaap:PropertyPlantAndEquipmentGross"
PPE_N = "us-gaap:PropertyPlantAndEquipmentNet"
CASH = "us-gaap:CashAndCashEquivalentsAtCarryingValue"
RCASH = "us-gaap:RestrictedCashAndCashEquivalents"
PREPAID = "us-gaap:PrepaidExpenseAndOtherAssets"
DEBT = "us-gaap:DebtAndCapitalLeaseObligations"
ACC_SAL = "us-gaap:EmployeeRelatedLiabilitiesCurrentAndNoncurrent / us-gaap:AccruedSalariesCurrent"
ACC_OTH = "us-gaap:OtherAccruedLiabilitiesCurrentAndNoncurrent"
ACC = "us-gaap:AccruedLiabilitiesCurrentAndNoncurrent"
TAXP = "us-gaap:TaxesPayableCurrentAndNoncurrent"
CFO = "us-gaap:NetCashProvidedByUsedInOperatingActivitiesContinuingOperations"
DEP_REE = "htz:CostOfServicesDepreciationAsAdjustedForGainLossUponDisposal"
DEP_PPE = "us-gaap:DepreciationNonproduction"
CAPEX_REE = "htz:PaymentsToAcquireRevenueEarningEquipment"
PROC_REE = "htz:ProceedsFromSaleOfRevenueEarningEquipment"
CAPEX_PPE = "us-gaap:PaymentsToAcquireOtherPropertyPlantAndEquipment"
ACQ = "us-gaap:PaymentsToAcquireBusinessesNetOfCashAcquired"
PROV = "us-gaap:ProvisionForDoubtfulAccounts"
REC_REE = "htz:SalesOfRevenueEarningEquipmentEarnedButNotReceived"

FY = ["2009-12-31", "2010-12-31", "2011-12-31", "2012-12-31", "2013-12-31"]
DAYS = {"2009-12-31": 365, "2010-12-31": 365, "2011-12-31": 365, "2012-12-31": 366, "2013-12-31": 365}
QDAYS = 92  # Oct-Dec


def cost(IS, p):  # cost line used: direct operating + depreciation of revenue earning equipment and lease charges
    return IS[DO][p] + IS[DEPLC][p]


def pct(x):
    return None if x is None else 100 * x


out = {}  # measure -> {column: value}


def put(measure, col, val):
    out.setdefault(measure, {})[col] = val


# ---------- allowance for doubtful accounts (annual reports; USD millions) ----------
# FY2012 10-K balance sheet p.72: $25,113 and $20,282 (thousands). FY2013 10-K balance sheet p.75: $30.8 and $29.3.
# FY2010 year-end balance from Schedule II of the FY2012 10-K, p.140: 19,708.
ALLOW = {"2010-12-31": 19.708, "2011-12-31": 20.282, "2012-12-31": 25.113, "2013-12-31": 30.8}
ALLOW_2012_REVISED = 29.3


def page_at(text, pos):
    m = None
    for m in re.finditer(r"\[p\. ([^\]]+)\]", text[:pos]):
        pass
    return m.group(1) if m else None


def check_source(fname, needle, page):
    t = open(fname, encoding="utf-8").read()
    i = t.find(needle)
    ok = i >= 0 and page_at(t, i) == str(page)
    print(f"  source check [{fname} p.{page}] {'OK' if ok else 'FAIL'}: {needle[:70]}")
    return ok


print("Allowance source checks")
check_source(AR12, "Receivables, less allowance for doubtful accounts of $25,113 and $20,282", 72)
check_source(AR13, "Receivables, less allowance for doubtful accounts of $30.8 and $29.3", 75)
check_source(AR12, "Year ended December 31, 2010 | 21,268 | 19,667 | (695) | (20,532) | (a) | 19,708", 140)

# ---------- annual measures ----------
for i, p in enumerate(FY):
    col = "FY" + p[:4]
    prev = FY[i - 1] if i > 0 else None
    rev, c = ISa[REV][p], cost(ISa, p)
    put("Gross margin %", col, pct((rev - c) / rev))
    put("Operating margin %", col, pct((rev - c - ISa[SGA][p]) / rev))
    put("SG&A % revenue", col, pct(ISa[SGA][p] / rev))
    put("Fleet depreciation and lease charges % revenue", col, pct(ISa[DEPLC][p] / rev))
    net_fleet_capex = -CFa[CAPEX_REE][p] - CFa[PROC_REE][p]
    fcf = CFa[CFO][p] - net_fleet_capex - (-CFa[CAPEX_PPE][p])
    put("Free cash flow $m", col, fcf)
    put("Net income $m", col, ISa[NI][p])
    put("Operating cash flow $m", col, CFa[CFO][p])
    put("Acquisitions outflow $m", col, -CFa[ACQ][p])
    if p in BSa[AR] and BSa[AR][p] is not None:
        put("DSO days", col, BSa[AR][p] / rev * DAYS[p])
        put("Days inventory", col, BSa[INV][p] / c * DAYS[p])
        put("Days payable", col, BSa[AP][p] / c * DAYS[p])
        put("Allowance % gross receivables", col, pct(ALLOW[p] / (BSa[AR][p] + ALLOW[p])))
        put("Accum. depreciation % gross cars", col, pct(-BSa[CARS_AD][p] / BSa[CARS][p]))
    if prev and prev in BSa[AR] and BSa[AR][prev] is not None:
        g = lambda a, b: a / b - 1
        rg, cg = g(rev, ISa[REV][prev]), g(c, cost(ISa, prev))
        put("Revenue growth %", col, pct(rg))
        put("Receivables growth - revenue growth, pts", col, pct(g(BSa[AR][p], BSa[AR][prev]) - rg))
        put("Inventory growth - cost growth, pts", col, pct(g(BSa[INV][p], BSa[INV][prev]) - cg))
        put("Payables growth - cost growth, pts", col, pct(g(BSa[AP][p], BSa[AP][prev]) - cg))
        avg_ta = (BSa[TA][p] + BSa[TA][prev]) / 2
        put("Accruals % avg assets", col, pct((ISa[NI][p] - CFa[CFO][p]) / avg_ta))
        avg_ree = (BSa[CARS][p] + BSa[OEQ][p] + BSa[CARS][prev] + BSa[OEQ][prev]) / 2
        put("Depreciation rate, rental equipment %", col, pct(CFa[DEP_REE][p] / avg_ree))
        avg_ppe = (BSa[PPE_G][p] + BSa[PPE_G][prev]) / 2
        put("Depreciation rate, property and equipment %", col, pct(CFa[DEP_PPE][p] / avg_ppe))
        put("Prepaid and other assets growth %", col, pct(g(BSa[PREPAID][p], BSa[PREPAID][prev])))
        put("Gross cars growth %", col, pct(g(BSa[CARS][p], BSa[CARS][prev])))
        put("Accum. depreciation on cars growth %", col, pct(g(BSa[CARS_AD][p], BSa[CARS_AD][prev])))
        put("Provision for doubtful accounts growth %", col, pct(g(CFa[PROV][p], CFa[PROV][prev])))


# ---------- Beneish M-score (annual pairs) ----------
def accrued_total(p):
    if p == "2010-12-31":
        return BSa[ACC_SAL][p] + BSa[ACC_OTH][p]
    if p == "2011-12-31":
        return BSa[ACC][p]  # 1,128.5 = salaries 426.7 + other 701.8
    if p == "2012-12-31":
        return BSa[ACC][p] + BSa[ACC_SAL][p]  # 739.6 + 440.9 (shown later as one line, 1,180.5)
    return BSa[ACC][p]


def beneish(p, q):  # p = current year, q = prior year
    r1, r0 = ISa[REV][p], ISa[REV][q]
    dsri = (BSa[AR][p] / r1) / (BSa[AR][q] / r0)
    gm1, gm0 = (r1 - cost(ISa, p)) / r1, (r0 - cost(ISa, q)) / r0
    gmi = gm0 / gm1

    def aq(x):  # 1 - (current-type assets + net rental equipment + net property) / total assets
        ca = BSa[CASH][x] + BSa[RCASH][x] + BSa[AR][x] + BSa[INV][x] + BSa[PREPAID][x]
        return 1 - (ca + BSa[REE_NET][x] + BSa[PPE_N][x]) / BSa[TA][x]

    aqi = aq(p) / aq(q)
    sgi = r1 / r0

    def deprate(x):
        d = CFa[DEP_REE][x] + CFa[DEP_PPE][x]
        return d / (d + BSa[REE_NET][x] + BSa[PPE_N][x])

    depi = deprate(q) / deprate(p)
    sgai = (ISa[SGA][p] / r1) / (ISa[SGA][q] / r0)
    tata = (ISa[NI][p] - CFa[CFO][p]) / BSa[TA][p]

    def lev(x):  # debt + payables + accrued liabilities + accrued taxes, over total assets
        return (BSa[DEBT][x] + BSa[AP][x] + accrued_total(x) + BSa[TAXP][x]) / BSa[TA][x]

    lvgi = lev(p) / lev(q)
    m = (-4.84 + 0.920 * dsri + 0.528 * gmi + 0.404 * aqi + 0.892 * sgi + 0.115 * depi
         - 0.172 * sgai + 4.679 * tata - 0.327 * lvgi)
    return m, dict(DSRI=dsri, GMI=gmi, AQI=aqi, SGI=sgi, DEPI=depi, SGAI=sgai, TATA=tata, LVGI=lvgi)


print("\nBeneish")
for p, q in [("2011-12-31", "2010-12-31"), ("2012-12-31", "2011-12-31"), ("2013-12-31", "2012-12-31")]:
    m, parts = beneish(p, q)
    put("Beneish M-score", "FY" + p[:4], m)
    print(" ", p[:4], round(m, 2), {k: round(v, 3) for k, v in parts.items()})

# ---------- latest quarter beside the same quarter a year earlier ----------
for p, prev_q_end, yago in [("2012-12-31", "2012-09-30", "2011-12-31"), ("2013-12-31", "2013-09-30", "2012-12-31")]:
    col = "Q4 " + p[:4]
    rev, c = ISq[REV][p], cost(ISq, p)
    put("Gross margin %", col, pct((rev - c) / rev))
    put("Operating margin %", col, pct((rev - c - ISq[SGA][p]) / rev))
    put("SG&A % revenue", col, pct(ISq[SGA][p] / rev))
    put("Fleet depreciation and lease charges % revenue", col, pct(ISq[DEPLC][p] / rev))
    put("DSO days", col, BSq[AR][p] / rev * QDAYS)
    put("Days inventory", col, BSq[INV][p] / c * QDAYS)
    put("Days payable", col, BSq[AP][p] / c * QDAYS)
    put("Allowance % gross receivables", col, pct(ALLOW[p] / (BSq[AR][p] + ALLOW[p])))
    g = lambda a, b: a / b - 1
    rg, cg = g(rev, ISq[REV][yago]), g(c, cost(ISq, yago))
    put("Revenue growth %", col, pct(rg))
    put("Receivables growth - revenue growth, pts", col, pct(g(BSq[AR][p], BSq[AR][yago]) - rg))
    put("Inventory growth - cost growth, pts", col, pct(g(BSq[INV][p], BSq[INV][yago]) - cg))
    put("Payables growth - cost growth, pts", col, pct(g(BSq[AP][p], BSq[AP][yago]) - cg))
    avg_ta = (BSq[TA][p] + BSq[TA][prev_q_end]) / 2
    put("Accruals % avg assets", col, pct((ISq[NI][p] - CFq[CFO][p]) / avg_ta))
    net_fleet_capex = -CFq[CAPEX_REE][p] - CFq[PROC_REE][p]
    put("Free cash flow $m", col, CFq[CFO][p] - net_fleet_capex - (-CFq[CAPEX_PPE][p]))
    put("Net income $m", col, ISq[NI][p])
    put("Operating cash flow $m", col, CFq[CFO][p])
    put("Acquisitions outflow $m", col, -CFq[ACQ][p])
    avg_ree = (BSq[CARS][p] + BSq[OEQ][p] + BSq[CARS][prev_q_end] + BSq[OEQ][prev_q_end]) / 2
    put("Depreciation rate, rental equipment %", col, pct(4 * CFq[DEP_REE][p] / avg_ree))  # annualised x4
    avg_ppe = (BSq[PPE_G][p] + BSq[PPE_G][prev_q_end]) / 2
    put("Depreciation rate, property and equipment %", col, pct(4 * CFq[DEP_PPE][p] / avg_ppe))
    put("SG&A growth %", col, pct(g(ISq[SGA][p], ISq[SGA][yago])))
    put("Direct operating growth %", col, pct(g(ISq[DO][p], ISq[DO][yago])))

cols = ["FY2009", "FY2010", "FY2011", "FY2012", "FY2013", "Q4 2012", "Q4 2013"]
print("\nMEASURES")
for mname, vals in out.items():
    print(f"{mname:48s}", " | ".join(f"{vals[c]:9.2f}" if c in vals and vals[c] is not None else "        -" for c in cols))

# ---------- supporting figures quoted in the flags ----------
print("\nSUPPORT")
g = lambda a, b: 100 * (a / b - 1)
print("Receivables FY12->FY13:", BSa[AR]["2012-12-31"], BSa[AR]["2013-12-31"], round(g(BSa[AR]["2013-12-31"], BSa[AR]["2012-12-31"]), 1), "%")
print("Revenue FY12->FY13 growth %:", round(g(ISa[REV]["2013-12-31"], ISa[REV]["2012-12-31"]), 1))
print("Cost line FY12->FY13 growth %:", round(g(cost(ISa, "2013-12-31"), cost(ISa, "2012-12-31")), 1))
print("Payables FY12->FY13:", BSa[AP]["2012-12-31"], BSa[AP]["2013-12-31"], round(g(BSa[AP]["2013-12-31"], BSa[AP]["2012-12-31"]), 1), "%")
print("Fleet sales in receivables (CF annual):", {p[:4]: CFa[REC_REE][p] for p in FY})
print("Prepaid and other assets:", {p[:4]: BSa[PREPAID][p] for p in FY[1:]})
print("Gross cars:", {p[:4]: BSa[CARS][p] for p in FY[1:]}, "Accum dep cars:", {p[:4]: BSa[CARS_AD][p] for p in FY[1:]})
print("Dep of REE (CF annual):", {p[:4]: CFa[DEP_REE][p] for p in FY})
print("Provision for doubtful accounts:", {p[:4]: CFa[PROV][p] for p in FY})
print("Allowance % gross receivables at 2012 on the revised figures (29.3 / (1,879.7+29.3)):",
      round(100 * ALLOW_2012_REVISED / (1879.7 + ALLOW_2012_REVISED), 2))
print("Q4 SG&A:", ISq[SGA]["2012-12-31"], ISq[SGA]["2013-12-31"], " Q4 direct operating:", ISq[DO]["2012-12-31"], ISq[DO]["2013-12-31"])
print("Q4 revenue:", ISq[REV]["2012-12-31"], ISq[REV]["2013-12-31"])

# ---------- Q4 2013 by subtraction: does it mix original and revised figures? ----------
# FY2013 10-K Note 2 (p.87) revises the nine months ended 30 Sept 2013. The spreadsheet's Q4 2013 is the FY2013 10-K
# less the nine months as first reported in the Q3 10-Q, so the nine-month adjustments land in the derived quarter.
# The results release (Table 1) gives the quarter on the revised basis.
print("\nQ4 2013 derived (spreadsheet) vs company's revised-basis quarter (release Table 1)")
rel13 = {"rev": 2556.3, "do": 1439.3, "dep": 643.0, "sga": 222.2, "ni": -0.6}
rel12 = {"rev": 2319.7, "do": 1257.7, "dep": 550.0, "sga": 331.1, "ni": -36.8}
check_source(ER13, "Total revenues | $2,556.3 | $2,319.7 | 100.0% | 100.0%", "unnumbered")
check_source(ER13, "Direct operating | 1,439.3 | 1,257.7 | 56.3% | 54.2%", "unnumbered")
check_source(ER13, "Depreciation of revenue earning equipment and lease charges | 643.0 | 550.0 | 25.2% | 23.7%", "unnumbered")
check_source(ER13, "Selling, general and administrative | 222.2 | 331.1 | 8.7% | 14.3%", "unnumbered")
check_source(AR13, "Depreciation of revenue earning equipment and lease charges | 676.7 | (1.8) | 674.9 | 1,904.8 | (h) | (22.3) | 1,882.5", 87)
for name, r in [("Q4 2013", rel13), ("Q4 2012", rel12)]:
    gm = 100 * (r["rev"] - r["do"] - r["dep"]) / r["rev"]
    om = 100 * (r["rev"] - r["do"] - r["dep"] - r["sga"]) / r["rev"]
    print(f"  release {name}: gross margin {gm:.1f}%, operating margin {om:.1f}%, fleet dep % revenue {100*r['dep']/r['rev']:.1f}%")
p = "2013-12-31"
print("  derived Q4 2013 lines:", ISq[REV][p], ISq[DO][p], ISq[DEPLC][p], ISq[SGA][p], ISq[NI][p])
print("  differences (derived - release): revenue %.1f, direct operating %.1f, fleet depreciation %.1f, SG&A %.1f, net income %.1f" % (
    ISq[REV][p] - rel13["rev"], ISq[DO][p] - rel13["do"], ISq[DEPLC][p] - rel13["dep"], ISq[SGA][p] - rel13["sga"], ISq[NI][p] - rel13["ni"]))
# rental-equipment depreciation rate for Q4 2013 on the revised basis: cash-flow depreciation of revenue earning
# equipment for the nine months was 1,847.1 as first reported (FY 2,445.0 less derived Q4 597.9); the nine-month
# income-statement line was revised by -22.3 (p.87).
nine_m_orig = CFa[DEP_REE][p] - CFq[DEP_REE][p]
q4_dep_rev = CFa[DEP_REE][p] - (nine_m_orig - 22.3)
avg_ree = (BSq[CARS][p] + BSq[OEQ][p] + BSq[CARS]["2013-09-30"] + BSq[OEQ]["2013-09-30"]) / 2
print("  Q4 2013 fleet depreciation (cash flow line) derived %.1f; adding back the 22.3 nine-month adjustment %.1f; annualised rate %.1f%% vs %.1f%%" % (
    CFq[DEP_REE][p], q4_dep_rev, 400 * CFq[DEP_REE][p] / avg_ree, 400 * q4_dep_rev / avg_ree))

# ---------- Revisions sheet: changes that moved income or equity ----------
print("\nREVISIONS that moved income or equity")
ws = wb["Revisions"]
keys = ("us-gaap:ProfitLoss", "us-gaap:IncomeLossFromContinuingOperationsBeforeIncomeTaxes", "us-gaap:StockholdersEquity",
        "us-gaap:RetainedEarningsAccumulatedDeficit", "us-gaap:Revenues", "us-gaap:DirectOperatingCosts",
        "htz:CostOfServicesDepreciationAndLeaseCharges", "us-gaap:SellingGeneralAndAdministrativeExpense",
        "us-gaap:NetCashProvidedByUsedInOperatingActivitiesContinuingOperations")
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] in ("Income statement", "Balance sheet", "Cash flow") and any(str(r[2]).startswith(k) for k in keys) and ("12 months" in r[3] or "as at" in r[3]):
        if r[0] == "Cash flow" and r[2] == "us-gaap:ProfitLoss":
            continue
        print(f"  {r[0]:16s} {r[1][:60]:60s} {r[3]:28s} {r[4]:>10.1f} -> {r[6]:>10.1f}  ({r[8]:+.1f})  {r[7][:40]}")

# ---------- page checks for the figures the reply cites from the text files ----------
print()
print("PAGE CHECKS for the reply")
check_source(AR13, "Income before income taxes | 328.3 | (7.0) | 321.3 | 612.4 | (11.6) | 600.8", 87)
check_source(AR13, "Direct operating | 1,525.4 | 14.1 | 1,539.5 | 4,282.6 | (g) | 30.1 | 4,312.7", 87)
print("Derived Q4 2013 pre-tax %.1f vs release 62.3: difference %.1f" % (
    ISq["us-gaap:IncomeLossFromContinuingOperationsBeforeIncomeTaxesMinorityInterestAndIncomeLossFromEquityMethodInvestments"]["2013-12-31"],
    ISq["us-gaap:IncomeLossFromContinuingOperationsBeforeIncomeTaxesMinorityInterestAndIncomeLossFromEquityMethodInvestments"]["2013-12-31"] - 62.3))
print("The reply contains no quotations from the text files (line labels from the spreadsheet only).")
