"""Writes pack-template/statements-template.xlsx from the line names in PACK-SPEC.md.

    python make_template.py

The line names here must match PACK-SPEC.md exactly; the prompts compute by name.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = Path(__file__).parent
OUT = HERE / "pack-template" / "statements-template.xlsx"

IS = ["Revenue", "Cost of goods sold", "Gross profit", "Selling, general and administrative",
      "Research and development", "Depreciation (excluding amortization of intangibles)",
      "Depreciation of equipment held for rental or lease",
      "Amortization of intangibles", "Restructuring and other charges", "Operating income",
      "Interest expense", "Interest income", "Other income (expense)", "Pre-tax income",
      "Income tax expense", "Income from continuing operations", "Discontinued operations",
      "Net income", "Diluted shares", "Diluted EPS"]
BS = ["Cash and equivalents", "Short-term investments", "Trade receivables, net",
      "Allowance for credit losses", "Contract assets", "Inventory, net", "Other current assets",
      "Total current assets", "PP&E, gross", "Accumulated depreciation", "PP&E, net",
      "Equipment held for rental or lease, gross",
      "Accumulated depreciation, equipment held for rental or lease",
      "Equipment held for rental or lease, net",
      "Right-of-use assets", "Goodwill", "Intangible assets, net",
      "Capitalized software or development costs", "Capitalized contract costs",
      "Long-term investments and advances", "Other non-current assets", "Total assets",
      "Trade payables", "Supplier finance obligations", "Accrued liabilities",
      "Contract liabilities", "Provisions, current", "Income taxes payable",
      "Current portion of long-term debt", "Current lease liabilities",
      "Total current liabilities", "Long-term debt", "Non-current lease liabilities",
      "Provisions, non-current", "Other non-current liabilities", "Total liabilities",
      "Preferred stock", "Common equity", "Non-controlling interests", "Total equity"]
CF = ["Net income", "Depreciation and amortization", "Stock-based compensation",
      "Change in receivables", "Change in inventory", "Change in payables",
      "Change in accrued liabilities", "Change in contract balances", "Other operating changes",
      "Cash from operations", "Capital expenditures (PP&E)",
      "Purchases of equipment held for rental or lease",
      "Proceeds from disposal of equipment held for rental or lease",
      "Capitalized software or development costs", "Acquisitions, net of cash",
      "Proceeds from disposals", "Cash from investing", "Debt issued", "Debt repaid",
      "Stock issued", "Stock repurchased", "Dividends paid", "Lease principal payments",
      "Cash from financing", "Interest paid", "Interest paid classified in (Operating / Financing)",
      "Interest received", "Interest received classified in (Operating / Investing)",
      "Dividends received", "Dividends received classified in (Operating / Investing)",
      "Dividends paid classified in (Operating / Financing)",
      "Income taxes paid", "Effect of exchange rates", "Net change in cash"]
NOTES = ["Allowance: opening", "Allowance: provision", "Allowance: write-offs",
         "Allowance: recoveries", "Allowance: other", "Allowance: closing",
         "Provision matrix (IFRS 7.35M/35N): see PROVISION_MATRIX sheet",
         "Inventory: raw materials", "Inventory: work in process", "Inventory: finished goods",
         "Inventory: supplies", "Inventory: reserve or write-down allowance",
         "Inventory: write-downs recognized", "Inventory: reversals of write-downs",
         "Inventories recognized as expense (IAS 2.36(d))",
         "LIFO reserve", "LIFO liquidation effect",
         "Supplier finance: confirmed obligations outstanding", "Supplier finance: added",
         "Supplier finance: settled", "Supplier finance: paid by finance providers",
         "Supplier finance: non-cash changes", "Supplier finance: payment terms (text)",
         "Supplier finance: range of payment due dates (text)",
         "Trade payables: range of payment due dates (text)",
         "Receivables derecognized: amount outstanding",
         "Receivables transferred, not derecognized", "Liability associated with transferred receivables",
         "Interest incurred", "Interest capitalized",
         "Depreciation of right-of-use assets", "Interest on lease liabilities",
         "Warranty: opening", "Warranty: accruals", "Warranty: changes for pre-existing",
         "Warranty: payments", "Warranty: closing", "Restructuring: opening",
         "Restructuring: charges", "Restructuring: payments", "Restructuring: adjustments",
         "Restructuring: closing", "Provisions (all classes): see PROVISIONS sheet",
         "Revenue from obligations satisfied in prior periods",
         "Revenue recognized from opening contract liabilities",
         "Development: additions from internal development", "Research and development expensed",
         "Amortization of internally generated intangibles", "Employees at year end"]
PROVISIONS = ["Class", "Opening", "Additions", "Used", "Unused amounts reversed",
              "Unwinding of discount", "Other", "Closing"]
MATRIX = ["Past-due band", "Gross carrying amount", "Expected loss rate", "Loss allowance"]

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="223A55")


def sheet(wb, name, lines, periods, note):
    ws = wb.create_sheet(name)
    ws.append(["Line"] + periods)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for line in lines:
        ws.append([line] + [None] * len(periods))
    ws.column_dimensions["A"].width = 52
    ws.freeze_panes = "B2"
    ws.cell(row=len(lines) + 3, column=1, value=note).font = Font(italic=True, color="7A8090")


def main():
    wb = Workbook()
    wb.remove(wb.active)
    q = [f"Q{i}" for i in range(1, 13)]
    a = [f"FY{i}" for i in range(1, 6)]
    qnote = "Replace Q1..Q12 with the company's fiscal quarter labels, oldest first. Cash flow is discrete quarters."
    anote = "Replace FY1..FY5 with fiscal year labels, oldest first."
    sheet(wb, "IS_Q", IS, q, qnote)
    sheet(wb, "BS_Q", BS, q, qnote)
    sheet(wb, "CF_Q", CF, q, qnote)
    sheet(wb, "IS_A", IS, a, anote)
    sheet(wb, "BS_A", BS, a, anote)
    sheet(wb, "CF_A", CF, a, anote)
    sheet(wb, "NOTES", NOTES, q + a, "Quarterly where disclosed; otherwise annual columns only. Leave blank if not reported.")
    ng = wb.create_sheet("NONGAAP")
    ng.append(["Measure", "Reconciling line (company's label)"] + q)
    for c in ng[1]:
        c.font, c.fill = HEAD, FILL
    ng.column_dimensions["A"].width = 28
    ng.column_dimensions["B"].width = 46
    for name, cols, note in (
        ("PROVISIONS", PROVISIONS, "IFRS: one row per class of provision per fiscal year (IAS 37.84). Add a Year column value in Class, e.g. 'Warranty FY2025'."),
        ("PROVISION_MATRIX", MATRIX, "IFRS: one block per year-end, from the credit-risk note (IFRS 7.35M/35N). Prefix the band with the year."),
    ):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = HEAD, FILL
        ws.column_dimensions["A"].width = 34
        ws.cell(row=3, column=1, value=note).font = Font(italic=True, color="7A8090")
    ex = wb.create_sheet("EXTRA")
    ex.append(["Statement", "Company label"] + q + a)
    for c in ex[1]:
        c.font, c.fill = HEAD, FILL
    ex.column_dimensions["B"].width = 52
    rs = wb.create_sheet("RESTATED")
    rs.append(["Line", "Year restated", "As shown in the later annual report", "Later report (file ID)"])
    for c in rs[1]:
        c.font, c.fill = HEAD, FILL
    rs.column_dimensions["A"].width = 46
    rs.cell(row=3, column=1, value="Optional. The _A sheets hold figures as originally reported; record here the comparative as reprinted a year later.").font = Font(italic=True, color="7A8090")
    loc = wb.create_sheet("LOCATORS")
    loc.append(["Sheet", "Line", "Period", "Source file ID", "Page or note", "Derived (Y/N)", "How derived"])
    for c in loc[1]:
        c.font, c.fill = HEAD, FILL
    for col, w in zip("ABCDEFG", (10, 46, 12, 14, 18, 12, 40)):
        loc.column_dimensions[col].width = w
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
